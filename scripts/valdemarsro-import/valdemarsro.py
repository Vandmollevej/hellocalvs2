from __future__ import annotations

import re
import sys
import unicodedata
from dataclasses import dataclass, field
from fractions import Fraction
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from playwright.sync_api import BrowserContext, Page, sync_playwright


# ============================================================
# Fixed Hello Cal scraper structure
# ============================================================
RETAILER = "valdemarsro"
ROOT_DIR = Path(r"C:\Users\Peter\Desktop\Hello Cal\Productdatabase\Valdemarsro")
IMAGES_DIR = ROOT_DIR / "Images"
TEMP_DIR = ROOT_DIR / "Temp"
CHECKPOINT_PATTERN = f"{RETAILER}_checkpoint_*.xlsx"
RECIPE_FILE = ROOT_DIR / f"{RETAILER}.xlsx"
INGREDIENT_FILE = ROOT_DIR / f"{RETAILER}_ingredients.xlsx"

BASE_URL = "https://www.valdemarsro.dk"
# The recipe overview lists every recipe category. Categories are discovered
# from it (and from each category page) and scraped one at a time as found.
RECIPES_ROOT_URL = f"{BASE_URL}/opskrifter/"
CHECKPOINT_EVERY = 20
HEADLESS = True
PAGE_TIMEOUT_MS = 45_000
NAVIGATION_TIMEOUT_MS = 60_000
REQUEST_TIMEOUT_MS = 45_000
MAX_RETRIES = 3
MAX_PAGES_PER_CATEGORY = 200
POLITE_DELAY_MS = 600

# No recipe limit. This scraper is intentionally configured to collect all discoverable recipes.
#
# Nutrition (kcal etc.) is only shown to Valdemarsro Premium members and is
# NOT scraped. Recipes get their nutrition from valdemarsro_match.py, which
# matches the ingredients against Frida and the other product databases.

# Site pages that are linked like categories but are not recipe listings.
NON_CATEGORY_SLUGS = {
    "opskrifter", "opskriftsoversigt", "madplan", "online-indkoebsliste", "valdemarsro-premium",
    "raavare-indeks", "ofte-stillede-spoergsmaal-til-valdemarsro-opskrifter", "kogeboeger",
    "kogebogsindeks", "boger", "boerneboeger-anbefaling", "indretning", "diy", "anmeldereksemplar",
    "inaktive-indlaeg", "i-haven",
}

RECIPE_HEADERS = [
    "HelloCal_Title",
    "Recipe Name",
    "Category",
    "Servings",
    "Servings Unit",
    "Total Time",
    "Work Time",
    "Freezable",
    "Ingredient Count",
    "Ingredients",
    "Instructions",
    "Image File",
    "Image URL",
    "Source URL",
    "Parse Status",
]

INGREDIENT_HEADERS = [
    "HelloCal_Title",
    "Line No",
    "Section",
    "Raw Text",
    "Quantity",
    "Unit",
    "Ingredient Name",
    "Note",
    "Source URL",
]


@dataclass
class ScrapeState:
    recipes: list[dict[str, Any]] = field(default_factory=list)
    ingredients: list[dict[str, Any]] = field(default_factory=list)
    category_queue: list[str] = field(default_factory=list)
    seen_categories: set[str] = field(default_factory=set)
    processed_urls: set[str] = field(default_factory=set)
    skipped_urls: set[str] = field(default_factory=set)
    failed_urls: set[str] = field(default_factory=set)


# ============================================================
# General helpers
# ============================================================
def ensure_directories() -> None:
    ROOT_DIR.mkdir(parents=True, exist_ok=True)
    IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    TEMP_DIR.mkdir(parents=True, exist_ok=True)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def safe_filename(text: str) -> str:
    text = unicodedata.normalize("NFKD", clean_text(text))
    text = text.encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9._-]+", "-", text).strip("-._")
    return text[:120] or "recipe"


def dedupe_preserve(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def canonicalize_url(url: str) -> str:
    parsed = urlparse(urljoin(BASE_URL, url))
    path = parsed.path if parsed.path.endswith("/") else parsed.path + "/"
    return f"https://www.valdemarsro.dk{path}"


def slug_of(url: str) -> str:
    return urlparse(url).path.strip("/").split("/")[0]


def is_category_url(url: str) -> bool:
    path = urlparse(url).path.strip("/")
    return bool(path) and "/" not in path and path not in NON_CATEGORY_SLUGS


# ============================================================
# Ingredient line parsing (shared with valdemarsro_match.py)
# ============================================================
UNICODE_FRACTIONS = {"½": "1/2", "¼": "1/4", "¾": "3/4", "⅓": "1/3", "⅔": "2/3", "⅛": "1/8"}

# Danish units as written on Valdemarsro -> canonical unit.
UNIT_ALIASES = {
    "g": "g", "gr": "g", "gr.": "g", "gram": "g",
    "kg": "kg", "kilo": "kg",
    "mg": "mg",
    "ml": "ml", "cl": "cl", "dl": "dl", "l": "l", "liter": "l",
    "spsk": "tbsp", "spsk.": "tbsp", "spiseskefuld": "tbsp", "spiseskefulde": "tbsp",
    "tsk": "tsp", "tsk.": "tsp", "teskefuld": "tsp", "teskefulde": "tsp",
    "knsp": "pinch", "knsp.": "pinch", "knivspids": "pinch", "nip": "pinch",
    "stk": "pcs", "stk.": "pcs", "styk": "pcs",
    "fed": "clove",
    "dåse": "can", "dåser": "can", "ds": "can", "ds.": "can",
    "pakke": "pack", "pakker": "pack", "pk": "pack", "pk.": "pack", "bakke": "pack", "bakker": "pack",
    "pose": "bag", "poser": "bag",
    "bundt": "bunch", "bundter": "bunch", "bdt": "bunch", "bdt.": "bunch",
    "håndfuld": "handful", "håndfulde": "handful", "håndfuldt": "handful",
    "stængel": "stalk", "stængler": "stalk",
    "skive": "slice", "skiver": "slice",
    "glas": "jar", "flaske": "bottle", "flasker": "bottle",
    "blad": "leaf", "blade": "leaf",
    "kvist": "sprig", "kviste": "sprig",
    "drys": "pinch", "sjat": "splash", "stænk": "splash",
}

QUANTITY_RE = re.compile(
    r"^(?:ca\.?\s*|cirka\s+|omkring\s+)?"
    r"(?P<q>\d+\s+\d+/\d+|\d+/\d+|\d+(?:[.,]\d+)?)(?![\d/])"
    r"(?:\s*(?:-|–|til)\s*(?P<q2>\d+/\d+|\d+(?:[.,]\d+)?))?\s*",
    re.IGNORECASE,
)


def parse_number(text: str) -> float:
    total = 0.0
    for part in text.replace(",", ".").split():
        total += float(Fraction(part)) if "/" in part else float(part)
    return total


def parse_ingredient_line(raw: str) -> dict[str, Any]:
    """Split "2 spsk olivenolie, gerne ekstra jomfru" into quantity, unit,
    ingredient name and note. A range ("2-3 løg") uses its midpoint."""
    text = clean_text(raw)
    for symbol, replacement in UNICODE_FRACTIONS.items():
        text = re.sub(rf"(\d)\s*{symbol}", rf"\1 {replacement}", text)
        text = text.replace(symbol, replacement)

    quantity: float | str = ""
    match = QUANTITY_RE.match(text)
    if match:
        low = parse_number(match.group("q"))
        high = parse_number(match.group("q2")) if match.group("q2") else low
        quantity = round((low + high) / 2, 3)
        text = text[match.end():]

    unit = ""
    first, _, rest = text.partition(" ")
    if first.lower() in UNIT_ALIASES and (quantity != "" or first.lower() in {"knsp", "knivspids", "nip", "drys", "sjat"}):
        unit = UNIT_ALIASES[first.lower()]
        text = rest

    name, _, note = text.partition(",")
    return {
        "Quantity": quantity,
        "Unit": unit,
        "Ingredient Name": clean_text(name),
        "Note": clean_text(note),
    }


# ============================================================
# Workbook helpers
# ============================================================
def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def append_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    style_header(ws)


def save_workbook(wb: Workbook, path: Path) -> bool:
    try:
        wb.save(path)
        return True
    except PermissionError:
        print(f"Could not save {path.name} - file is open elsewhere. Skipped.")
        return False


def save_final_files(state: ScrapeState) -> None:
    wb = Workbook()
    append_sheet(wb, "Recipes", RECIPE_HEADERS, state.recipes, first=True)
    save_workbook(wb, RECIPE_FILE)
    wb = Workbook()
    append_sheet(wb, "Ingredients", INGREDIENT_HEADERS, state.ingredients, first=True)
    save_workbook(wb, INGREDIENT_FILE)


def save_temp(state: ScrapeState) -> None:
    wb = Workbook()
    append_sheet(wb, "Recipes", RECIPE_HEADERS, state.recipes, first=True)
    append_sheet(wb, "Ingredients", INGREDIENT_HEADERS, state.ingredients)
    ws_state = wb.create_sheet("State")
    ws_state.append(["Type", "URL"])
    for url in state.category_queue:
        ws_state.append(["category_queued", url])
    for kind, urls in (
        ("category_seen", state.seen_categories),
        ("processed", state.processed_urls),
        ("skipped", state.skipped_urls),
        ("failed", state.failed_urls),
    ):
        for url in sorted(urls):
            ws_state.append([kind, url])
    style_header(ws_state)

    checkpoint_file = TEMP_DIR / f"{RETAILER}_checkpoint_{len(state.recipes):07d}.xlsx"
    if save_workbook(wb, checkpoint_file):
        print(f"Checkpoint saved: {checkpoint_file} ({len(state.recipes)} recipes)")


def read_sheet(wb, name: str) -> list[dict[str, Any]]:
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    headers = [clean_text(x) for x in next(rows, [])]
    result = []
    for values in rows:
        row = {headers[i]: values[i] if i < len(values) and values[i] is not None else "" for i in range(len(headers))}
        if any(clean_text(v) for v in row.values()):
            result.append(row)
    return result


def load_temp() -> ScrapeState:
    state = ScrapeState()
    checkpoints = sorted(TEMP_DIR.glob(CHECKPOINT_PATTERN))
    if not checkpoints:
        return state
    try:
        wb = load_workbook(checkpoints[-1], read_only=True, data_only=True)
        state.recipes = read_sheet(wb, "Recipes")
        state.ingredients = read_sheet(wb, "Ingredients")
        for row in read_sheet(wb, "State"):
            kind, url = clean_text(row.get("Type")).lower(), clean_text(row.get("URL"))
            if not url:
                continue
            if kind == "category_queued":
                state.category_queue.append(url)
            elif kind == "category_seen":
                state.seen_categories.add(url)
            elif kind == "processed":
                state.processed_urls.add(url)
            elif kind == "skipped":
                state.skipped_urls.add(url)
            elif kind == "failed":
                state.failed_urls.add(url)
        wb.close()
        state.category_queue = dedupe_preserve(u for u in state.category_queue if u not in state.seen_categories)
        print(
            f"Checkpoint loaded: {len(state.recipes)} recipes, "
            f"{len(state.seen_categories)} categories done, {len(state.category_queue)} categories queued"
        )
        return state
    except Exception as exc:
        print(f"Could not load checkpoint. Starting fresh. Reason: {exc}")
        return ScrapeState()


# ============================================================
# Browser helpers
# ============================================================
def create_context(browser) -> BrowserContext:
    context = browser.new_context(
        locale="da-DK",
        timezone_id="Europe/Copenhagen",
        viewport={"width": 1440, "height": 1100},
        user_agent=(
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/151.0.0.0 Safari/537.36"
        ),
    )
    context.set_default_timeout(PAGE_TIMEOUT_MS)
    context.set_default_navigation_timeout(NAVIGATION_TIMEOUT_MS)
    return context


def handle_cookies(page: Page) -> None:
    # Decline non-essential cookies when the banner offers it.
    candidates = ["Kun nødvendige", "Afvis alle", "Afvis", "Tillad kun nødvendige", "Nej tak", "Decline"]
    for label in candidates:
        try:
            locator = page.get_by_role("button", name=re.compile(re.escape(label), re.IGNORECASE))
            if locator.count() > 0 and locator.first.is_visible():
                locator.first.click(timeout=2_000)
                page.wait_for_timeout(500)
                return
        except Exception:
            pass


def goto_with_retry(page: Page, url: str) -> bool:
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=NAVIGATION_TIMEOUT_MS)
            page.wait_for_timeout(POLITE_DELAY_MS)
            handle_cookies(page)
            return True
        except Exception as exc:
            print(f"Navigation attempt {attempt}/{MAX_RETRIES} failed: {url} - {exc}")
            if attempt < MAX_RETRIES:
                page.wait_for_timeout(1_500 * attempt)
    return False


def hrefs(page: Page, selector: str) -> list[str]:
    try:
        links = page.eval_on_selector_all(selector, "els => els.map(a => a.href).filter(Boolean)")
    except Exception:
        return []
    return dedupe_preserve(canonicalize_url(url) for url in links if "valdemarsro.dk" in url)


def recipe_links(page: Page) -> list[str]:
    # Recipe cards are framed tiles with the "mobile-aspect-1" modifier.
    return hrefs(page, "a.category-item-framed.mobile-aspect-1")


def category_links(page: Page) -> list[str]:
    links = hrefs(
        page,
        "a.category-item-framed:not(.mobile-aspect-1), a.category-button, a.swiper-slide.item:not(.category-item-framed)",
    )
    return [url for url in links if is_category_url(url)]


def next_page_url(page: Page) -> str:
    try:
        href = page.eval_on_selector('link[rel="next"]', "el => el.href")
        return canonicalize_url(href) if href else ""
    except Exception:
        return ""


def download_image(context: BrowserContext, url: str, key: str) -> str:
    if not url:
        return ""
    try:
        response = context.request.get(url, timeout=REQUEST_TIMEOUT_MS)
        if not response.ok:
            print(f"Image download failed ({response.status}): {url}")
            return ""
        extension = Path(urlparse(url).path).suffix.lower()
        if extension not in {".jpg", ".jpeg", ".png", ".webp"}:
            extension = ".jpg"
        filename = f"{safe_filename(key)}{extension}"
        path = IMAGES_DIR / filename
        if not path.exists():
            path.write_bytes(response.body())
        return filename
    except Exception as exc:
        print(f"Image download failed: {url} - {exc}")
        return ""


# ============================================================
# Recipe extraction
# ============================================================
EXTRACT_JS = """
() => {
  const text = (el) => (el ? el.textContent.replace(/\\s+/g, ' ').trim() : '');
  const stats = {};
  document.querySelectorAll('.recipe-stats .recipe-stat').forEach((s) => {
    stats[text(s.querySelector('.recipe-stat-label'))] = text(s.querySelector('strong'));
  });
  const lines = [];
  let section = '';
  document.querySelectorAll('ul.ingredientlist li').forEach((li) => {
    if (li.classList.contains('ingredient-header')) { section = text(li); return; }
    const raw = text(li);
    if (raw) lines.push({ section, raw });
  });
  const steps = [...document.querySelectorAll('[itemprop="recipeInstructions"] p, [itemprop="recipeInstructions"] li')]
    .map(text).filter(Boolean);
  const og = document.querySelector('meta[property="og:image"]');
  const recipeImg = document.querySelector('.recipe-image img, img.recipe-image');
  return {
    name: text(document.querySelector('h2[itemprop="name"]')) || text(document.querySelector('.recipe-print-header-title')) || text(document.querySelector('h1')),
    yield: text(document.querySelector('[itemprop="recipeYield"]')),
    stats,
    lines,
    steps,
    image: (og && og.content) || (recipeImg && recipeImg.src) || '',
  };
}
"""


def parse_servings(value: str, fallback: str) -> tuple[Any, str]:
    match = re.match(r"\s*(\d+(?:[.,]\d+)?)\s*(.*)", value or "")
    if match:
        return parse_number(match.group(1)), clean_text(match.group(2))
    match = re.match(r"\s*(\d+(?:[.,]\d+)?)", fallback or "")
    return (parse_number(match.group(1)), "") if match else ("", "")


def extract_recipe(page: Page, context: BrowserContext, url: str, category: str) -> tuple[dict[str, Any], list[dict[str, Any]]] | None:
    data = page.evaluate(EXTRACT_JS)
    if not data["lines"]:
        return None

    name = clean_text(data["name"])
    key = name or slug_of(url)
    servings, servings_unit = parse_servings(data["stats"].get("Antal", ""), data["yield"])
    image_url = clean_text(data["image"])
    image_file = download_image(context, image_url, slug_of(url))

    ingredients = []
    for index, line in enumerate(data["lines"], start=1):
        parsed = parse_ingredient_line(line["raw"])
        ingredients.append(
            {
                "HelloCal_Title": key,
                "Line No": index,
                "Section": clean_text(line["section"]),
                "Raw Text": clean_text(line["raw"]),
                **parsed,
                "Source URL": url,
            }
        )

    recipe = {
        "HelloCal_Title": key,
        "Recipe Name": name,
        "Category": category,
        "Servings": servings,
        "Servings Unit": servings_unit,
        "Total Time": data["stats"].get("Tid i alt", ""),
        "Work Time": data["stats"].get("Arbejdstid", ""),
        "Freezable": data["stats"].get("Kan fryses", ""),
        "Ingredient Count": len(ingredients),
        "Ingredients": "\n".join(
            (f"[{row['Section']}] " if row["Section"] else "") + row["Raw Text"] for row in ingredients
        ),
        "Instructions": "\n".join(f"{i}. {step}" for i, step in enumerate(data["steps"], start=1)),
        "Image File": image_file,
        "Image URL": image_url,
        "Source URL": url,
        "Parse Status": "ok" if data["steps"] else "no instructions",
    }
    return recipe, ingredients


# ============================================================
# Crawl: one category at a time, each recipe scraped as soon as it is found
# ============================================================
def scrape_recipe(page: Page, context: BrowserContext, state: ScrapeState, url: str, category: str) -> None:
    if url in state.processed_urls or url in state.skipped_urls:
        return
    if not goto_with_retry(page, url):
        state.failed_urls.add(url)
        return
    try:
        result = extract_recipe(page, context, url, category)
    except Exception as exc:
        print(f"  Parse failed: {url} - {exc}")
        state.failed_urls.add(url)
        return
    if result is None:
        state.skipped_urls.add(url)
        print(f"  Skipped (no ingredient list): {url}")
        return

    recipe, ingredients = result
    state.recipes.append(recipe)
    state.ingredients.extend(ingredients)
    state.processed_urls.add(url)
    state.failed_urls.discard(url)
    print(f"  [{len(state.recipes)}] {recipe['Recipe Name']} ({len(ingredients)} ingredients)")
    if len(state.recipes) % CHECKPOINT_EVERY == 0:
        save_temp(state)


def crawl_category(page: Page, context: BrowserContext, state: ScrapeState, category_url: str) -> None:
    category = slug_of(category_url)
    listing_url = category_url
    seen_on_category: set[str] = set()
    for page_no in range(1, MAX_PAGES_PER_CATEGORY + 1):
        if not goto_with_retry(page, listing_url):
            break
        if page_no == 1:
            for sub in category_links(page):
                if sub not in state.seen_categories and sub not in state.category_queue:
                    state.category_queue.append(sub)

        found = [url for url in recipe_links(page) if url not in seen_on_category]
        next_url = next_page_url(page)
        print(f"Category '{category}' page {page_no}: {len(found)} new recipe links")
        if not found:
            break
        seen_on_category.update(found)
        for url in found:
            scrape_recipe(page, context, state, url, category)
        if not next_url or next_url == listing_url:
            break
        listing_url = next_url


def crawl_and_scrape(page: Page, context: BrowserContext, state: ScrapeState) -> None:
    if not state.category_queue and not state.seen_categories:
        if goto_with_retry(page, RECIPES_ROOT_URL):
            state.category_queue = category_links(page)
            print(f"Discovered {len(state.category_queue)} categories on {RECIPES_ROOT_URL}")

    try:
        while state.category_queue:
            category_url = state.category_queue.pop(0)
            if category_url in state.seen_categories:
                continue
            print(f"\n=== Category: {category_url} ({len(state.category_queue)} left in queue)")
            crawl_category(page, context, state, category_url)
            state.seen_categories.add(category_url)
            save_temp(state)

        for url in sorted(state.failed_urls):
            print(f"Retrying failed recipe: {url}")
            scrape_recipe(page, context, state, url, "")
    except KeyboardInterrupt:
        save_temp(state)
        save_final_files(state)
        raise
    except Exception:
        save_temp(state)
        raise


def main() -> None:
    # Show progress lines immediately in the terminal instead of buffering them.
    sys.stdout.reconfigure(line_buffering=True)
    ensure_directories()
    state = load_temp()

    print("Valdemarsro scraper started")
    print(f"Root: {ROOT_DIR}")
    print(f"Images: {IMAGES_DIR}")
    print(f"Temp checkpoint folder: {TEMP_DIR} ({CHECKPOINT_PATTERN})")
    print("Recipe limit: none")
    print("Discovery mode: one category at a time, each recipe scraped as it is found")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=HEADLESS)
        context = create_context(browser)
        page = context.new_page()
        try:
            crawl_and_scrape(page, context, state)
            save_temp(state)
            save_final_files(state)
            print(f"Completed. Recipes saved: {len(state.recipes)}")
            print(f"Recipe file: {RECIPE_FILE}")
            print(f"Ingredient file: {INGREDIENT_FILE}")
            print(f"Images saved directly in: {IMAGES_DIR}")
            print("Next step: run valdemarsro_match.py to calculate calories from the ingredients.")
        finally:
            context.close()
            browser.close()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("Stopped by user.")
        sys.exit(130)
