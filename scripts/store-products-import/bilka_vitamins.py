"""
HELLO CAL - BILKA VITAMINS AND MINERALS

Add-on to bilka.py: visits every product in bilka_product_information.xlsx,
clicks the panel "Info om vitaminer og mineraler" (bilka.py never opened it)
and writes the declared vitamins and minerals per 100 g/ml to
bilka_vitamins.xlsx next to it. The product sheets are not touched.
build_data.py reads the file and the store import stores the values as the
producer's own (docs/DECISIONS.md 2026-10-02).

Run from VS Code:  py bilka_vitamins.py            (all products, resumes)
                   py bilka_vitamins.py --limit 20 (a short test)
"""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

# ============================================================
# Fixed Hello Cal scraper structure
# ============================================================
RETAILER = "bilka"
NAS_BILKA_DIR = Path(
    r"\\192.168.1.90\Hello Cal\Arkiv - historiske kilde- og importfiler\Oprydning 2026-09-29\Productdatabase\Bilka"
)
# The script works from the Bilka folder itself (next to bilka.py) or from the
# repository; BILKA_DIR overrides both.
HERE = Path(__file__).resolve().parent
ROOT_DIR = Path(os.environ["BILKA_DIR"]) if os.environ.get("BILKA_DIR") else (
    HERE if (HERE / f"{RETAILER}_product_information.xlsx").exists() else NAS_BILKA_DIR
)
TEMP_DIR = ROOT_DIR / "Temp"
CHECKPOINT_PATTERN = f"{RETAILER}_vitamins_checkpoint_*.xlsx"
PRODUCT_INFO_FILE = ROOT_DIR / f"{RETAILER}_product_information.xlsx"
VITAMINS_FILE = ROOT_DIR / f"{RETAILER}_vitamins.xlsx"

CHECKPOINT_EVERY = 25
PAGE_TIMEOUT_MS = 30_000
BETWEEN_PRODUCTS_MS = 250
CONTEXT_RESTART_EVERY = 100
HEADLESS = True

PANEL_HEADING = "Info om vitaminer og mineraler"

# Column (per 100 g/ml, in the unit the app uses - src/lib/nutrients.ts) and
# the names Bilka uses for it. build_data.py reads exactly these columns.
NUTRIENTS = [
    ("Vitamin A µg per 100 g", "µg", ["vitamin a"]),
    ("Vitamin D µg per 100 g", "µg", ["vitamin d", "vitamin d2", "vitamin d3", "d-vitamin"]),
    ("Vitamin E mg per 100 g", "mg", ["vitamin e"]),
    ("Vitamin K µg per 100 g", "µg", ["vitamin k", "vitamin k1", "vitamin k2"]),
    ("Vitamin C mg per 100 g", "mg", ["vitamin c"]),
    ("Thiamin B1 mg per 100 g", "mg", ["thiamin", "tiamin", "vitamin b1"]),
    ("Riboflavin B2 mg per 100 g", "mg", ["riboflavin", "vitamin b2"]),
    ("Niacin B3 mg per 100 g", "mg", ["niacin", "vitamin b3"]),
    ("Pantothenic Acid B5 mg per 100 g", "mg", ["pantothensyre", "pantotensyre", "pantothenic acid", "vitamin b5"]),
    ("Vitamin B6 mg per 100 g", "mg", ["vitamin b6"]),
    ("Biotin B7 µg per 100 g", "µg", ["biotin", "vitamin b7"]),
    ("Folate B9 µg per 100 g", "µg", ["folsyre", "folat", "folacin", "folic acid", "folate", "vitamin b9"]),
    ("Vitamin B12 µg per 100 g", "µg", ["vitamin b12"]),
    ("Potassium mg per 100 g", "mg", ["kalium", "potassium"]),
    ("Calcium mg per 100 g", "mg", ["calcium", "kalcium"]),
    ("Phosphorus mg per 100 g", "mg", ["fosfor", "phosphor", "phosphorus"]),
    ("Magnesium mg per 100 g", "mg", ["magnesium"]),
    ("Iron mg per 100 g", "mg", ["jern", "iron"]),
    ("Zinc mg per 100 g", "mg", ["zink", "zinc"]),
    ("Copper mg per 100 g", "mg", ["kobber", "copper"]),
    ("Manganese mg per 100 g", "mg", ["mangan", "manganese"]),
    ("Selenium µg per 100 g", "µg", ["selen", "selenium"]),
    ("Iodine µg per 100 g", "µg", ["jod", "iodine"]),
    ("Sodium mg per 100 g", "mg", ["natrium", "sodium"]),
]
NAME_TO_COLUMN = {name: (column, unit) for column, unit, names in NUTRIENTS for name in names}
TO_MICROGRAMS = {"g": 1_000_000.0, "mg": 1_000.0, "µg": 1.0}
UNIT_ALIASES = {"g": "g", "mg": "mg", "µg": "µg", "μg": "µg", "mcg": "µg", "ug": "µg"}

ROW_RE = re.compile(
    r"(?P<name>vitamin\s+[a-k]\d{0,2}|d-vitamin|thiamin|tiamin|riboflavin|niacin|pantot?hensyre|pantothenic acid|biotin"
    r"|folsyre|folat|folacin|folic acid|folate|kalium|potassium|kalcium|calcium|fosfor|phosphor(?:us)?|magnesium"
    r"|jern|iron|zink|zinc|kobber|copper|mangan(?:ese)?|selen(?:ium)?|jod|iodine|natrium|sodium)"
    r"(?:\s*\([^)]*\))?\s*:?\s*<?\s*(?P<value>\d+(?:[.,]\d+)?)\s*(?P<unit>mg|µg|μg|mcg|ug|g)\b",
    re.IGNORECASE,
)
BASIS_RE = re.compile(r"pr\.?\s*100\s*(g|ml)\b", re.IGNORECASE)
PER_SERVING_RE = re.compile(r"pr\.?\s*(portion|dosis|tablet|kapsel|stk|glas|serving)", re.IGNORECASE)

BASE_HEADERS = ["EAN", "HelloCal_Title", "Source URL", "Basis", "Status"]
HEADERS = BASE_HEADERS + [column for column, _, _ in NUTRIENTS] + ["Other", "Raw"]

TEMP_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# General helpers
# ============================================================
def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def number(text: str) -> float:
    return float(text.replace(",", "."))


def convert(value: float, unit: str, target: str) -> float:
    return round(value * TO_MICROGRAMS[unit] / TO_MICROGRAMS[target], 6)


# ============================================================
# Panel parsing
# ============================================================
def parse_panel(text: str) -> tuple[str, str, dict[str, float], list[str]]:
    """Returns (basis, status, values per column, other rows as text)."""
    flat = clean(text)
    if not flat:
        return "", "EMPTY_PANEL", {}, []
    basis_match = BASIS_RE.search(flat)
    basis = f"100 {basis_match.group(1).lower()}" if basis_match else ""
    if not basis and PER_SERVING_RE.search(flat):
        return "", "NOT_PER_100", {}, []

    values: dict[str, float] = {}
    other: list[str] = []
    for match in ROW_RE.finditer(flat):
        name = re.sub(r"\s+", " ", match.group("name").lower())
        unit = UNIT_ALIASES.get(match.group("unit").lower().replace("μ", "µ"), "")
        value = number(match.group("value"))
        mapped = NAME_TO_COLUMN.get(name)
        if not mapped or not unit:
            other.append(f"{match.group('name')}: {match.group('value')} {match.group('unit')}")
            continue
        column, target = mapped
        # Bilka sometimes writes mcg for a mineral in mg ("Calcium 124 mcg" in
        # milk): 10 or more micrograms of an mg-nutrient is that typo.
        if unit == "µg" and target == "mg" and value >= 10:
            other.append(f"{match.group('name')}: {match.group('value')} {match.group('unit')} read as mg")
            unit = "mg"
        values.setdefault(column, convert(value, unit, target))
    if not values and not other:
        return basis, "NO_VALUES", {}, []
    return basis, "OK" if basis else "OK_BASIS_UNKNOWN", values, other


# ============================================================
# Product page
# ============================================================
def read_vitamin_panel(page, url: str) -> tuple[str, str]:
    """Opens the product page, clicks the vitamin panel open like a visitor
    and returns (status, panel text)."""
    page.goto(url, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
    page.wait_for_timeout(500)

    button = page.get_by_role("button", name=PANEL_HEADING, exact=False)
    if not button.count():
        return "NO_PANEL", ""
    button.first.click(timeout=5000)
    page.wait_for_timeout(300)

    panel = button.first.locator("xpath=ancestor::div[contains(@class, 'expansion-panel')][1]")
    text = panel.inner_text() if panel.count() else ""
    if not re.search(r"\d", text or "") and panel.count():
        text = panel.text_content() or ""
    return "FOUND", text.replace(PANEL_HEADING, "", 1)


def build_row(product: dict[str, Any], status: str, text: str) -> dict[str, Any]:
    row = {
        "EAN": product["EAN"],
        "HelloCal_Title": product["HelloCal_Title"],
        "Source URL": product["Source URL"],
        "Basis": "",
        "Status": status,
        "Other": "",
        "Raw": clean(text)[:2000],
    }
    if status == "FOUND":
        basis, parsed_status, values, other = parse_panel(text)
        row.update(values)
        row["Basis"] = basis
        row["Status"] = parsed_status
        row["Other"] = "; ".join(other)
    return row


# ============================================================
# Workbook helpers
# ============================================================
def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def save_workbook_atomically(wb: Workbook, path: Path) -> None:
    temp_path = path.with_name(f"{path.stem}.tmp{path.suffix}")
    wb.save(temp_path)
    os.replace(temp_path, path)


def rows_workbook(rows: list[dict[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vitamins and minerals"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in HEADERS])
    style_header(ws)
    return wb


def save_files(rows: list[dict[str, Any]]) -> None:
    checkpoint = TEMP_DIR / f"{RETAILER}_vitamins_checkpoint_{len(rows):07d}.xlsx"
    try:
        save_workbook_atomically(rows_workbook(rows), checkpoint)
        save_workbook_atomically(rows_workbook(rows), VITAMINS_FILE)
        print(f"Checkpoint saved: {checkpoint.name} ({len(rows)} products)", flush=True)
    except PermissionError:
        print(f"Could not save {VITAMINS_FILE.name} - file is open elsewhere. Skipped.", flush=True)


def load_checkpoint() -> list[dict[str, Any]]:
    checkpoints = sorted(TEMP_DIR.glob(CHECKPOINT_PATTERN))
    if not checkpoints:
        return []
    try:
        wb = load_workbook(checkpoints[-1], read_only=True, data_only=True)
        sheet = wb.worksheets[0].iter_rows(values_only=True)
        headers = [clean(h) for h in next(sheet)]
        rows = [dict(zip(headers, values)) for values in sheet if any(v not in (None, "") for v in values)]
        wb.close()
        print(f"Checkpoint loaded: {checkpoints[-1].name} ({len(rows)} products done)", flush=True)
        return rows
    except Exception as exc:
        print(f"Could not load checkpoint. Starting fresh. Reason: {exc}", flush=True)
        return []


def load_products() -> list[dict[str, Any]]:
    wb = load_workbook(PRODUCT_INFO_FILE, read_only=True, data_only=True)
    sheet = wb.worksheets[0].iter_rows(values_only=True)
    headers = [clean(h) for h in next(sheet)]
    products = []
    for values in sheet:
        row = dict(zip(headers, values))
        url = clean(row.get("Source URL"))
        if url:
            products.append({"EAN": clean(row.get("EAN")), "HelloCal_Title": clean(row.get("HelloCal_Title")), "Source URL": url})
    wb.close()
    return products


# ============================================================
# Main
# ============================================================
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=0, help="stop after this many new products (a test run)")
    args = parser.parse_args()

    print()
    print("HELLO CAL - BILKA VITAMINS AND MINERALS", flush=True)
    print(f"Products: {PRODUCT_INFO_FILE}", flush=True)
    print(f"Output: {VITAMINS_FILE}", flush=True)
    print(f"Temp: {TEMP_DIR}", flush=True)

    products = load_products()
    rows = load_checkpoint()
    done = {clean(row.get("Source URL")) for row in rows}
    todo = [p for p in products if p["Source URL"] not in done]
    if args.limit:
        todo = todo[: args.limit]
    print(f"{len(products)} products, {len(done)} already done, {len(todo)} to visit now", flush=True)

    since_checkpoint = 0
    with sync_playwright() as p:
        browser = None
        context = None
        page = None
        session_count = 0

        def open_session():
            nonlocal browser, context, page, session_count
            for closable in (context, browser):
                try:
                    if closable is not None:
                        closable.close()
                except Exception:
                    pass
            browser = p.chromium.launch(headless=HEADLESS)
            context = browser.new_context(locale="da-DK", viewport={"width": 1440, "height": 1000})
            page = context.new_page()
            session_count = 0
            print("  [New browser session]", flush=True)

        open_session()
        try:
            for index, product in enumerate(todo, start=1):
                if session_count >= CONTEXT_RESTART_EVERY:
                    open_session()
                print(f"[{index}/{len(todo)}] {product['Source URL']}", flush=True)
                row = None
                for attempt in range(2):
                    try:
                        if attempt == 1:
                            print("  Retrying in fresh browser session...", flush=True)
                            open_session()
                        status, text = read_vitamin_panel(page, product["Source URL"])
                        row = build_row(product, status, text)
                        break
                    except PlaywrightTimeoutError:
                        continue
                    except Exception as error:
                        print(f"  ERROR: {error}", flush=True)
                        break
                if row is None:
                    row = build_row(product, "ERROR", "")
                rows.append(row)
                session_count += 1
                since_checkpoint += 1
                found = [header.split(" per ")[0] for header in HEADERS[len(BASE_HEADERS):-2] if row.get(header) not in (None, "")]
                print(f"  {row['Status']}: {', '.join(found) or '-'}", flush=True)

                if since_checkpoint >= CHECKPOINT_EVERY:
                    save_files(rows)
                    since_checkpoint = 0
                if BETWEEN_PRODUCTS_MS:
                    page.wait_for_timeout(BETWEEN_PRODUCTS_MS)
        except KeyboardInterrupt:
            print()
            print("CTRL+C - SAVING BEFORE STOP...", flush=True)
        finally:
            for closable in (context, browser):
                try:
                    if closable is not None:
                        closable.close()
                except Exception:
                    pass

    save_files(rows)
    with_values = sum(1 for row in rows if str(row.get("Status", "")).startswith("OK"))
    print()
    print(f"DONE: {len(rows)} products visited, {with_values} with vitamins or minerals", flush=True)
    print(f"Output: {VITAMINS_FILE}", flush=True)


if __name__ == "__main__":
    main()
