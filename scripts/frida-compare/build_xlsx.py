"""Frida-sammenligning: Rema/Bilka-varer grupperet under deres Frida-reference.

Bygger et Excel-ark, hvor hver gruppe starter med Frida-varen og under den de
butiksvarer, der er vurderet som samme fødevare. Sidste kolonne er den største
afvigelse i energifordeling (protein/kulhydrat/fedt, procentpoint). Se README.md.

Kør:  py build_xlsx.py [ud.xlsx]
"""
import json, glob, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.join(HERE, '..', '..') + os.sep
DATA = os.path.join(HERE, 'data')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'Frida-sammenligning.xlsx')
fr = {f['external_id']: f for f in json.load(open(R + 'scripts/frida-import/sheet/frida_sheet.json', encoding='utf-8'))}
rema = {r['ean']: r for r in json.load(open(R + 'scripts/rema1000-import/data/rema1000_products.json', encoding='utf-8'))}

def num(x):
    if x is None or x == '': return None
    try: return float(str(x).replace(',', '.'))
    except ValueError: return None

# Kolonner: (overskrift, frida-nøgle, butiks-nøgle)
BASE = [('kcal', 'kcal', 'kcal'), ('Protein g', 'protein', 'protein'),
        ('Kulhydrat g', 'carbs', 'carbs'), ('Fedt g', 'fat', 'fat')]
BASE_KEYS = None
EXTRA = [('Mættet fedt g', 'saturatedFat'), ('Sukker g', 'sugar'), ('Kostfibre g', 'fiber'), ('Salt g', 'salt')]
# Vitaminer/mineraler (kun vist, hvis mindst én butiksvare har tallet)
MICRO = [('Vitamin A µg', 'vitaminA'), ('Vitamin D µg', 'vitaminD'), ('Vitamin E mg', 'vitaminE'),
         ('Vitamin K µg', 'vitaminK'), ('Vitamin C mg', 'vitaminC'), ('Thiamin (B1) mg', 'vitaminB1'),
         ('Riboflavin (B2) mg', 'vitaminB2'), ('Niacin (B3) mg', 'vitaminB3'), ('Pantotensyre (B5) mg', 'vitaminB5'),
         ('Vitamin B6 mg', 'vitaminB6'), ('Biotin (B7) µg', 'vitaminB7'), ('Folat (B9) µg', 'vitaminB9'),
         ('Vitamin B12 µg', 'vitaminB12'), ('Calcium mg', 'calcium'), ('Jern mg', 'iron'), ('Magnesium mg', 'magnesium'),
         ('Kalium mg', 'potassium'), ('Fosfor mg', 'phosphorus'), ('Zink mg', 'zinc'), ('Jod µg', 'iodine'),
         ('Selen µg', 'selenium'), ('Kobber mg', 'copper'), ('Mangan mg', 'manganese')]

BASE_KEYS = [(h, k) for h, k, _ in BASE] + EXTRA

def frida_vals(f):
    v = {k: num(f.get(k)) for _, k, _ in BASE}
    v.update({k: num((f.get('micros') or {}).get(k)) for _, k in EXTRA + MICRO})
    return v

def rema_vals(r):
    n = r['nutrition']
    return {'kcal': num(n.get('energyKcal')), 'protein': num(n.get('protein')), 'carbs': num(n.get('carbohydrate')),
            'fat': num(n.get('fat')), 'saturatedFat': num(n.get('saturatedFat')), 'sugar': num(n.get('sugars')),
            'fiber': num(n.get('fibre')), 'salt': num(n.get('salt'))}

def macro_kcal(v):
    if None in (v.get('protein'), v.get('carbs'), v.get('fat')): return 0
    return 4 * v['protein'] + 4 * v['carbs'] + 9 * v['fat']

# Butiksrækker: [(kilde, navn, id, værdier, note, frida_id)]
stores = []
for path in [os.path.join(DATA, 'rema_matches.json')]:  # matchet 2026-10-10, se README
    for m in json.load(open(path, encoding='utf-8')):
        r = rema.get(m['ean']); f = fr.get(str(m['frida_id']))
        if not r or not f: print('ukendt id', m, file=sys.stderr); continue
        name = r['helloCalTitle'] + (f" ({r['quantity']})" if r.get('quantity') else '')
        stores.append(('Rema 1000', name, r['ean'], rema_vals(r), m.get('note', ''), f['external_id']))
# Bilka: bilka_extract.py skriver bilka_values.json; matchningen skriver bilka_matches.json (se README).
bilka_values_path = os.path.join(DATA, 'bilka_values.json')
bilka_matches_path = os.path.join(DATA, 'bilka_matches.json')
if os.path.isfile(bilka_values_path) and os.path.isfile(bilka_matches_path):
    bilka = json.load(open(bilka_values_path, encoding='utf-8'))
    for m in json.load(open(bilka_matches_path, encoding='utf-8')):
        b = bilka.get(m['ean']); f = fr.get(str(m['frida_id']))
        if not b or not f: print('ukendt id', m, file=sys.stderr); continue
        stores.append(('Bilka', b['name'], m['ean'], b['values'], m.get('note', ''), f['external_id']))

ALCOHOL_CATS = {'Cider', 'Likør', 'Spiritus', 'Vine', 'Hedvine', 'Øl og andre maltdrikke'}
skipped = []; alcohol = []
groups = {}
for s in stores:
    f = fr[s[5]]
    if f['category'] in ALCOHOL_CATS or (f.get('tags') or {}).get('isAlcohol'):
        alcohol.append(s); continue
    if macro_kcal(s[3]) < 5 or macro_kcal(frida_vals(f)) < 5:
        skipped.append(s); continue
    groups.setdefault(s[5], []).append(s)

def dev(a, b):
    ta, tb = macro_kcal(a), macro_kcal(b)
    return max(abs(4 * a['protein'] / ta - 4 * b['protein'] / tb), abs(4 * a['carbs'] / ta - 4 * b['carbs'] / tb),
               abs(9 * a['fat'] / ta - 9 * b['fat'] / tb))

for fid, rows in groups.items():
    fv = frida_vals(fr[fid])
    rows.sort(key=lambda s: -dev(s[3], fv))
order = sorted(groups, key=lambda fid: (-max(dev(s[3], frida_vals(fr[fid])) for s in groups[fid]), fr[fid]['name']))

micro_cols = [(h, k) for h, k in MICRO if any(s[3].get(k) is not None for s in stores)]
extra_cols = [(h, k) for h, k in EXTRA if any(s[3].get(k) is not None for s in stores)]

wb = Workbook(); ws = wb.active; ws.title = 'Sammenligning'
FONT = 'Arial'
heads = (['Kilde', 'Vare', 'FoodID / EAN'] + [h for h, _, _ in BASE] + ['Protein E%', 'Kulhydrat E%', 'Fedt E%']
         + [h for h, _ in extra_cols] + [h for h, _ in micro_cols] + ['Matchnote', 'Afviger mest på', 'Afvigelse %'])
col = {h: i + 1 for i, h in enumerate(heads)}
ws.append(heads)
for c in ws[1]:
    c.font = Font(name=FONT, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F5233')
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
ws.row_dimensions[1].height = 32
ws.freeze_panes = 'D2'

FRIDA_FILL = PatternFill('solid', fgColor='E8F0E3')
cP, cK, cF = L(col['Protein g']), L(col['Kulhydrat g']), L(col['Fedt g'])
eP, eK, eF = L(col['Protein E%']), L(col['Kulhydrat E%']), L(col['Fedt E%'])
row = 2
for fid in order:
    f = fr[fid]; fv = frida_vals(f); fr_row = row
    entries = [('Frida', f['name'], f['external_id'], fv, f"Reference · {f['category']}")] + [s[:5] for s in groups[fid]]
    for kilde, name, ident, v, note in entries:
        ws.cell(row, col['Kilde'], kilde); ws.cell(row, col['Vare'], name); ws.cell(row, col['FoodID / EAN'], str(ident))
        for h, k, _ in BASE: ws.cell(row, col[h], v.get(k))
        for h, k in extra_cols + micro_cols: ws.cell(row, col[h], v.get(k))
        tot = f'(4*{cP}{row}+4*{cK}{row}+9*{cF}{row})'
        ws.cell(row, col['Protein E%'], f'=4*{cP}{row}/{tot}')
        ws.cell(row, col['Kulhydrat E%'], f'=4*{cK}{row}/{tot}')
        ws.cell(row, col['Fedt E%'], f'=9*{cF}{row}/{tot}')
        ws.cell(row, col['Matchnote'], note)
        if kilde != 'Frida':
            d = [f'ABS({c}{row}-{c}${fr_row})' for c in (eP, eK, eF)]
            dv = L(col['Afvigelse %'])
            ws.cell(row, col['Afvigelse %'], f'=MAX({",".join(d)})')
            ws.cell(row, col['Afviger mest på'],
                    f'=IF({dv}{row}={d[0]},"Protein",IF({dv}{row}={d[1]},"Kulhydrat","Fedt"))')
        for c in ws[row]:
            c.font = Font(name=FONT, bold=(kilde == 'Frida'))
            if kilde == 'Frida': c.fill = FRIDA_FILL
        if kilde == 'Rema 1000':  # 0,5 g er REMA's pladsholder for "under 0,5 g"
            for h in ('Protein g', 'Kulhydrat g', 'Fedt g', 'Mættet fedt g', 'Sukker g', 'Kostfibre g'):
                if h in col and v.get(dict(BASE_KEYS)[h]) == 0.5:
                    ws.cell(row, col[h]).font = Font(name=FONT, italic=True, color='8C8C8C')
        row += 1
    row += 1  # tom række mellem grupper

last = row
for h in heads:
    letter = L(col[h])
    fmt = '0.0%' if h.endswith('E%') or h == 'Afvigelse %' else ('0' if h == 'kcal' else '0.0##')
    if h in ('Kilde', 'Vare', 'FoodID / EAN', 'Matchnote', 'Afviger mest på'): fmt = None
    if fmt:
        for r in range(2, last): ws[f'{letter}{r}'].number_format = fmt
    ws.column_dimensions[letter].width = {'Vare': 52, 'Matchnote': 44, 'FoodID / EAN': 15, 'Kilde': 10,
                                          'Afviger mest på': 13}.get(h, 11)
dv = L(col['Afvigelse %'])
rng = f'{dv}2:{dv}{last}'
ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['0.15'], fill=PatternFill('solid', fgColor='F4B6B6')))
ws.conditional_formatting.add(rng, CellIsRule(operator='between', formula=['0.075', '0.15'], fill=PatternFill('solid', fgColor='FBE3A6')))
for c in ws[f'{dv}2:{dv}{last}']:
    c[0].font = Font(name=FONT, bold=True)

m = wb.create_sheet('Metode')
lines = [
    ('Frida ↔ butik: energifordeling', True),
    ('', False),
    ('Hver gruppe starter med Frida-referencen (grøn, fed). Under den står de butiksvarer, der er vurderet som samme fødevare (ca. 90 % ens).', False),
    ('Matchet er lavet ud fra navn, type, tilstand og variant — uden at kigge på næringstal, så store afvigelser ikke er sorteret fra.', False),
    ('Naturlige forskelle (mærke, økologi, størrelse, lille smagsvariant) er tilladt; se Matchnote.', False),
    ('', False),
    ('Energifordeling (E%) = andel af energien fra protein, kulhydrat og fedt, beregnet ens for alle rækker: 4 kcal/g protein, 4 kcal/g kulhydrat, 9 kcal/g fedt.', False),
    ('Afvigelse % = største forskel i procentpoint mellem butiksvarens og Fridas E% for protein, kulhydrat eller fedt.', False),
    ('Farver: rød ≥ 15 procentpoint, gul 7,5–15 procentpoint. Det er en indikator for, hvad der bør undersøges nærmere — ikke en fejlmelding.', False),
    ('Grupperne er sorteret med størst afvigelse først; inden for gruppen også efter afvigelse.', False),
    ('Alkohol, fibre og polyoler indgår ikke i E% (står ikke på alle deklarationer). For øl/vin er E% derfor kun fra makroer.', False),
    ('Alkoholholdige drikke (øl, vin, hedvin, cider, spiritus) er udeladt: det meste af deres energi kommer fra alkohol, så fordelingen mellem protein, kulhydrat og fedt er støj.', False),
    ('Grå kursiv = præcis 0,5 g hos butikken. REMA 1000-datasættet bruger 0,5 som pladsholder for "under 0,5 g" (fedt hos 601 varer, kostfibre hos 1.361) — det kan give kunstige udslag på magre varer.', False),
    ('Kcal og gram står som deklareret; en stor kcal-forskel (fx 0–1 kcal) peger ofte på en fejl i butiksdata eller på tal pr. tilberedt portion (kaffe).', False),
    ('Varer med under 5 kcal fra makroer (vand, kaffe, te, light-sodavand) er udeladt, da fordelingen ikke giver mening.', False),
    ('Kun kolonner med næringsdata fra mindst én butiksvare er vist. REMA 1000 deklarerer ikke vitaminer/mineraler i datasættet; vitaminkolonnerne kommer fra Bilka.', False),
    ('', False),
    ('Kilder: Frida-arket (scripts/frida-import/sheet/frida_sheet.json, DTU Frida 6.1), REMA 1000-datasættet (scripts/rema1000-import/data/rema1000_products.json) og Bilkas egne ark (bilka_product_information.xlsx + bilka_vitamins.xlsx), når de er tilføjet.', False),
    (f'Antal: {len(groups)} Frida-grupper, {sum(len(g) for g in groups.values())} butiksvarer. Udeladt pga. ~0 kcal: {len(skipped)}. Udeladt alkohol: {len(alcohol)}.', False),
]
for t, b in lines:
    m.append([t]); m.cell(m.max_row, 1).font = Font(name=FONT, bold=b, size=13 if b else 10)
m.column_dimensions['A'].width = 140
wb.calculation.fullCalcOnLoad = True  # Excel beregner formlerne ved åbning
wb.save(OUT)
print('grupper', len(groups), 'varer', sum(len(g) for g in groups.values()), 'udeladt', len(skipped))
