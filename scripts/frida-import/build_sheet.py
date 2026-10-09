"""
Bygger scripts/frida-import/sheet/frida_sheet.json ud fra det færdige Frida-ark
(Frida-ark/frida.xlsx — titler er formler, så arket skal være gemt i Excel,
så de beregnede værdier ligger i filen) og Fridas officielle datasæt
(data/Frida_FCDB_6.1_full.xlsx) til næring (kcal, protein, kulhydrat, fedt og
mikrodata pr. 100 g).

Kør lokalt:  python build_sheet.py "<sti til frida.xlsx>" [<sti til Frida_FCDB xlsx>]

agent.py (apply_sheet) læser JSON-filen og opdaterer/opretter/sletter
Frida-varerne i databasen — se docs/FRIDA.md.
"""

import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))

# Samme parameter-mapping som agent.py (holdes i sync dér).
PARAM_KCAL, PARAM_PROTEIN, PARAM_CARBS, PARAM_FAT = 356, 218, 170, 141
MACRO_PARAMS = {PARAM_KCAL, PARAM_PROTEIN, PARAM_CARBS, PARAM_FAT}
MICRO_PARAMS = {
    "saturatedFat": [248], "unsaturatedFat": [247, 251], "transFat": [261], "cholesterol": [115],
    "sugar": [245], "fiber": [168], "salt": [327], "sodium": [201], "potassium": [165],
    "calcium": [108], "magnesium": [184], "iron": [162], "zinc": [274], "copper": [166],
    "manganese": [187], "selenium": [230], "phosphorus": [214], "iodine": [163],
    "vitaminA": [12], "vitaminC": [47], "vitaminD": [126], "vitaminE": [135], "vitaminK": [442],
    "vitaminB1": [37], "vitaminB2": [39], "vitaminB3": [294], "vitaminB5": [210], "vitaminB6": [40],
    "vitaminB7": [42], "vitaminB9": [143], "vitaminB12": [38],
}
WANTED = MACRO_PARAMS | {pid for ids in MICRO_PARAMS.values() for pid in ids}

# Ark-kolonne → nøgle i products."dietaryTags" (kun udfyldte gemmes).
TAG_COLUMNS = {
    "_is_frozen": "isFrozen", "_is_raw": "isRaw", "_is_cooked": "isCooked", "_is_processed": "isProcessed",
    "_is_light": "isLight", "_is_country_of_origen": "countryOfOrigin", "_is_organic": "isOrganic",
    "_is_sweeteners": "isSweeteners", "_is_meat": "isMeat", "_is_vegan": "isVegan",
    "_is_sugar_free": "isSugarFree", "_is_fat": "isFat", "_is_alcohol_free": "isAlcoholFree",
    "_is_decaf": "isDecaf",
}


def parse_nutrition(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = {}
    for name_dk, _name_en, food_id, *_ in wb["Food"].iter_rows(min_row=2, values_only=True):
        if food_id is not None:
            names[food_id] = name_dk
    values = {}
    for row in wb["Data_Normalised"].iter_rows(min_row=2, values_only=True):
        food_id, param_id, res_val = row[0], row[3], row[7]
        if param_id in WANTED and isinstance(res_val, (int, float)):
            values.setdefault(food_id, {})[param_id] = float(res_val)
    out = {}
    for food_id, params in values.items():
        if not MACRO_PARAMS.issubset(params):
            continue
        micros = {
            key: round(sum(params[pid] for pid in ids), 4)
            for key, ids in MICRO_PARAMS.items()
            if all(pid in params for pid in ids)
        }
        out[str(food_id)] = {
            "kcal": round(params[PARAM_KCAL], 2), "protein": round(params[PARAM_PROTEIN], 2),
            "carbs": round(params[PARAM_CARBS], 2), "fat": round(params[PARAM_FAT], 2),
            "micros": micros, "originalName": names.get(food_id),
        }
    return out


def clean(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip()
    return value or None


def main():
    sheet_path = sys.argv[1]
    nutrition_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), "data", "Frida_FCDB_6.1_full.xlsx")
    nutrition = parse_nutrition(nutrition_path)

    ws = openpyxl.load_workbook(sheet_path, data_only=True).active
    header = [c.value for c in ws[1]]
    # "_is_alcohol" findes to gange i arket: første = tekst ("indeholder alkohol"), anden = procent.
    cols = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        key = h if h not in cols else h + "_2"
        cols[key] = i

    seen = {}
    rows, missing = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        food_id = clean(row[cols["FoodID"]])
        singular = clean(row[cols["Product title singular"]])
        plural = clean(row[cols["Product title plural"]]) or singular
        if not food_id or not singular:
            continue
        seen[food_id] = seen.get(food_id, 0) + 1
        external_id = food_id if seen[food_id] == 1 else f"{food_id}-{seen[food_id]}"
        nut = nutrition.get(food_id)
        if not nut:
            missing.append(food_id)
            continue
        tags = {}
        for column, key in TAG_COLUMNS.items():
            value = clean(row[cols[column]]) if column in cols else None
            if value:
                tags[key] = value
        alcohol_pct = clean(row[cols["_is_alcohol_2"]]) if "_is_alcohol_2" in cols else None
        alcohol_text = clean(row[cols["_is_alcohol"]]) if "_is_alcohol" in cols else None
        if alcohol_text:
            tags["isAlcohol"] = alcohol_text
        if alcohol_pct:
            tags["alcoholPercent"] = alcohol_pct
        keywords = [k for k in (clean(row[cols[f"Keyword {n}"]]) for n in (1, 2, 3)) if k]
        rows.append(
            {
                "external_id": external_id,
                "food_id": food_id,
                "name": singular,
                "name_plural": plural,
                "product_type": clean(row[cols["Product Type"]]),
                "variant": clean(row[cols["Variation"]]),
                "category": clean(row[cols["Category"]]),
                "keywords": keywords,
                "tags": tags,
                **{k: nut[k] for k in ("kcal", "protein", "carbs", "fat", "micros")},
            }
        )

    out_dir = os.path.join(os.path.dirname(__file__), "sheet")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "frida_sheet.json")
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(rows, fh, ensure_ascii=False, separators=(",", ":"))
    print(f"skrev {len(rows)} varer til {out_path}; uden næring i datasættet (springes over): {sorted(set(missing))}")


if __name__ == "__main__":
    main()
