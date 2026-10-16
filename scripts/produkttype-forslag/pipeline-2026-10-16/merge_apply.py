"""Samler workflow-output (batches/*.out.json) til final_changes.json i samme format som mech_changes.json
(bruges af apply_mech.ps1 -Json ...). Rører kun celler hvor værdien faktisk ændres; sletter rækker med remove=true."""
import collections
import json
import os
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rule_check as rc

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = r"C:\Users\Peter\Desktop\Hello Cal"
SHEETS = {
    "DRK": "Produkter/drk.xlsx", "Aarstiderne": "Produkter/aarstiderne.xlsx", "Wolt": "Produkter/DK/wolt.xlsx",
    "DM": "Produkter/dm.xlsx", "EDEKA": "Produkter/edeka.xlsx", "REWE": "Produkter/rewe.xlsx",
}
LANG = {"DRK": "da", "Aarstiderne": "da", "Wolt": "da", "DM": "de", "EDEKA": "de", "REWE": "de"}
GENERIC_SB = {
    "bio", "vegan", "vegansk", "zero", "light", "original", "classic", "classico", "extra", "natur", "naturell", "premium", "pur", "mild",
    "økologisk", "øko", "organic", "frisch", "frisk", "fein", "gold", "silber", "red", "blue", "green", "white", "black", "plus", "mini", "maxi", "xl", "xxl",
    "glutenfrei", "glutenfri", "laktosefrei", "laktosefri", "zuckerfrei", "sukkerfri", "alkoholfrei", "alkoholfri", "protein", "high protein",
}
FLAGCOL = {"organic": "_is_organic", "alcohol_free": "_is_alcohol_free", "sugar_free": "_is_sugar_free", "glutenfree": "_is_glutenfree",
           "lactose_free": "_is_lactose_free", "vegan": "_is_vegan"}


def load(rel):
    dst = os.path.join(tempfile.gettempdir(), "mg_" + os.path.basename(rel))
    shutil.copy2(os.path.join(ROOT, rel), dst)
    ws = openpyxl.load_workbook(dst, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    h = [str(x) if x is not None else "" for x in next(it)]
    return h, list(it)


def sv(v):
    return None if v in (None, "") else str(v).strip()


def main():
    manifest = json.load(open(os.path.join(HERE, "manifest.json"), encoding="utf-8"))
    by_sheet = collections.defaultdict(list)
    for m in manifest:
        by_sheet[m["sheet"]].append(m)
    out = []
    removed_export = {}
    for sheet, items in by_sheet.items():
        rel = SHEETS[sheet]
        h, rows = load(rel)
        idx = {}
        for i, n in enumerate(h):
            idx.setdefault(n, i)
        alc = [i for i, n in enumerate(h) if n == "_is_alcohol"]
        kw = [i for i, n in enumerate(h) if n.startswith("Keyword ")]
        outputs = {}
        inputs = {}
        problems = 0
        for m in items:
            if not os.path.exists(m["output"]):
                problems += 1
                continue
            try:
                for i in json.load(open(m["input"], encoding="utf-8")):
                    inputs[i["row"]] = i
                for o in json.load(open(m["output"], encoding="utf-8")):
                    outputs[o["row"]] = o
            except Exception as e:  # noqa: BLE001
                print("fejl i", m["output"], e)
                problems += 1
        changes, delete, delete_orig, stats = [], [], [], collections.Counter()
        for n, r in enumerate(rows, start=2):
            o = outputs.get(n)
            if o is None:
                continue
            orig = r[idx["Original Title"]]
            inp = inputs.get(n)
            if inp is None or sv(inp.get("orig_title")) != sv(orig):
                stats["sprunget_over_titel_ændret"] += 1
                continue
            if o.get("remove"):
                delete.append(n)
                delete_orig.append(str(orig))
                removed_export.setdefault(sheet, []).append(list(r))
                stats["fjernes"] += 1
                continue
            cols = {}
            errs_now, _w = rc.check_row(inp, o)
            if any(e.startswith("V12b") for e in errs_now):
                o = dict(o, country=None)
                stats["land_droppet_ikke_belagt"] += 1
            if any(e.startswith("V9b") for e in errs_now):
                o = dict(o, alcohol_percent=None)
                stats["alkoholprocent_droppet_ikke_belagt"] += 1

            def setcol(name, val, only_if_empty=False, never_clear=False):
                if name not in idx:
                    return
                ci = idx[name]
                cur = sv(r[ci])
                new = sv(val)
                if new == cur:
                    return
                if new is None and never_clear:
                    return
                if only_if_empty and cur is not None:
                    return
                cols[str(ci)] = new
                stats[name] += 1

            setcol("Brand", o.get("brand"), never_clear=True)
            cur_sb = sv(r[idx["Subbrand"]]) if "Subbrand" in idx else None
            if sv(o.get("subbrand")) is None and cur_sb is not None and cur_sb.lower() not in GENERIC_SB:
                pass  # ryd aldrig et ægte subbrand (kun generiske ord som Bio/Vegan/Zero fjernes)
            else:
                setcol("Subbrand", o.get("subbrand"))
            setcol("Co-brand", o.get("cobrand"), never_clear=True)
            setcol("Product Type", o.get("product_type"), never_clear=True)
            setcol("Variation", o.get("variation"))
            if LANG[sheet] == "da":  # de tyske ark har bevidst tomt flertal (ingen tysk bøjning)
                plural = o.get("product_type_plural") or o.get("product_type")
                setcol("Product type plural", plural, never_clear=True)
            newkw = [k for k in (o.get("keywords") or []) if k and str(k).strip().lower() not in ("allergen", "allergener")][:5]
            curkw = [str(r[i]) for i in kw if r[i] not in (None, "")]
            if newkw != curkw:
                for j, ci in enumerate(kw):
                    val = newkw[j] if j < len(newkw) else None
                    if sv(r[ci]) != sv(val):
                        cols[str(ci)] = val
                stats["keywords"] += 1
            for k, c in FLAGCOL.items():
                if o.get("flags", {}).get(k):
                    setcol(c, o["flags"][k], only_if_empty=True)
            if o.get("country"):
                setcol("_is_country_of_origen", o["country"], only_if_empty=True)
            if o.get("alcohol_percent") and len(alc) > 1 and sv(r[alc[1]]) is None:
                cols[str(alc[1])] = o["alcohol_percent"]
                stats["alkohol_procent"] += 1
            if cols:
                changes.append({"row": n, "orig": str(orig), "cols": cols})
        out.append({"sheet": sheet, "file": rel, "header": h, "rows": len(rows), "changes": changes, "delete_rows": delete, "delete_orig": delete_orig,
                    "stats": dict(stats), "problems": problems})
        print(sheet, len(rows), "rækker:", len(changes), "med ændringer,", len(delete), "fjernes,", dict(stats), "manglende batches:", problems)
    json.dump(out, open(os.path.join(HERE, "final_changes.json"), "w", encoding="utf-8"), ensure_ascii=False)
    if removed_export:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet, rws in removed_export.items():
            h = [x for x in out if x["sheet"] == sheet][0]["header"]
            ws = wb.create_sheet(sheet[:30])
            ws.append(h)
            for r in rws:
                ws.append(r)
        dest = os.path.join(ROOT, "Produkter", "Blacklistede produkter", "fjernede_raekker_ikke_mad_2026-10-16.xlsx")
        wb.save(dest)
        print("fjernede rækker gemt:", dest)


if __name__ == "__main__":
    main()
