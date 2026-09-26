"""Calculate calories for the recipes scraped by one of the recipe-site
scrapers. Run: python recipe_sites_match.py arla   (or coop, rema1000, meny,
hjerteforeningen, tv2)

Uses the same Frida + store-database matcher as Valdemarsro
(scripts/valdemarsro-import/valdemarsro_match.py): every ingredient line is
converted to grams and matched; nothing is guessed. Lines that cannot be
matched go to the "Review" sheet and can be fixed in
<site>_manual_matches.xlsx (created on first run).

When the site itself states energy per person (Hjerteforeningen), that value
is used as the final calories; otherwise the calculated value is used.
"""

from __future__ import annotations

import importlib
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "valdemarsro-import"))

import valdemarsro_match as vm  # noqa: E402

from recipe_sites_common import append_sheet, clean_text, read_sheet, save_workbook  # noqa: E402

SITES = ["arla", "coop", "rema1000", "meny", "hjerteforeningen", "tv2"]


def load_manual_matches(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        wb = Workbook()
        append_sheet(wb, "Manual Matches", vm.MANUAL_HEADERS, [], first=True)
        save_workbook(wb, path)
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = read_sheet(wb, wb.sheetnames[0])
    wb.close()
    return {vm.normalize(row.get("Ingredient Name")): row for row in rows if clean_text(row.get("Ingredient Name"))}


def site_per_serving(recipe: dict[str, Any]) -> float | None:
    basis = clean_text(recipe.get("Site Nutrition Basis")).lower()
    kcal = vm.to_float(recipe.get("Site Energy kcal"))
    if kcal and ("person" in basis or "portion" in basis or "serving" in basis):
        return kcal
    return None


def main() -> None:
    if len(sys.argv) < 2 or sys.argv[1] not in SITES:
        print(f"Usage: python recipe_sites_match.py <{'|'.join(SITES)}>")
        sys.exit(1)
    site = importlib.import_module(sys.argv[1]).SITE
    nutrition_file = site.root_dir / f"{site.retailer}_nutrition.xlsx"
    manual_file = site.root_dir / f"{site.retailer}_manual_matches.xlsx"
    if not site.recipe_file.exists() or not site.ingredient_file.exists():
        print(f"Run {sys.argv[1]}.py first - missing {site.recipe_file.name} / {site.ingredient_file.name}")
        sys.exit(1)

    wb = load_workbook(site.recipe_file, read_only=True, data_only=True)
    recipes = read_sheet(wb, "Recipes")
    wb.close()
    wb = load_workbook(site.ingredient_file, read_only=True, data_only=True)
    lines = read_sheet(wb, "Ingredients")
    wb.close()
    print(f"{len(recipes)} recipes, {len(lines)} ingredient lines")

    matcher = vm.Matcher(vm.load_frida(), vm.load_store_products())
    manual_matches = load_manual_matches(manual_file)

    match_rows: list[dict[str, Any]] = []
    review: dict[str, dict[str, Any]] = {}
    totals: dict[str, dict[str, float]] = {}

    for line in lines:
        title = clean_text(line.get("HelloCal_Title"))
        name = clean_text(line.get("Ingredient Name"))
        manual = manual_matches.get(vm.normalize(name))
        total = totals.setdefault(title, {"kcal": 0.0, "protein": 0.0, "carbs": 0.0, "fat": 0.0, "grams": 0.0,
                                          "matched": 0, "review": 0})
        row = {**line, "Grams": "", "Grams Basis": "", "Match Source": "", "Matched Name": "", "Score": "",
               "Energy kcal": "", "Status": ""}

        if manual and clean_text(manual.get("Ignore")):
            row["Status"] = "ignored (manual)"
            total["matched"] += 1
            match_rows.append(row)
            continue

        grams, basis = vm.to_grams(line, manual)
        food, score = (None, 0.0)
        volume = clean_text(line.get("Unit")) in vm.UNIT_ML
        if manual and clean_text(manual.get("Reference Name")):
            food = matcher.by_name.get(clean_text(manual["Reference Name"]).lower())
            score = 1.0 if food else 0.0
        elif grams:
            food, score = matcher.match(name, volume=volume)

        row.update({"Grams Basis": basis, "Score": round(score, 2)})
        if grams is not None:
            row["Grams"] = round(grams, 1)

        if grams == 0:
            row["Status"] = "ok (0 g)"
            total["matched"] += 1
        elif grams is not None and food:
            factor = grams / 100
            total["kcal"] += food.kcal * factor
            total["protein"] += food.protein * factor
            total["carbs"] += food.carbs * factor
            total["fat"] += food.fat * factor
            total["grams"] += grams
            total["matched"] += 1
            row.update({"Match Source": food.source, "Matched Name": food.name,
                        "Energy kcal": round(food.kcal * factor, 1), "Status": "ok"})
        else:
            reason = basis if grams is None else "no confident match"
            row["Status"] = f"review: {reason}"
            total["review"] += 1
            entry = review.setdefault(vm.normalize(name) or name, {
                "Ingredient Name": name, "Reason": reason, "Best Guess": "", "Best Guess Score": "",
                "Example Line": clean_text(line.get("Raw Text")), "Recipe Count": 0,
            })
            entry["Recipe Count"] += 1
            if grams is not None:
                guess, guess_score = matcher.match(name, volume=volume)
                if guess:
                    entry["Best Guess"], entry["Best Guess Score"] = guess.name, round(guess_score, 2)
        match_rows.append(row)

    recipe_rows = []
    for recipe in recipes:
        title = clean_text(recipe.get("HelloCal_Title"))
        total = totals.get(title, {"kcal": 0, "protein": 0, "carbs": 0, "fat": 0, "grams": 0, "matched": 0, "review": 0})
        servings = vm.to_float(recipe.get("Servings")) or 0
        complete = total["review"] == 0 and total["matched"] > 0
        site_kcal = site_per_serving(recipe)
        row = {
            "HelloCal_Title": title,
            "Meal Type": recipe.get("Meal Type"),
            "Child Friendly": recipe.get("Child Friendly"),
            "Servings": recipe.get("Servings"),
            "Servings Unit": recipe.get("Servings Unit"),
            "Matched Lines": total["matched"],
            "Review Lines": total["review"],
            "Match Status": "complete" if complete else "needs review",
            "Total Weight g": "", "Energy kcal total": "", "Energy kcal per serving": "",
            "Protein per serving": "", "Carbohydrate per serving": "", "Fat per serving": "",
            "Energy kcal per 100 g": "",
            "Site Nutrition Basis": recipe.get("Site Nutrition Basis"),
            "Site Energy kcal": recipe.get("Site Energy kcal"),
            "Final kcal per serving": "",
            "Final Source": "",
            "Source URL": recipe.get("Source URL"),
        }
        if complete:
            per = servings if servings > 0 else 1
            row.update({
                "Total Weight g": round(total["grams"]),
                "Energy kcal total": round(total["kcal"]),
                "Energy kcal per serving": round(total["kcal"] / per),
                "Protein per serving": round(total["protein"] / per, 1),
                "Carbohydrate per serving": round(total["carbs"] / per, 1),
                "Fat per serving": round(total["fat"] / per, 1),
                "Energy kcal per 100 g": round(total["kcal"] / total["grams"] * 100) if total["grams"] else "",
                "Final kcal per serving": round(total["kcal"] / per),
                "Final Source": "calculated",
            })
        if site_kcal:
            row.update({"Final kcal per serving": round(site_kcal), "Final Source": "site"})
        recipe_rows.append(row)

    review_rows = sorted(review.values(), key=lambda r: -r["Recipe Count"])
    wb = Workbook()
    append_sheet(wb, "Recipe Nutrition", list(recipe_rows[0].keys()) if recipe_rows else ["HelloCal_Title"], recipe_rows, first=True)
    append_sheet(wb, "Ingredient Matches", list(match_rows[0].keys()) if match_rows else ["HelloCal_Title"], match_rows)
    append_sheet(wb, "Review", ["Ingredient Name", "Reason", "Best Guess", "Best Guess Score", "Example Line", "Recipe Count"], review_rows)
    if save_workbook(wb, nutrition_file):
        final = sum(1 for r in recipe_rows if r["Final kcal per serving"] != "")
        print(f"Saved: {nutrition_file}")
        print(f"Recipes with calories: {final}/{len(recipe_rows)}")
        print(f"Ingredients to review: {len(review_rows)} (add them to {manual_file.name} and run again)")


if __name__ == "__main__":
    main()
