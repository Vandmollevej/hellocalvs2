"""Gendan Variation-celler som optimer_ark.py ved en fejl afkortede (indledende mit/ohne/aus/in blev fjernet, 2026-10-11).
Brug: py -3 gendan_variation.py <optimering_*.csv> [butik ...] [--dry]
Kun celler hvor 'foer' begynder med et funktionsord og cellen stadig har 'efter'-vaerdien gendannes. Backup foer skrivning.
"""
import csv, datetime, os, shutil, sys, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BACKUP = os.path.join(HERE, 'backup')
SHEETS = {'dm': 'dm.xlsx', 'edeka': 'edeka.xlsx', 'rewe': 'rewe.xlsx', 'wolt': 'wolt.xlsx'}
LEAD = ('mit ', 'ohne ', 'in ', 'im ', 'aus ', 'von ', 'vom ', 'für ', 'auf ', 'zum ', 'zur ', 'und ', 'nach ', 'bei ', 'med ', 'uden ', 'i ', 'af ', 'på ', 'fra ')


def main(args):
    dry = '--dry' in args
    a = [x for x in args if not x.startswith('--')]
    src, stores = a[0], a[1:] or list(SHEETS)
    rows = list(csv.reader(open(src, encoding='utf-8-sig'), delimiter=';'))[1:]
    todo = {}
    for st, rn, col, before, after, rule in rows:
        if rule == 'variation_oprydning' and col == 'Variation' and st in stores and before.lower().startswith(LEAD):
            todo.setdefault(st, []).append((int(rn), before, after))
    for st, items in todo.items():
        path = os.path.join(HERE, SHEETS[st])
        lock = os.path.join(HERE, '~$' + SHEETS[st])
        if os.path.exists(lock) and not dry:
            print(f'{st}: SPRUNGET OVER - aaben i Excel')
            continue
        mtime = os.path.getmtime(path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        hdr = [c.value for c in ws[1]]
        vi = hdr.index('Variation') + 1
        n = skipped = 0
        for rn, before, after in items:
            c = ws.cell(row=rn, column=vi)
            cur = '' if c.value is None else str(c.value)
            if cur == after:
                c.value = before
                n += 1
            else:
                skipped += 1
        fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP'))
        os.close(fd)
        wb.calculation.fullCalcOnLoad = True
        wb.save(tmp)
        wb.close()
        if dry:
            print(f'{st}: proeve: {n} gendannet, {skipped} uaendret -> {tmp}')
            continue
        if os.path.getmtime(path) != mtime or os.path.exists(lock):
            print(f'{st}: SPRUNGET OVER - blev aendret/aabnet imens')
            os.remove(tmp)
            continue
        stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
        shutil.copy2(path, os.path.join(BACKUP, f'{os.path.splitext(SHEETS[st])[0]}_{stamp}_foer_gendan-variation.xlsx'))
        shutil.move(tmp, path)
        print(f'{st}: {n} Variation-celler gendannet, {skipped} uaendret (anden vaerdi nu) -> {path}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1:])
