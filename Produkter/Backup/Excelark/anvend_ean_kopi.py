"""Kopier manglende stregkoder mellem 'samme produkt'-raekker (gengangerrapporten, brugerens ja 2026-10-16).
Brug: py -3 anvend_ean_kopi.py <ean_plan.json> [ark ...] [--dry]

Regler (praecision foerst - en forkert EAN er vaerre end en manglende):
- Planen er bygget af gengangerrapporten: kun 'samme produkt' med sikkerhed 5 + titellighed >= 0.6 (eller sikkerhed 4 + >= 0.9),
  kun inden for samme land, kun gyldige GS1-stregkoder (ikke 2xxxxxxxxxxxx), ikke hvis flere forskellige EAN er mulige, og
  aldrig en EAN der allerede findes i maalarket. Der slettes og sammenlaegges ingen raekker.
- Maalraekken findes via Source URL (rækkenumre forskydes). EAN skrives KUN hvis cellen stadig er tom. Gemt som tekst.
- Foer skrivning: backup i Produkter/Backup. Efter: kun EAN-celler maa vaere aendret og titelformlerne (A/B) skal vaere identiske.
- Arket springes over hvis det er aabent i Excel (~$-fil) eller aendres undervejs - koer igen senere.
"""
import collections, csv, datetime, json, os, re, shutil, sys, tempfile
import openpyxl

P = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))  # ...\Produkter
BACKUP = os.path.join(P, 'Backup')
SHEETS = {
    'nemlig': os.path.join(P, 'Nemlig', 'nemlig.xlsx'), 'wolt': os.path.join(P, 'DK', 'wolt.xlsx'),
    'edeka': os.path.join(P, 'edeka.xlsx'), 'drk': os.path.join(P, 'drk.xlsx'), 'aarstiderne': os.path.join(P, 'aarstiderne.xlsx'),
    'bilka': os.path.join(P, 'Bilka', 'bilka.xlsx'), 'rema': os.path.join(P, 'rema 1000', 'rema1000_version 2.xlsx'),
    'spar': os.path.join(P, 'SPAR', 'spar.xlsx'), 'dm': os.path.join(P, 'dm.xlsx'), 'rewe': os.path.join(P, 'rewe.xlsx'),
}
v = lambda x: getattr(x, 'text', x)


def formulas(ws, ia, ib):
    out = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        a, b = row[ia].value, row[ib].value
        out.append((v(a), getattr(b, 'ref', None), v(b)))
    return out


def run(name, items, dry, log):
    path = SHEETS[name]
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock) and not dry:
        print(f'{name}: SPRUNGET OVER - {os.path.basename(path)} er aaben i Excel')
        return False
    mtime = os.path.getmtime(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    ie, iu = hdr.index('EAN'), hdr.index('Source URL')
    before = formulas(ws, 0, 1)
    url2row = {}
    for r in range(2, ws.max_row + 1):
        u = ws.cell(row=r, column=iu + 1).value
        if u:
            url2row[str(u)] = r
    eans_now = {re.sub(r'\D', '', str(ws.cell(row=r, column=ie + 1).value or '')).lstrip('0') for r in range(2, ws.max_row + 1)} - {''}
    done = skipped = 0
    for it in items:
        r = url2row.get(it['url'])
        if not r:
            skipped += 1
            log.append([name, '', it['url'], it['ean'], it['donor_sheet'], it['pair'], it['conf'], 'SPRUNGET OVER: raekke findes ikke'])
            continue
        c = ws.cell(row=r, column=ie + 1)
        if c.value not in (None, ''):
            skipped += 1
            log.append([name, r, it['url'], it['ean'], it['donor_sheet'], it['pair'], it['conf'], f'SPRUNGET OVER: har allerede EAN {c.value}'])
            continue
        if it['ean'].lstrip('0') in eans_now:
            skipped += 1
            log.append([name, r, it['url'], it['ean'], it['donor_sheet'], it['pair'], it['conf'], 'SPRUNGET OVER: EAN findes allerede i arket'])
            continue
        c.value = str(it['ean'])
        c.number_format = '@'
        eans_now.add(it['ean'].lstrip('0'))
        done += 1
        log.append([name, r, it['url'], it['ean'], it['donor_sheet'], it['pair'], it['conf'], 'SKREVET'])
    wb.calculation.fullCalcOnLoad = True
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP'))
    os.close(fd)
    wb.save(tmp)
    wb.close()
    # efterkontrol paa den gemte kopi: formler uaendrede, kun EAN-celler aendret
    chk = openpyxl.load_workbook(tmp)
    cws = chk.active
    after = formulas(cws, 0, 1)
    ok = after == before
    changed_other = 0
    src = openpyxl.load_workbook(path, read_only=True)
    for r1, r2 in zip(src.active.iter_rows(min_row=2, values_only=True), cws.iter_rows(min_row=2, values_only=True)):
        for k, (x, y) in enumerate(zip(r1, r2)):
            if k < 2 or k == ie:
                continue
            if (v(x) or None) != (v(y) or None):
                changed_other += 1
    src.close()
    chk.close()
    print(f'{name}: skrevet {done}, sprunget over {skipped}; titelformler uaendrede: {ok}; andre celler aendret: {changed_other}')
    if not ok or changed_other:
        print(f'{name}: STOPPET - efterkontrollen fejlede, intet gemt')
        os.remove(tmp)
        return False
    if dry:
        print(f'{name}: proeve -> {tmp}')
        return True
    if os.path.getmtime(path) != mtime or os.path.exists(lock):
        print(f'{name}: SPRUNGET OVER - arket blev aendret/aabnet imens - koer igen')
        os.remove(tmp)
        return False
    stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
    os.makedirs(BACKUP, exist_ok=True)
    shutil.copy2(path, os.path.join(BACKUP, f'{os.path.splitext(os.path.basename(path))[0]}_{stamp}_foer_ean-kopi.xlsx'))
    shutil.move(tmp, path)
    print(f'{name}: gemt -> {path}')
    return True


def main(args):
    dry = '--dry' in args
    a = [x for x in args if not x.startswith('--')]
    plan = json.load(open(a[0], encoding='utf-8'))
    only = a[1:]
    by = collections.defaultdict(list)
    for p in plan:
        by[p['sheet']].append(p)
    log = []
    for name, items in by.items():
        if only and name not in only:
            continue
        run(name, items, dry, log)
    out = os.path.join(BACKUP if not dry else os.environ['TEMP'], f'ean_kopi_{datetime.datetime.now():%Y-%m-%d_%H%M}.csv')
    with open(out, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['ark', 'raekke', 'Source URL', 'EAN', 'donor-ark', 'par-id', 'sikkerhed', 'resultat'])
        w.writerows(log)
    print('log ->', out)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1:])
