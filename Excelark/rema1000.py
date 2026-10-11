# -*- coding: utf-8 -*-

import os
import re
import shutil
import tempfile
import unicodedata
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime

from openpyxl import load_workbook


TARGET_FILE = (
    r"C:\Users\Peter\Desktop\Hello Cal\Productdatabase"
    r"\REMA1000\rema1000_version 2.xlsx"
)

SNAPSHOT_FOLDER = (
    r"C:\Users\Peter\Desktop\Hello Cal\Productdatabase"
    r"\Excel Snapshops"
)

SHEET_NAME = "Products"


# ============================================================
# FIXED / PROTECTED COLUMNS
# ============================================================

PROTECTED_EXACT = {
    "original title",
    "subtitle",
    "category",
    "kategori",
    "produkttype",
    "product type",
    "producttype",
    "ean",
    "ean code",
    "ean-code",
    "barcode",
    "stregkode",
    "id",
    "product id",
    "productid",
}

PROTECTED_CONTAINS = {
    "url",
    "image",
    "billede",
    "source",
    "kilde",
    "raw",
    "created",
    "updated",
    "import",
}

NO_CAPITALIZATION = {
    "quantity",
    "amount",
    "%",
    "organic",
    "certificate",
}


# ============================================================
# TEXT / HEADER HELPERS
# ============================================================

def header_key(value):
    if not isinstance(value, str):
        return ""

    return re.sub(
        r"\s+",
        " ",
        value.strip().casefold()
    )


def is_protected_header(value):
    key = header_key(value)

    if key in PROTECTED_EXACT:
        return True

    return any(
        part in key
        for part in PROTECTED_CONTAINS
    )


def find_header(headers, *names):
    wanted = {
        header_key(name)
        for name in names
    }

    for col, value in headers.items():
        if header_key(value) in wanted:
            return col

    return None


def find_amount_header(headers):
    for col, value in headers.items():
        key = header_key(value)

        if (
            key == "amount"
            or key.startswith("amount ")
            or key.startswith("amount/")
        ):
            return col

    return None


def normalize_unicode(value):
    if not isinstance(value, str):
        return value

    return unicodedata.normalize(
        "NFKC",
        value
    ).replace(
        "\u00a0",
        " "
    )


def same_text(a, b):
    return (
        "" if a is None else str(a)
    ) == (
        "" if b is None else str(b)
    )


# ============================================================
# SNAPSHOTS
# ============================================================

def make_snapshot(source_path, label):
    os.makedirs(
        SNAPSHOT_FOLDER,
        exist_ok=True
    )

    base = os.path.basename(
        source_path
    )

    stem, ext = os.path.splitext(
        base
    )

    stamp = datetime.now().strftime(
        "%Y-%m-%d_%H-%M-%S-%f"
    )

    destination = os.path.join(
        SNAPSHOT_FOLDER,
        f"{stem}_{stamp}_{label}{ext}",
    )

    shutil.copy2(
        source_path,
        destination
    )

    print(
        f"Snapshot: {destination}"
    )

    return destination


# ============================================================
# UNIT NORMALIZATION
# ============================================================

UNIT_ALIASES = {
    "g": "g",
    "gr": "g",
    "gram": "g",
    "grams": "g",

    "kg": "kg",
    "mg": "mg",

    "ug": "ug",
    "\u00b5g": "ug",

    "ml": "ml",
    "cl": "cl",
    "dl": "dl",

    "l": "L",
    "ltr": "L",
    "liter": "L",
    "litre": "L",

    "stk": "stk",
    "st": "stk",
    "styk": "stk",
    "stykker": "stk",

    "pak": "pak",
    "pakke": "pak",
    "pakker": "pak",
    "pk": "pak",
    "pack": "pak",
}


UNIT_PATTERN = (
    r"(?:"
    r"stykker|styk|stk|st|"
    r"pakker|pakke|pak|pk|pack|"
    r"grams|gram|gr|"
    r"kg|mg|\u00b5g|ug|g|"
    r"ml|cl|dl|"
    r"ltr|liter|litre|l"
    r")"
)

NUMBER_PATTERN = (
    r"\d+(?:[.,]\d+)?"
)


RANGE_QUANTITY_RE = re.compile(
    rf"(?<!\d)"
    rf"({NUMBER_PATTERN})"
    rf"\s*[-\u2013\u2014]\s*"
    rf"({NUMBER_PATTERN})"
    rf"\s*({UNIT_PATTERN})"
    rf"\.?",
    re.IGNORECASE,
)


EXACT_RANGE_QUANTITY_RE = re.compile(
    rf"^\s*"
    rf"({NUMBER_PATTERN})"
    rf"\s*[-\u2013\u2014]\s*"
    rf"({NUMBER_PATTERN})"
    rf"\s*({UNIT_PATTERN})"
    rf"\.?\s*$",
    re.IGNORECASE,
)


EXACT_QUANTITY_RE = re.compile(
    rf"^\s*"
    rf"({NUMBER_PATTERN})"
    rf"\s*({UNIT_PATTERN})"
    rf"\.?\s*$",
    re.IGNORECASE,
)


def normalize_number_display(value):
    return re.sub(
        r"\s+",
        "",
        str(value)
    )


def normalize_unit(value):
    key = (
        str(value)
        .strip()
        .casefold()
        .rstrip(".")
    )

    return UNIT_ALIASES.get(
        key,
        key
    )


def normalize_exact_quantity(value):
    if not isinstance(value, str):
        return value

    match = EXACT_RANGE_QUANTITY_RE.fullmatch(
        value
    )

    if match:
        a = normalize_number_display(
            match.group(1)
        )

        b = normalize_number_display(
            match.group(2)
        )

        unit = normalize_unit(
            match.group(3)
        )

        return f"{a}-{b}{unit}"

    match = EXACT_QUANTITY_RE.fullmatch(
        value
    )

    if match:
        number = normalize_number_display(
            match.group(1)
        )

        unit = normalize_unit(
            match.group(2)
        )

        return f"{number}{unit}"

    return value


def extract_range_from_subtitle(value):
    if not isinstance(value, str):
        return None

    match = RANGE_QUANTITY_RE.search(
        value
    )

    if not match:
        return None

    a = normalize_number_display(
        match.group(1)
    )

    b = normalize_number_display(
        match.group(2)
    )

    unit = normalize_unit(
        match.group(3)
    )

    return f"{a}-{b}{unit}"


def parse_exact_stk(value):
    if not isinstance(value, str):
        return None

    match = re.fullmatch(
        rf"\s*"
        rf"({NUMBER_PATTERN})"
        rf"\s*"
        rf"(?:stk|st|styk|stykker)"
        rf"\.?\s*",
        value,
        re.IGNORECASE,
    )

    if not match:
        return None

    return (
        f"{normalize_number_display(match.group(1))}"
        f"stk"
    )


# ============================================================
# BRAND PACKAGE COUNT -> AMOUNT
# ============================================================

BRAND_PACKAGE_RE = re.compile(
    rf"(?<!\w)"
    rf"({NUMBER_PATTERN})"
    rf"\s*[-/]?\s*"
    rf"(stk|st|styk|stykker|"
    rf"pak|pakke|pakker|pk|pack)"
    rf"\.?"
    rf"(?!\w)",
    re.IGNORECASE,
)


def extract_package_from_brand(value):
    if not isinstance(value, str):
        return None

    match = BRAND_PACKAGE_RE.search(
        value
    )

    if not match:
        return None

    number = normalize_number_display(
        match.group(1)
    )

    unit = normalize_unit(
        match.group(2)
    )

    amount = (
        f"{number}{unit}"
    )

    remaining = (
        value[:match.start()]
        + " "
        + value[match.end():]
    )

    remaining = re.sub(
        r"\s+",
        " ",
        remaining
    ).strip()

    remaining = remaining.strip(
        " -/,;:"
    )

    return (
        amount,
        remaining or None
    )


# ============================================================
# ORGANIC
# ============================================================

ORGANIC_KEYS = {
    "\u00f8ko",
    "\u00f8kologi",
    "\u00f8kologisk",
    "\u00f8kologiske",
}

ORGANIC_CANONICAL = (
    "\u00d8kologisk"
)


def is_organic_marker(value):
    if not isinstance(value, str):
        return False

    key = (
        normalize_unicode(value)
        .strip()
        .casefold()
        .rstrip(".")
    )

    return key in ORGANIC_KEYS


# ============================================================
# CERTIFICATES
# ============================================================

NOEGLEHUL = (
    "N\u00f8glehul"
)

NOEGLEHULSMAERKET = (
    "N\u00f8glehulsm\u00e6rket"
)

DYRENES = (
    "Anbefalet af Dyrenes Beskyttelse"
)


def normalize_certificate(value):
    if not isinstance(value, str):
        return None

    text = re.sub(
        r"\s+",
        " ",
        normalize_unicode(value).strip()
    )

    if not text:
        return None

    key = (
        text.casefold()
        .rstrip(".")
    )

    if key == "n\u00f8glehul":
        return NOEGLEHUL

    if key == "n\u00f8glehulsm\u00e6rket":
        return NOEGLEHULSMAERKET

    if (
        key
        == "anbefalet af dyrenes beskyttelse"
    ):
        return DYRENES

    without_number = re.sub(
        r"^\d+\s*[\.)]\s*",
        "",
        key
    )

    if without_number.startswith(
        "bedre dyrevelf\u00e6rd"
    ):
        return text.rstrip(
            " ."
        )

    return None


def split_csv_values(value):
    if (
        not isinstance(value, str)
        or not value.strip()
    ):
        return []

    return [
        part.strip()
        for part in value.split(",")
        if part.strip()
    ]


def merge_unique_csv(
    existing,
    additions
):
    result = []
    seen = set()

    values = (
        split_csv_values(existing)
        + list(additions)
    )

    for value in values:
        clean = re.sub(
            r"\s+",
            " ",
            str(value).strip()
        )

        if not clean:
            continue

        key = clean.casefold()

        if key in seen:
            continue

        seen.add(
            key
        )

        result.append(
            clean
        )

    if not result:
        return None

    return ", ".join(
        result
    )


# ============================================================
# SIZE -> VARIATION
# ============================================================

SIZE_PART = (
    r"(?:xxxs|xxs|xs|s|m|l|xl|xxl|xxxl)"
)

SIZE_RE = re.compile(
    rf"^\s*"
    rf"({SIZE_PART})"
    rf"\s*/\s*"
    rf"({SIZE_PART})"
    rf"\s*$",
    re.IGNORECASE,
)

SIZE_LABEL_RE = re.compile(
    r"^\s*str\.?\s*$",
    re.IGNORECASE,
)


def normalize_size(value):
    if not isinstance(value, str):
        return None

    match = SIZE_RE.fullmatch(
        value
    )

    if not match:
        return None

    return (
        f"{match.group(1).casefold()}/"
        f"{match.group(2).casefold()}"
    )


def is_size_label(value):
    return (
        isinstance(value, str)
        and bool(
            SIZE_LABEL_RE.fullmatch(
                value
            )
        )
    )


# ============================================================
# PERCENTAGES
# ============================================================

PERCENT_RE = re.compile(
    r"(?<!\d)"
    r"(\d+(?:\s*[.,]\s*\d+)?)"
    r"\s*"
    r"[%\uff05\ufe6a]"
)


def percent_canon(
    number_text
):
    return (
        re.sub(
            r"\s+",
            "",
            number_text
        )
        .replace(",", ".")
    )


def percent_display(
    number_text
):
    number = re.sub(
        r"\s+",
        "",
        number_text
    )

    return f"{number}%"


def extract_percent_canons(
    value
):
    canons = set()

    if not isinstance(value, str):
        return canons

    for match in PERCENT_RE.finditer(
        value
    ):
        canons.add(
            percent_canon(
                match.group(1)
            )
        )

    return canons


def merge_percent_destination(
    existing,
    additions
):
    values = []
    seen = set()

    if (
        isinstance(existing, str)
        and existing.strip()
    ):
        matches = list(
            PERCENT_RE.finditer(
                existing
            )
        )

        if matches:
            for match in matches:
                canon = percent_canon(
                    match.group(1)
                )

                if canon in seen:
                    continue

                seen.add(
                    canon
                )

                values.append(
                    percent_display(
                        match.group(1)
                    )
                )

        else:
            raw = existing.strip()

            if raw:
                values.append(
                    raw
                )

                seen.add(
                    raw.casefold()
                )

    for display, canon in additions:
        if canon in seen:
            continue

        seen.add(
            canon
        )

        values.append(
            display
        )

    if not values:
        return None

    return ", ".join(
        values
    )


def remove_only_new_percentages(
    value,
    destination_canons
):
    if not isinstance(value, str):
        return value, []

    added = []

    def repl(match):
        canon = percent_canon(
            match.group(1)
        )

        # Already exists in % column:
        # DO NOT remove it from source.
        if canon in destination_canons:
            return match.group(0)

        display = percent_display(
            match.group(1)
        )

        destination_canons.add(
            canon
        )

        added.append(
            (
                display,
                canon
            )
        )

        return ""

    result = PERCENT_RE.sub(
        repl,
        value
    )

    if not added:
        return value, []

    result = re.sub(
        r"\(\s*\)",
        "",
        result
    )

    result = re.sub(
        r"\[\s*\]",
        "",
        result
    )

    result = re.sub(
        r"\s{2,}",
        " ",
        result
    ).strip()

    result = re.sub(
        r"^[\s,;:/|+\-\u2013\u2014]+",
        "",
        result
    )

    result = re.sub(
        r"[\s,;:/|+\-\u2013\u2014]+$",
        "",
        result
    ).strip()

    return (
        result or None,
        added
    )


# ============================================================
# ORIGINAL TITLE ENDING IN PERCENT
# ============================================================

ORIGINAL_PERCENT_RE = re.compile(
    r"^\s*"
    r"(.+?\S)"
    r"\s+"
    r"(\d+(?:\s*[.,]\s*\d+)?)"
    r"\s*"
    r"[%\uff05\ufe6a]"
    r"\s*$",
    re.IGNORECASE,
)


def original_title_percent_brand(
    value
):
    if not isinstance(value, str):
        return None

    match = ORIGINAL_PERCENT_RE.fullmatch(
        normalize_unicode(value)
    )

    if not match:
        return None

    text = re.sub(
        r"\s+",
        " ",
        match.group(1).strip()
    )

    percent = percent_display(
        match.group(2)
    )

    return (
        f"{text} {percent}"
    )


# ============================================================
# GENERIC TRIM + CAPITALIZATION
# ============================================================

def is_outer_trim_char(
    char
):
    if char.isspace():
        return True

    category = unicodedata.category(
        char
    )

    return (
        category.startswith("P")
        or category.startswith("S")
    )


def trim_outer_specials(
    value
):
    if not isinstance(value, str):
        return value

    value = normalize_unicode(
        value
    )

    start = 0
    end = len(value)

    while (
        start < end
        and is_outer_trim_char(
            value[start]
        )
    ):
        start += 1

    while (
        end > start
        and is_outer_trim_char(
            value[end - 1]
        )
    ):
        end -= 1

    result = value[
        start:end
    ].strip()

    return (
        result
        if result
        else None
    )


WORD_RE = re.compile(
    r"[^\W\d_]+",
    re.UNICODE
)


def capitalize_non_upper_words(
    value
):
    if not isinstance(value, str):
        return value

    def repl(match):
        word = match.group(0)

        start = match.start()
        end = match.end()

        prev_char = (
            value[start - 1]
            if start > 0
            else ""
        )

        next_char = (
            value[end]
            if end < len(value)
            else ""
        )

        # 500g / 6pak etc.
        if (
            prev_char.isdigit()
            or next_char.isdigit()
        ):
            return word

        # l/xl etc.
        if (
            prev_char == "/"
            or next_char == "/"
        ):
            return word

        # REMA / XL etc.
        if word.isupper():
            return word

        return (
            word[0].upper()
            + word[1:]
        )

    return WORD_RE.sub(
        repl,
        value
    )


# ============================================================
# BRAND PREFIX -> VARIATION
# ============================================================

def brand_tokens(value):
    if not isinstance(value, str):
        return []

    return [
        token
        for token in re.split(
            r"\s+",
            value.strip()
        )
        if token
    ]


def brand_key(tokens):
    return tuple(
        token.casefold()
        for token in tokens
    )


def detect_brand_variations(
    work,
    brand_col,
    first_data_row,
    last_row
):
    exact_counts = Counter()
    row_tokens = {}

    for row in range(
        first_data_row,
        last_row + 1
    ):
        tokens = brand_tokens(
            work[row][brand_col]
        )

        if not tokens:
            continue

        row_tokens[
            row
        ] = tokens

        exact_counts[
            brand_key(tokens)
        ] += 1

    candidates = defaultdict(
        list
    )

    for row, tokens in row_tokens.items():

        if len(tokens) < 2:
            continue

        for length in range(
            1,
            len(tokens)
        ):
            prefix = tokens[:length]
            suffix = tokens[length:]

            candidates[
                brand_key(prefix)
            ].append(
                {
                    "row": row,
                    "prefix": prefix,
                    "suffix": suffix,
                    "length": length,
                }
            )

    valid_prefixes = set()

    for (
        prefix_key,
        items
    ) in candidates.items():

        base_exists = (
            exact_counts[
                prefix_key
            ] > 0
        )

        distinct_suffixes = {
            brand_key(
                item["suffix"]
            )
            for item in items
        }

        if (
            base_exists
            or len(
                distinct_suffixes
            ) >= 2
        ):
            valid_prefixes.add(
                prefix_key
            )

    matches = {}

    for prefix_key in valid_prefixes:

        for item in candidates[
            prefix_key
        ]:
            row = item["row"]

            current = matches.get(
                row
            )

            if (
                current is None
                or item["length"]
                > current["length"]
            ):
                matches[row] = {
                    "brand":
                        " ".join(
                            item["prefix"]
                        ),

                    "variation":
                        " ".join(
                            item["suffix"]
                        ),

                    "length":
                        item["length"],
                }

    return matches


# ============================================================
# MAIN
# ============================================================

def main():

    if not os.path.isfile(
        TARGET_FILE
    ):
        raise RuntimeError(
            f"Filen findes ikke: "
            f"{TARGET_FILE}"
        )


    # ========================================================
    # SNAPSHOT BEFORE EVERYTHING
    # ========================================================

    make_snapshot(
        TARGET_FILE,
        "FOER_ALLE_REGLER"
    )


    # ========================================================
    # LOAD
    # ========================================================

    workbook = load_workbook(
        TARGET_FILE,
        data_only=False,
        keep_links=True
    )


    if (
        SHEET_NAME
        not in workbook.sheetnames
    ):
        raise RuntimeError(
            f'Arket "{SHEET_NAME}" '
            f"blev ikke fundet."
        )


    sheet = workbook[
        SHEET_NAME
    ]


    # ========================================================
    # FIND HEADER ROW
    # ========================================================

    header_row = None

    for row in range(
        1,
        min(
            sheet.max_row,
            10
        ) + 1
    ):
        row_values = [
            sheet.cell(
                row=row,
                column=col
            ).value

            for col in range(
                1,
                sheet.max_column + 1
            )
        ]

        if any(
            header_key(value)
            == "original title"

            for value
            in row_values
        ):
            header_row = row
            break


    if header_row is None:
        raise RuntimeError(
            'Kunne ikke finde '
            '"Original title".'
        )


    headers = {
        col:
            sheet.cell(
                row=header_row,
                column=col
            ).value

        for col in range(
            1,
            sheet.max_column + 1
        )
    }


    # ========================================================
    # REQUIRED COLUMNS
    # ========================================================

    original_title_col = (
        find_header(
            headers,
            "Original title"
        )
    )

    subtitle_col = (
        find_header(
            headers,
            "Subtitle"
        )
    )

    quantity_col = (
        find_header(
            headers,
            "Quantity"
        )
    )

    amount_col = (
        find_amount_header(
            headers
        )
    )

    brand_col = (
        find_header(
            headers,
            "Brand"
        )
    )

    variation_col = (
        find_header(
            headers,
            "Variation",
            "Variety",
            "Variaty"
        )
    )

    organic_col = (
        find_header(
            headers,
            "Organic"
        )
    )

    certificate_col = (
        find_header(
            headers,
            "Certificate"
        )
    )

    percent_col = (
        find_header(
            headers,
            "%"
        )
    )


    required = {
        "Original title":
            original_title_col,

        "Subtitle":
            subtitle_col,

        "Quantity":
            quantity_col,

        "Amount":
            amount_col,

        "Brand":
            brand_col,

        "Variation/Variety/Variaty":
            variation_col,

        "Organic":
            organic_col,

        "Certificate":
            certificate_col,

        "%":
            percent_col,
    }


    missing = [
        name

        for name, col
        in required.items()

        if col is None
    ]


    if missing:
        raise RuntimeError(
            "Mangler kolonner: "
            + ", ".join(
                missing
            )
        )


    first_data_row = (
        header_row + 1
    )

    last_row = sheet.max_row

    max_col = sheet.max_column

    original_max_row = (
        sheet.max_row
    )

    original_max_col = (
        sheet.max_column
    )


    # ========================================================
    # READ EVERYTHING INTO MEMORY
    # ========================================================

    original = {
        row:
            [None]
            + [
                sheet.cell(
                    row=row,
                    column=col
                ).value

                for col in range(
                    1,
                    max_col + 1
                )
            ]

        for row in range(
            first_data_row,
            last_row + 1
        )
    }


    work = deepcopy(
        original
    )


    formula_cells = {
        (
            row,
            col
        )

        for row in range(
            first_data_row,
            last_row + 1
        )

        for col in range(
            1,
            max_col + 1
        )

        if (
            sheet.cell(
                row=row,
                column=col
            ).data_type
            == "f"
        )
    }


    def protected_col(col):
        return is_protected_header(
            headers[col]
        )


    def set_value(
        row,
        col,
        value,
        rule_name
    ):

        if protected_col(col):
            raise RuntimeError(
                f"SIKKERHEDSSTOP: "
                f"Reglen {rule_name} "
                f"forsogte at aendre "
                f"{headers[col]} "
                f"paa raekke {row}."
            )


        if (
            row,
            col
        ) in formula_cells:

            raise RuntimeError(
                f"SIKKERHEDSSTOP: "
                f"Reglen {rule_name} "
                f"forsogte at overskrive "
                f"en formel i "
                f"{headers[col]} "
                f"paa raekke {row}."
            )


        work[row][col] = value


    stats = Counter()


    # ========================================================
    # 1. SUBTITLE RANGE -> QUANTITY
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        found = extract_range_from_subtitle(
            work[row][subtitle_col]
        )

        if (
            found is not None
            and not same_text(
                work[row][quantity_col],
                found
            )
        ):
            set_value(
                row,
                quantity_col,
                found,
                "subtitle_range_to_quantity"
            )

            stats[
                "subtitle_range_to_quantity"
            ] += 1


    # ========================================================
    # 2. NORMALIZE QUANTITY
    #    STK -> AMOUNT
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        value = work[
            row
        ][
            quantity_col
        ]


        stk = parse_exact_stk(
            value
        )


        if stk is not None:

            set_value(
                row,
                amount_col,
                stk,
                "quantity_stk_to_amount"
            )

            set_value(
                row,
                quantity_col,
                None,
                "quantity_stk_to_amount"
            )

            stats[
                "quantity_stk_to_amount"
            ] += 1

            continue


        normalized = (
            normalize_exact_quantity(
                value
            )
        )


        if not same_text(
            normalized,
            value
        ):
            set_value(
                row,
                quantity_col,
                normalized,
                "normalize_quantity"
            )

            stats[
                "normalize_quantity"
            ] += 1


    # ========================================================
    # 3. STK / PAK IN BRAND -> AMOUNT
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        result = (
            extract_package_from_brand(
                work[row][brand_col]
            )
        )


        if result is None:
            continue


        amount, remaining_brand = (
            result
        )


        set_value(
            row,
            amount_col,
            amount,
            "brand_package_to_amount"
        )

        set_value(
            row,
            brand_col,
            remaining_brand,
            "brand_package_to_amount"
        )


        stats[
            "brand_package_to_amount"
        ] += 1


    # ========================================================
    # 4. BRAND PREFIX -> VARIATION
    # ========================================================

    brand_matches = (
        detect_brand_variations(
            work,
            brand_col,
            first_data_row,
            last_row
        )
    )


    for row, match in (
        brand_matches.items()
    ):

        old_brand = work[
            row
        ][
            brand_col
        ]


        if (
            not isinstance(
                old_brand,
                str
            )
            or not old_brand.strip()
        ):
            continue


        suffix = (
            match["variation"]
            .strip()
        )


        if not suffix:
            continue


        set_value(
            row,
            brand_col,
            match["brand"],
            "brand_prefix_to_variation"
        )


        existing = work[
            row
        ][
            variation_col
        ]


        merged = merge_unique_csv(
            existing,
            [
                suffix
            ]
        )


        set_value(
            row,
            variation_col,
            merged,
            "brand_prefix_to_variation"
        )


        stats[
            "brand_prefix_to_variation"
        ] += 1


    # ========================================================
    # 5. ORIGINAL TITLE ENDING IN %
    #
    # CURRENT BRAND -> VARIATION
    # ORIGINAL TITLE -> BRAND
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        new_brand = (
            original_title_percent_brand(
                work[
                    row
                ][
                    original_title_col
                ]
            )
        )


        if new_brand is None:
            continue


        old_brand = work[
            row
        ][
            brand_col
        ]


        if (
            isinstance(
                old_brand,
                str
            )
            and old_brand.strip()
        ):

            existing = work[
                row
            ][
                variation_col
            ]


            merged = merge_unique_csv(
                existing,
                [
                    old_brand.strip()
                ]
            )


            set_value(
                row,
                variation_col,
                merged,
                "original_percent_brand"
            )


        set_value(
            row,
            brand_col,
            new_brand,
            "original_percent_brand"
        )


        stats[
            "original_percent_brand"
        ] += 1


    # ========================================================
    # 6. ORGANIC
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        sources = []

        found = False


        for col in range(
            1,
            max_col + 1
        ):

            if (
                protected_col(col)
                or (
                    row,
                    col
                ) in formula_cells
            ):
                continue


            if is_organic_marker(
                work[row][col]
            ):

                found = True


                if col != organic_col:
                    sources.append(
                        col
                    )


        if found:

            set_value(
                row,
                organic_col,
                ORGANIC_CANONICAL,
                "organic"
            )


            for col in sources:

                set_value(
                    row,
                    col,
                    None,
                    "organic"
                )


            stats[
                "organic"
            ] += 1


    # ========================================================
    # 7. CERTIFICATES
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        additions = []

        sources = []


        for col in range(
            1,
            max_col + 1
        ):

            if col == certificate_col:
                continue


            if (
                protected_col(col)
                or (
                    row,
                    col
                ) in formula_cells
            ):
                continue


            cert = normalize_certificate(
                work[row][col]
            )


            if cert is None:
                continue


            additions.append(
                cert
            )

            sources.append(
                col
            )


        if additions:

            merged = merge_unique_csv(
                work[
                    row
                ][
                    certificate_col
                ],
                additions
            )


            set_value(
                row,
                certificate_col,
                merged,
                "certificate"
            )


            for col in sources:

                set_value(
                    row,
                    col,
                    None,
                    "certificate"
                )


            stats[
                "certificate"
            ] += 1


    # ========================================================
    # 8. STR. + SIZE -> VARIATION
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        labels = []

        sizes = []


        for col in range(
            1,
            max_col + 1
        ):

            if (
                protected_col(col)
                or (
                    row,
                    col
                ) in formula_cells
            ):
                continue


            value = work[
                row
            ][
                col
            ]


            if is_size_label(
                value
            ):
                labels.append(
                    col
                )


            size = normalize_size(
                value
            )


            if size is not None:
                sizes.append(
                    (
                        col,
                        size
                    )
                )


        if (
            len(labels) == 1
            and len(sizes) == 1
        ):

            label_col = labels[0]

            size_col, size_value = (
                sizes[0]
            )


            if label_col != size_col:

                set_value(
                    row,
                    variation_col,
                    f"Str. {size_value}",
                    "size_to_variation"
                )


                if (
                    label_col
                    != variation_col
                ):

                    set_value(
                        row,
                        label_col,
                        None,
                        "size_to_variation"
                    )


                if (
                    size_col
                    != variation_col
                ):

                    set_value(
                        row,
                        size_col,
                        None,
                        "size_to_variation"
                    )


                stats[
                    "size_to_variation"
                ] += 1


    # ========================================================
    # 9. GLOBAL PERCENTAGE EXTRACTION
    #
    # IMPORTANT:
    # IF SAME % IS ALREADY IN % COLUMN,
    # SOURCE TEXT IS NOT CHANGED.
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        existing_destination = work[
            row
        ][
            percent_col
        ]


        destination_canons = (
            extract_percent_canons(
                existing_destination
            )
        )


        additions = []


        for col in range(
            1,
            max_col + 1
        ):

            if col == percent_col:
                continue


            if (
                protected_col(col)
                or (
                    row,
                    col
                ) in formula_cells
            ):
                continue


            value = work[
                row
            ][
                col
            ]


            if not isinstance(
                value,
                str
            ):
                continue


            (
                new_value,
                newly_added
            ) = (
                remove_only_new_percentages(
                    value,
                    destination_canons
                )
            )


            if newly_added:

                set_value(
                    row,
                    col,
                    new_value,
                    "percent_to_percent_column"
                )


                additions.extend(
                    newly_added
                )


        if additions:

            new_destination = (
                merge_percent_destination(
                    existing_destination,
                    additions
                )
            )


            set_value(
                row,
                percent_col,
                new_destination,
                "percent_to_percent_column"
            )


            stats[
                "percent_values_moved"
            ] += len(
                additions
            )


            stats[
                "percent_rows"
            ] += 1


    # ========================================================
    # 10. TRIM OUTER SPECIALS + CAPITALIZATION
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        for col in range(
            1,
            max_col + 1
        ):

            if (
                protected_col(col)
                or (
                    row,
                    col
                ) in formula_cells
            ):
                continue


            value = work[
                row
            ][
                col
            ]


            if not isinstance(
                value,
                str
            ):
                continue


            key = header_key(
                headers[col]
            )


            cleaned = value


            # Do not strip % sign from % column
            if col != percent_col:

                cleaned = (
                    trim_outer_specials(
                        cleaned
                    )
                )


            if (
                cleaned is not None
                and key
                not in NO_CAPITALIZATION
            ):

                cleaned = (
                    capitalize_non_upper_words(
                        cleaned
                    )
                )


            if not same_text(
                cleaned,
                value
            ):

                set_value(
                    row,
                    col,
                    cleaned,
                    "generic_cleanup"
                )


                stats[
                    "generic_cleanup"
                ] += 1


    # ========================================================
    # EXACT CHANGED CELLS
    # ========================================================

    changes = []


    for row in range(
        first_data_row,
        last_row + 1
    ):

        for col in range(
            1,
            max_col + 1
        ):

            old = original[
                row
            ][
                col
            ]

            new = work[
                row
            ][
                col
            ]


            if same_text(
                old,
                new
            ):
                continue


            if protected_col(
                col
            ):
                raise RuntimeError(
                    f"SIKKERHEDSSTOP: "
                    f"Beskyttet kolonne "
                    f"{headers[col]} "
                    f"paa raekke {row} "
                    f"er planlagt til aendring."
                )


            if (
                row,
                col
            ) in formula_cells:

                raise RuntimeError(
                    f"SIKKERHEDSSTOP: "
                    f"Formelcelle "
                    f"{headers[col]} "
                    f"paa raekke {row} "
                    f"er planlagt til aendring."
                )


            changes.append(
                (
                    row,
                    col,
                    new
                )
            )


    # ========================================================
    # REPORT BEFORE WRITE
    # ========================================================

    print()

    print(
        f"Raekker i arket: "
        f"{last_row - first_data_row + 1}"
    )

    print(
        f"Celler der skal aendres: "
        f"{len(changes)}"
    )

    print(
        f"Subtitle range -> Quantity: "
        f"{stats['subtitle_range_to_quantity']}"
    )

    print(
        f"Quantity normaliseret: "
        f"{stats['normalize_quantity']}"
    )

    print(
        f"stk Quantity -> Amount: "
        f"{stats['quantity_stk_to_amount']}"
    )

    print(
        f"stk/pak Brand -> Amount: "
        f"{stats['brand_package_to_amount']}"
    )

    print(
        f"Brand -> Variation: "
        f"{stats['brand_prefix_to_variation']}"
    )

    print(
        f"Original title % -> Brand: "
        f"{stats['original_percent_brand']}"
    )

    print(
        f"Organic: "
        f"{stats['organic']}"
    )

    print(
        f"Certificate: "
        f"{stats['certificate']}"
    )

    print(
        f"Str. + size -> Variation: "
        f"{stats['size_to_variation']}"
    )

    print(
        f"Procentvaerdier flyttet: "
        f"{stats['percent_values_moved']}"
    )

    print(
        f"Generisk oprydning: "
        f"{stats['generic_cleanup']}"
    )

    print()


    # ========================================================
    # APPLY EXACT CELL CHANGES
    #
    # NO ROW OPERATIONS EXIST IN THIS SCRIPT.
    # ========================================================

    for (
        row,
        col,
        new_value
    ) in changes:

        sheet.cell(
            row=row,
            column=col
        ).value = new_value


    # ========================================================
    # SAFETY CHECK BEFORE SAVE
    # ========================================================

    if (
        sheet.max_row
        != original_max_row

        or sheet.max_column
        != original_max_col
    ):

        raise RuntimeError(
            "SIKKERHEDSSTOP: "
            "Antal raekker eller kolonner "
            "har aendret sig. "
            "Filen gemmes ikke."
        )


    # ========================================================
    # VERIFY ALL PROTECTED COLUMNS
    # ========================================================

    for row in range(
        first_data_row,
        last_row + 1
    ):

        for col in range(
            1,
            max_col + 1
        ):

            if not protected_col(
                col
            ):
                continue


            before = original[
                row
            ][
                col
            ]


            after = sheet.cell(
                row=row,
                column=col
            ).value


            if not same_text(
                before,
                after
            ):

                raise RuntimeError(
                    f"SIKKERHEDSSTOP: "
                    f"Beskyttet kolonne "
                    f"{headers[col]} "
                    f"blev aendret "
                    f"paa raekke {row}."
                )


    # ========================================================
    # SAVE TO TEMP FILE FIRST
    # ========================================================

    target_dir = os.path.dirname(
        TARGET_FILE
    )


    fd, temp_path = (
        tempfile.mkstemp(
            prefix=(
                "rema1000_version2_cleanup_"
            ),
            suffix=".xlsx",
            dir=target_dir,
        )
    )

    os.close(
        fd
    )


    try:

        workbook.save(
            temp_path
        )

        workbook.close()


        # ====================================================
        # REOPEN TEMP FILE AND VALIDATE
        # ====================================================

        check = load_workbook(
            temp_path,
            data_only=False,
            keep_links=True,
            read_only=False,
        )


        check_sheet = check[
            SHEET_NAME
        ]


        if (
            check_sheet.max_row
            != original_max_row

            or check_sheet.max_column
            != original_max_col
        ):

            check.close()

            raise RuntimeError(
                "SIKKERHEDSSTOP: "
                "Temp-filen har forkert "
                "antal raekker/kolonner."
            )


        # ====================================================
        # VERIFY PROTECTED FIELDS AGAIN
        # ====================================================

        for row in range(
            first_data_row,
            last_row + 1
        ):

            for col in range(
                1,
                max_col + 1
            ):

                if not protected_col(
                    col
                ):
                    continue


                before = original[
                    row
                ][
                    col
                ]


                after = check_sheet.cell(
                    row=row,
                    column=col
                ).value


                if not same_text(
                    before,
                    after
                ):

                    check.close()

                    raise RuntimeError(
                        f"SIKKERHEDSSTOP: "
                        f"Temp-filen aendrede "
                        f"{headers[col]} "
                        f"paa raekke {row}."
                    )


        check.close()


        # ====================================================
        # ATOMIC REPLACE
        # ====================================================

        os.replace(
            temp_path,
            TARGET_FILE
        )


    finally:

        if os.path.exists(
            temp_path
        ):

            try:
                os.remove(
                    temp_path
                )

            except OSError:
                pass


    # ========================================================
    # SNAPSHOT AFTER SUCCESSFUL SAVE
    # ========================================================

    make_snapshot(
        TARGET_FILE,
        "EFTER_ALLE_REGLER"
    )


    # ========================================================
    # DONE
    # ========================================================

    print()

    print(
        "FAERDIG"
    )

    print(
        f"Skrevne celler: "
        f"{len(changes)}"
    )

    print(
        "Ingen raekker er indsat, "
        "slettet, flyttet eller sorteret."
    )

    print(
        "Original title, Subtitle, "
        "Category og Produkttype er urorte."
    )


if __name__ == "__main__":
    main()