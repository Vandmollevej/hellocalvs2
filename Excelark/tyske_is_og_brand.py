"""Tyske ark (dm, EDEKA, REWE): udfyld _is_-kolonnerne og brand/subbrand (brugerens opgave 2026-10-10).
Brug: python tyske_is_og_brand.py [dm edeka rewe] [--dry]   (uden butik: alle tre; --dry skriver kun til TEMP)

- _is_-kolonnerne udfyldes med tyske ord ud fra originalteksten (Product Name, Original Title, Subtitle) og vores
  egne felter (category, productType, variant, keyword1-5). Kun tomme celler udfyldes - intet overskrives.
  Ordene er de samme som omlaeg_til_bilka_kolonner.py bruger (bio, glutenfrei, laktosefrei, vegan, zuckerfrei,
  süßungsmittel, vollkorn, alkoholfrei, entkoffeiniert, frozen) + tilstande (geräuchert, gekocht, roh, frisch ...),
  fedt-/alkoholprocent med punktum ("3.5%") og koedtype (rind, schwein, hähnchen ...).
- brand/subbrand: maerker vi kender fra de danske ark (Bilka, REMA, SPAR, Nemlig, Wolt, DRK, Aarstiderne) kopieres
  1:1 med den danske stavemaade. Maerket skal staa forrest i den tyske titel (eller vaere Manufacturer); subbrand
  kopieres kun, hvis det samme danske subbrand ogsaa staar i titlen. Et eksisterende brand overskrives kun, naar det
  er samme maerke med en anden stavemaade, eller naar raekken er markeret BRAND_GUESS i Flag.
- Originalkolonnerne roeres aldrig (Excelark/NAVNEREGLER.md). Formlerne i A/B bevares.
- Arket gemmes oven i sig selv; foerst laegges en kopi i Excelark/backup. En CSV med alle brand-match laegges
  ved siden af backuppen, saa de kan gennemses.
"""
import collections, csv, datetime, os, re, shutil, sys, tempfile, unicodedata
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BACKUP = os.path.join(HERE, 'backup')
P_ = lambda *a: os.path.join(ROOT, *a)
GERMAN = {
    'dm': P_('Excelark', 'dm_ny.xlsx'),
    'edeka': P_('Excelark', 'edeka_ny.xlsx'),
    'rewe': P_('Excelark', 'rewe_ny.xlsx'),
}
DANISH = [  # Bilka er master: ved flere stavemaader vinder Bilkas
    P_('Excelark', 'bilka.xlsx'),
    P_('Produkter', 'rema1000_version 2.xlsx'),
    P_('Produkter', 'SPAR', 'spar.xlsx'),
    P_('Produkter', 'Nemlig', 'nemlig.xlsx'),
    P_('Excelark', 'wolt_ny.xlsx'),
    P_('Excelark', 'drk.xlsx'),
    P_('Excelark', 'aarstiderne.xlsx'),
]
ORIG = ['Original Title', 'Product Name', 'Subtitle', 'Source URL', 'Image File', 'Parse Status',
        'Manufacturer', 'Servings', 'Price', 'Venue', 'Subcategory', 'Vare', 'Variant', 'quantity']
TITLE_COLS = ['Product Name', 'Original Title']
TEXT_COLS = ['Product Name', 'Original Title', 'Subtitle', 'category', 'productType', 'variant'] + [f'keyword{n}' for n in range(1, 6)]

# ---------- _is_-regler (tyske ord) ----------
E = r'(?:e[rsnm]?)?'  # tyske boejningsendelser: vegan, vegane, veganer ...
W = lambda p: re.compile(r'(?<![\wäöüß])' + p + r'(?![\wäöüß])', re.I)
FLAGS = [  # kolonne, vaerdi, moenster (foerste kolonne der findes i arket bruges)
    ('_is_organic', 'bio', W(r'(?:bio(?:logisch' + E + r')?|öko(?:logisch' + E + r')?|organic|demeter|bioland|naturland|dmbio|alnatura)')),
    ('_is_glutenfree', 'glutenfrei', W(r'(?:glutenfrei' + E + r'|ohne gluten|gluten[- ]?free)')),
    ('_is_lactose_free', 'laktosefrei', W(r'(?:la[ck]tosefrei' + E + r'|ohne la[ck]tose|la[ck]tose[- ]?free)')),
    ('_is_vegan', 'vegan', W(r'vegan' + E)),
    ('_is_whole_grain', 'vollkorn', re.compile(r'vollkorn', re.I)),
    ('_is_decaf', 'entkoffeiniert', W(r'(?:entkoffeiniert' + E + r'|koffeinfrei' + E + r'|ohne koffein|decaf)')),
    ('_is_sweeteners', 'süßungsmittel', W(r'(?:mit süßungsmittel[n]?|süßstoff' + E + r'|mit stevia|sucralose|aspartam)')),
    ('_is_light', 'light', W(r'(?:light|lite)')),
]
SUGAR_ADDED = W(r'(?:ohne zuckerzusatz|ohne zusatz von zucker|ohne zugesetzten zucker|ohne zucker[- ]?zusatz|kein zugesetzter zucker)')
SUGAR_FREE = W(r'(?:zuckerfrei' + E + r'|ohne zucker|zero(?: sugar)?|sugar[- ]?free)')
ALC_FREE = W(r'(?:alkoholfrei' + E + r'|ohne alkohol|0[.,]0\s*%(?:\s*vol)?)')
ALC_PCT = re.compile(r'(?<![\d.,])(\d{1,2}(?:[.,]\d{1,2})?)\s*%\s*vol', re.I)
ALC_CAT = W(r'(?:bier|biere|wein|weine|sekt|spirituose[n]?|likör|liköre|schnaps|whisky|whiskey|wodka|vodka|gin|rum|cider|prosecco|champagner)')
ALC_CAT_ONLY = re.compile(r'bier|wein|sekt|spirituos|likör|alkohol', re.I)  # i category (ikke "Weinessig" i titlen)
FAT = re.compile(r'(?<![\d.,])(\d{1,2}(?:[.,]\d{1,2})?)\s*%\s*fett(?!\s*i\.?\s*tr)|fett(?:gehalt)?\s*:?\s*(\d{1,2}(?:[.,]\d{1,2})?)\s*%', re.I)
FAT_DAIRY = re.compile(r'milch|joghurt|jogurt|quark|skyr|sahne|schmand|crème fraîche|creme fraiche|kefir|buttermilch', re.I)
FAT_DAIRY_PCT = re.compile(r'(?<![\d.,])(\d(?:[.,]\d{1,2})?|10)\s*%(?!\s*[a-zäöüß])', re.I)  # 3,5% uden ord efter (ikke 15% Erdbeeren)
NO_DAIRY_FAT = re.compile(r'schokolade|kakao|riegel|milchreis|eis\b|pulver', re.I)
COOKED = [  # vaerdi, moenster
    ('geräuchert', re.compile(r'(?<![\wäöüß])geräuchert' + E + r'(?![\wäöüß])|räucher(?!stäbchen)|rauchfleisch', re.I)),
    ('gekocht', re.compile(r'(?<![\wäöüß])gekocht' + E + r'(?![\wäöüß])|kochschinken', re.I)),
    ('getrocknet', W(r'(?:getrocknet' + E + r'|gedörrt' + E + ')')),
    ('gebraten', W(r'gebraten' + E)),
    ('frittiert', W(r'frittiert' + E)),
    ('geröstet', W(r'geröstet' + E)),
    ('gegrillt', W(r'gegrillt' + E)),
    ('gebacken', W(r'(?:gebacken' + E + r'|fertig gebacken' + E + ')')),
    ('gedämpft', W(r'gedämpft' + E)),
    ('eingelegt', W(r'eingelegt' + E)),
    ('gegart', W(r'(?:vorgegart|gegart)' + E)),
]
RAW = W(r'roh' + E)
FRESH = W(r'frisch' + E)
FRESH_CTX = re.compile(r'fleisch|hack|filet|steak|schnitzel|braten|gulasch|keule|brust|hähnchen|huhn|pute[n]?|ente|lachs|fisch|forelle|kabeljau|garnele|wurst', re.I)
MEAT = [  # vaerdi, moenster (rækkefoelge: det mest specifikke foerst)
    ('kalb', re.compile(r'kalb', re.I)),
    ('rind', re.compile(r'rind|beef|rinder', re.I)),
    ('schwein', re.compile(r'schwein|pork', re.I)),
    ('hähnchen', re.compile(r'hähnchen|hühnchen|huhn|chicken|geflügel', re.I)),
    ('pute', re.compile(r'(?<![\wäöüß])pute|truthahn', re.I)),
    ('lamm', re.compile(r'lamm', re.I)),
    ('ente', re.compile(r'(?<![\wäöüß])ente(?![\wäöüß])|entenbrust|entenkeule', re.I)),
    ('wild', re.compile(r'(?<![\wäöüß])wild(?![\wäöüß])|hirsch|reh(?![\wäöüß])|wildschwein', re.I)),
]
MEAT_CTX = re.compile(r'fleisch|hack|filet|steak|schnitzel|braten|gulasch|keule|brust|wurst|würstchen|salami|schinken|speck|frikadelle'
                      r'|bulette|kotelett|nacken|rippchen|spare ?ribs|leber|geschnetzeltes|burger|patty|patties|minutensteak|roulade|ragout', re.I)
NO_MEAT = re.compile(r'vegan|vegetarisch|veggie|fleischlos|geschmack|würze|würzmischung|fond|brühe|bouillon|futter|hund|katze|chips|snack'
                     r'|sauce|soße|dressing|gewürz|marinade|fix für|nudeln|suppe', re.I)
WELFARE = re.compile(r'haltungsform\s*(\d)|(initiative tierwohl|tierwohl)', re.I)
FROZEN = W(r'(?:tk|tiefgekühlt' + E + r'|tiefgefroren' + E + r'|tiefkühl\w*|gefroren' + E + ')')
FROZEN_CAT = re.compile(r'tiefk(?:ü|ue)hl|\btk\b|frost|gefroren', re.I)
COUNTRY = {
    'deutschland': 'deutschland', 'dänemark': 'dänemark', 'italien': 'italien', 'spanien': 'spanien', 'frankreich': 'frankreich',
    'griechenland': 'griechenland', 'österreich': 'österreich', 'den niederlanden': 'niederlande', 'holland': 'niederlande',
    'polen': 'polen', 'irland': 'irland', 'norwegen': 'norwegen', 'schottland': 'schottland', 'neuseeland': 'neuseeland',
    'argentinien': 'argentinien', 'chile': 'chile', 'peru': 'peru', 'marokko': 'marokko', 'ägypten': 'ägypten', 'israel': 'israel',
    'südafrika': 'südafrika', 'der türkei': 'türkei', 'portugal': 'portugal', 'belgien': 'belgien', 'der schweiz': 'schweiz',
}
COUNTRY_RE = re.compile(r'(?<![\wäöüß])aus (' + '|'.join(sorted(map(re.escape, COUNTRY), key=len, reverse=True)) + r')(?![\wäöüß])', re.I)


def pct(num):
    s = num.replace(',', '.')
    return (s[:-2] if s.endswith('.0') else s) + '%'


def is_values(row, hdr_ix):
    """{kolonne: vaerdi} for de _is_-felter der kan afgoeres af teksten"""
    g = lambda h: str(row[hdr_ix[h]]) if h in hdr_ix and row[hdr_ix[h]] not in (None, '') else ''
    text = ' | '.join(g(h) for h in TEXT_COLS if g(h))
    cat = g('category')
    out = {}
    for col, val, rx in FLAGS:
        if rx.search(text):
            out[col] = val
    # Bio-Siegel staar ofte som "Bio" i kategorien hos dm/REWE
    if re.search(r'(?<![\wäöüß])bio(?![\wäöüß])', cat, re.I):
        out['_is_organic'] = 'bio'
    if SUGAR_ADDED.search(text):
        out['_is_sugar_free'] = 'ohne zuckerzusatz'
    elif SUGAR_FREE.search(text):
        out['_is_sugar_free'] = 'zuckerfrei'
    m = ALC_PCT.search(text)
    alc_free = bool(ALC_FREE.search(text)) or bool(m and float(m.group(1).replace(',', '.')) == 0)
    if alc_free:
        out['_is_alcohol_free'] = 'alkoholfrei'
    elif m:
        out['_is_alcohol'] = 'enthält alkohol'
        out['_is_alcohol_pct'] = pct(m.group(1))
    elif ALC_CAT_ONLY.search(cat) or (ALC_CAT.search(g('productType')) and not re.search(r'essig|soße|sauce|gummi|praline', text, re.I)):
        out['_is_alcohol'] = 'enthält alkohol'
    m = FAT.search(text)
    if m:
        out['_is_fat'] = pct(m.group(1) or m.group(2))
    elif FAT_DAIRY.search(text) and not NO_DAIRY_FAT.search(text):
        m = FAT_DAIRY_PCT.search(text)
        if m:
            out['_is_fat'] = pct(m.group(1))
    cooked = [v for v, rx in COOKED if rx.search(text)]
    if cooked:
        out['_is_cooked'] = ' und '.join(cooked)
    if RAW.search(text) and not cooked:
        out['_is_raw'] = 'roh'
    elif FRESH.search(text) and FRESH_CTX.search(text) and not cooked:
        out['_is_raw'] = 'frisch'
    if MEAT_CTX.search(text) and not NO_MEAT.search(text):
        for v, rx in MEAT:
            if rx.search(text):
                out['_is_meat'] = v
                break
    m = WELFARE.search(text)
    if m:
        out['_is_animal_welfare'] = f'haltungsform {m.group(1)}' if m.group(1) else 'tierwohl'
    if FROZEN.search(text) or FROZEN_CAT.search(cat):
        out['_is_frozen'] = 'frozen'
    m = COUNTRY_RE.search(text)
    if m:
        out['_is_country_of_origen'] = COUNTRY[m.group(1).lower()]
    return out


# ---------- brands fra de danske ark ----------
STOP_BRANDS = {  # danske maerker der ogsaa er almindelige ord forrest i en tysk titel
    'bio', 'natur', 'nature', 'natural', 'frisch', 'fresh', 'gut', 'classic', 'klassik', 'premium', 'gold', 'extra', 'mini', 'maxi',
    'select', 'original', 'kids', 'family', 'light', 'zero', 'vital', 'sport', 'pure', 'das', 'der', 'die', 'mein', 'meine', 'unser',
    'beste', 'feine', 'feinste', 'vegan', 'delikatess', 'gourmet', 'king', 'star', 'sun', 'sonne', 'alpen', 'best', 'happy', 'organic',
    'eis', 'ice', 'kaffee', 'tee', 'brot', 'milch', 'butter', 'käse', 'wurst', 'salami', 'pizza', 'pasta', 'chips', 'snack', 'snacks',
    'fit', 'free', 'smart', 'basic', 'basics', 'daily', 'go', 'easy', 'quick', 'home', 'baby', 'junior', 'max', 'pro', 'plus', 'top',
    'grill', 'bbq', 'chili', 'curry', 'mango', 'apfel', 'tomate', 'tomaten', 'honig', 'zucker', 'salz', 'öl', 'wasser', 'saft',
    'ja', 'nein', 'my', 'mr', 'mrs', 'la', 'le', 'il', 'el', 'the', 'and', 'und', 'with', 'mit', 'for', 'für', 'new', 'neu',
}


def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'[®™©]', '', s).replace('’', "'").replace('`', "'")
    return re.sub(r'\s+', ' ', re.sub(r"[^a-z0-9&'+ ]", ' ', s)).strip()


STOP = {norm(s) for s in STOP_BRANDS} | {'oko', 'eko'}


def read_sheet(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    it = ws.iter_rows(values_only=True)
    hdr = [h.strip() if isinstance(h, str) else h for h in next(it)]
    return hdr, [r for r in it if any(v not in (None, '') for v in r)]


def danish_brands():
    """norm(brand) -> (dansk stavemaade, {norm(subbrand): dansk subbrand})"""
    spell = collections.defaultdict(collections.Counter)
    subs = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    for prio, path in enumerate(DANISH):
        if not os.path.exists(path):
            print(f'   (dansk ark findes ikke: {os.path.relpath(path, ROOT)})')
            continue
        hdr, rows = read_sheet(path)
        bi = next((hdr.index(h) for h in ('brand', 'Brand') if h in hdr), None)
        si = next((hdr.index(h) for h in ('subbrand', 'Subbrand') if h in hdr), None)
        if bi is None:
            print(f'   (ingen brand-kolonne i {os.path.relpath(path, ROOT)})')
            continue
        n = 0
        w = 10 ** (len(DANISH) - prio)  # Bilkas stavemaade vinder
        for r in rows:
            b = r[bi] if bi < len(r) else None
            if not isinstance(b, str) or not b.strip():
                continue
            b = b.strip()
            k = norm(b)
            if len(k) < 3 or k in STOP or not re.search(r'[a-z]', k):
                continue
            spell[k][b] += w
            n += 1
            s = r[si] if si is not None and si < len(r) else None
            if isinstance(s, str) and len(norm(s)) >= 3 and norm(s) != k:
                subs[k][norm(s)][s.strip()] += w
        print(f'   {os.path.relpath(path, ROOT)}: {n} raekker med brand')
    out = {k: (c.most_common(1)[0][0], {sk: sc.most_common(1)[0][0] for sk, sc in subs[k].items()}) for k, c in spell.items()}
    for k in list(out):  # Kellogg's = Kellogg = Kelloggs
        for alias in (re.sub(r"'s\b", '', k), k.replace("'", '')):
            if alias != k and len(alias) >= 3 and alias not in STOP:
                out.setdefault(alias, out[k])
    return out


def match_brand(row, ix, brands, by_len):
    """(dansk brand, dansk subbrand, hvordan) eller None"""
    g = lambda h: str(row[ix[h]]) if h in ix and row[ix[h]] not in (None, '') else ''
    titles = [norm(g(h)) for h in TITLE_COLS if g(h)]
    man = norm(g('Manufacturer'))
    hit = how = None
    if man and man in brands:
        hit, how = man, 'manufacturer'
    else:
        for k in by_len:  # laengste maerke foerst (Ben & Jerry's foer Ben)
            if any(t == k or t.startswith(k + ' ') for t in titles):
                hit, how = k, 'titel'
                break
    if not hit:
        return None
    b, subs = brands[hit]
    sub = None
    for sk in sorted(subs, key=len, reverse=True):
        if any(re.search(r'(?<![a-z0-9])' + re.escape(sk) + r'(?![a-z0-9])', t[len(hit):] if t.startswith(hit) else t) for t in titles):
            sub = subs[sk]
            break
    return b, sub, how


# ---------- arkene ----------
def run(name, path, brands, by_len, dry, report):
    if not os.path.exists(path):
        print(f'{name}: findes ikke - {path}')
        return
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock) and not dry:
        print(f'{name}: SPRUNGET OVER - {os.path.basename(path)} er aaben i Excel')
        return
    mtime = os.path.getmtime(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[1]]
    ix = {}
    for i, h in enumerate(hdr):
        ix.setdefault(h, i)
    alc = [i for i, h in enumerate(hdr) if h == '_is_alcohol']
    col_of = {h: i for h, i in ix.items() if isinstance(h, str) and h.startswith('_is_')}
    if len(alc) > 1:  # Bilka: 1. _is_alcohol = markering, 2. = procent
        col_of['_is_alcohol_pct'] = alc[1]
    missing = sorted({c for c, _, _ in FLAGS} - set(col_of))
    if missing:
        print(f'   {name}: kolonner der ikke findes i arket (springes over): {missing}')
    orig_ix = {ix[h] for h in ORIG if h in ix}
    stats = collections.Counter()
    for rn, row in enumerate(ws.iter_rows(min_row=2), 2):
        vals = [c.value for c in row]
        if not any(v not in (None, '') for v in vals):
            continue
        for h, v in is_values(vals, ix).items():
            i = col_of.get(h)
            if i is None or i in orig_ix:
                continue
            if vals[i] in (None, ''):
                row[i].value = v
                stats[h] += 1
        if 'brand' not in ix or ix['brand'] in orig_ix:
            continue
        m = match_brand(vals, ix, brands, by_len)
        if not m:
            continue
        b, sub, how = m
        bi, si = ix['brand'], ix.get('subbrand')
        old_b = vals[bi] if isinstance(vals[bi], str) and vals[bi].strip() else None
        guess = 'BRAND_GUESS' in str(vals[ix['Flag']] or '') if 'Flag' in ix else False
        ob, nb = norm(old_b or '').replace("'", ''), norm(b).replace("'", '')
        same = old_b and (ob == nb or ob.startswith(nb) or nb.startswith(ob))
        if old_b and not same and not guess:
            stats['brand_konflikt_urort'] += 1
            report.append([name, rn, vals[ix['Product Name']] if 'Product Name' in ix else '', old_b, b, sub or '', how, 'konflikt - ikke aendret'])
            continue
        if old_b != b:
            row[bi].value = b
            stats['brand_ny' if not old_b else 'brand_rettet'] += 1
        if sub and si is not None and vals[si] in (None, ''):
            row[si].value = sub
            stats['subbrand'] += 1
        report.append([name, rn, vals[ix['Product Name']] if 'Product Name' in ix else '', old_b or '', b, sub or '', how,
                       'uaendret' if old_b == b else 'skrevet'])
    wb.calculation.fullCalcOnLoad = True
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP'))
    os.close(fd)
    wb.save(tmp)
    if dry:
        print(f'{name}: proeve -> {tmp}')
    else:
        if os.path.getmtime(path) != mtime or os.path.exists(lock):
            print(f'{name}: SPRUNGET OVER - {os.path.basename(path)} blev aendret/aabnet imens')
            os.remove(tmp)
            return
        os.makedirs(BACKUP, exist_ok=True)
        stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
        shutil.copy2(path, os.path.join(BACKUP, f'{os.path.splitext(os.path.basename(path))[0]}_{stamp}_foer_is_og_brand.xlsx'))
        shutil.move(tmp, path)
        print(f'{name}: gemt -> {path}')
    print('   ', dict(stats))


def main(args):
    dry = '--dry' in args
    names = [a for a in args if not a.startswith('--')] or list(GERMAN)
    print('Maerker fra de danske ark:')
    brands = danish_brands()
    by_len = sorted(brands, key=len, reverse=True)
    print(f'   {len(brands)} maerker')
    report = []
    for n in names:
        run(n, GERMAN[n], brands, by_len, dry, report)
    if report:
        os.makedirs(BACKUP, exist_ok=True)
        out = os.path.join(BACKUP if not dry else (os.environ.get('TEMP') or tempfile.gettempdir()),
                           f'tyske_brands_{datetime.datetime.now():%Y-%m-%d_%H%M}.csv')
        with open(out, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(['butik', 'raekke', 'Product Name', 'brand foer', 'brand (dansk)', 'subbrand (dansk)', 'fundet via', 'resultat'])
            w.writerows(report)
        print(f'Brand-match ({len(report)}) -> {out}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1:])
