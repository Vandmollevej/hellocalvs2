"""REMA: stregkoder som tekst (ikke 5,70604E+12) og kun Bilkas kolonner (brugeren 2026-10-10: arkene skal vaere identiske).
Brug: py ret_rema_ean.py   (arket skal vaere lukket i Excel). Roerer kun EAN-kolonnen og de ekstra kolonner bagerst."""
import datetime, os, shutil, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, 'Produkter', 'rema1000_version 2.xlsx')
BILKA = os.path.join(HERE, 'bilka.xlsx')
BACKUP = os.path.join(ROOT, 'Backup')

lock = os.path.join(os.path.dirname(SRC), '~$' + os.path.basename(SRC))
if os.path.exists(lock):
    raise SystemExit('REMA er aaben i Excel - luk den foerst')
bhdr = [c.value for c in next(openpyxl.load_workbook(BILKA, read_only=True).active.iter_rows(max_row=1))]
wb = openpyxl.load_workbook(SRC)
ws = wb.active
hdr = [c.value for c in ws[1]]
assert hdr[:len(bhdr)] == bhdr, f'REMA starter ikke med Bilkas kolonner: {hdr}'
stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
shutil.copy2(SRC, os.path.join(BACKUP, f'rema1000_version 2_{stamp}_foer_ean-tekst.xlsx'))
ean = next(h for h in ('EAN', 'barcode') if h in bhdr)
col = hdr.index(ean) + 1
n = 0
for (c,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
        c.value = str(int(c.value))
        n += 1
    c.number_format = '@'
extra = hdr[len(bhdr):]
if extra:
    ws.delete_cols(len(bhdr) + 1, len(extra))
    ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(bhdr))}{ws.max_row}'
wb.save(SRC)
print(f'{n} stregkoder -> tekst; fjernet kolonner: {extra}')
