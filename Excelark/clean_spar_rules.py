# -*- coding: utf-8 -*-
import re
import openpyxl
from collections import Counter

SRC = r"C:\Users\Peter\Desktop\Hello Cal\Productdatabase\SPAR\spar.xlsx"
OUT = r"C:\Users\Peter\Desktop\Hello Cal\Productdatabase\SPAR\spar_cleaned_preview.xlsx"
REVIEW_OUT = r"C:\Users\Peter\Desktop\Hello Cal\Productdatabase\SPAR\spar_needs_review.txt"

BRANDS_MULTI = [
    "Coca-Cola", "Coca Cola", "Häagen-Dazs", "Haagen-Dazs", "Dr. Oetker", "Ben & Jerry's",
    "Ben & Jerrys", "Red Bull", "Capri-Sun", "Capri Sun", "Carlsberg 1883", "Cup A Soup",
    "Faxe Kondi", "K-Salat", "3-Stjernet",
]
BRANDS_SINGLE = [
    "Tuborg", "Carlsberg", "Heineken", "1664", "Royal", "Faxe", "Harboe", "Grimbergen",
    "Somersby", "Kims", "Toms", "Haribo", "Malaco", "Katjes", "Ga-Jol", "Stimorol",
    "Ritter", "Lindt", "Marabou", "Fazer", "Kinder", "Nestle", "Nestlé", "Nescafe",
    "Nescafé", "Starbucks", "Lavazza", "Pickwick", "Merrild", "Arla", "Alpro",
    "Naturli", "Naturli'", "Yoggi", "Yogi", "Actimel", "Castello", "Buko", "Wyke", "Tulip",
    "Amanda", "Beauvais", "Stryhns", "Wasa", "Oetker", "Knorr", "Heinz", "Mutti",
    "Pepsi", "Fanta", "7Up", "Monster", "Innocent", "Pringles", "Snickers", "Mars",
    "Twix", "Gullon", "Gullón", "Schär", "Schar", "Corny", "Odense", "Panda",
    "Absolut", "Aalborg", "Adelhardt", "Anarkist", "Ørbæk", "Øllingegaard",
    "Wonder", "X-Ray", "TUC", "Egekilde", "Nissin", "Pukka", "Bangs",
    "Mentos", "Underberg", "Finsbury", "Ancelin", "Alte", "Wanted", "An940",
    "Barilla", "Zenzo", "Skagen", "Rynkeby",
    "President", "Lurpak", "Kærgården", "Otly", "Oatly", "Sirena",
    "Milka", "Daim", "Freia", "Cloetta", "Läkerol",
    "Extra", "Chupa Chups", "M&M's",
    "Kelloggs", "Kellogg's", "Cheerios", "Weetabix", "Quaker",
    "Klovborg", "Herslev", "Hatting", "Rynkeby", "Cocio", "Pågen", "Galbani",
    "Svaneke", "Skælskør", "Gevalia", "Bki", "BKI", "Schweppes", "Bacardi",
    "Fever-Tree", "Baileys", "Cultura", "Skyr.is", "Änglamark", "John West",
    "Uncle Ben's", "Miracoli", "Dolmio", "Queen", "Graasten", "K&K", "Easis",
    "Spicefield", "Carletti", "Bähncke", "Svansø", "Valsemøllen", "Jensens",
    "Hansens", "Evers", "Castus", "Banderos", "Katjes", "Bähncke",
]
BRANDS_MULTI_SORTED = sorted(BRANDS_MULTI, key=len, reverse=True)
BRANDS_SINGLE_SET = {b.lower(): b for b in BRANDS_SINGLE}

BEVERAGE_CATEGORIES = {"Drikkevarer", "Vin og spiritus"}
GENERIC_STOP_SUBBRAND = {"med", "og", "i", "til", "the", "den", "det", "kk", "ds", "ds.", "l", "0,0"}
DIRECT_FALLBACK_CATEGORIES = {"Frugt og grønt", "Kød", "Fisk og skaldyr"}

UNIT_MAP = {"g": "g", "gr": "g", "gram": "g", "kg": "kg", "ml": "ml", "cl": "cl", "l": "l",
            "lt": "l", "stk": "stk", "bdt": "bdt", "ps": "ps"}

PACK_RE = re.compile(r"(\d+)\s*[-]?\s*(?:x|X)\s*(?=\d)|(\d+)\s*[-]?\s*(?:pak|pk|Pak|PAK|PK)\b", re.IGNORECASE)
QTY_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(kg|g|gram|gr|ml|cl|lt|l|stk|bdt|ps)\b", re.IGNORECASE)
PERCENT_RE = re.compile(r"\d+(?:[.,]\d+)?\s*%")
OKO_RE = re.compile(r"\b(øko|økologisk)\b", re.IGNORECASE)

NOISE_CODES = {
    "gs", "gb", "fp", "kk", "dgf", "dgs", "ec", "cdo", "il", "san", "them", "ie",
    "dk", "nl", "es", "se", "no", "de", "fr", "it", "pl", "be", "co", "eg",
    "s.m.", "s.m", "ry", "gestus", "nordthy", "omhu", "anbrud", "gb/øg", "kr",
    "m", "s", "k", "v", "f", "d", "u", "a", "b", "c", "l", "p", "t", "r", "n",
    "org", "orig", "exp", "del", "le",
}

# canonical Vare dictionary: phrase (lowercase) -> canonical term. Multi-word phrases first.
VARE_PHRASES = [
    ("hakket oksekød", "Hakket oksekød"), ("hakket kalv og flæsk", "Hakket kalv og flæsk"),
    ("cup a soup", "Suppe"), ("franske kartofler", "Chips"), ("gårdchips", "Chips"),
    ("india pale ale", "Øl"), ("pale ale", "Øl"), ("new england", "Øl"),
    ("vine eddike", "Eddike"), ("ris vineddike", "Eddike"),
]
VARE_WORDS = {
    "sauce": "Sauce", "lakrids": "Lakrids", "chips": "Chips", "snack": "Snack",
    "snacks": "Snack", "pizza": "Pizza", "pasta": "Pasta", "spaghetti": "Pasta",
    "ris": "Ris", "mælk": "Mælk", "letmælk": "Mælk", "ost": "Ost", "skiveost": "Ost",
    "gouda": "Ost", "havarti": "Ost", "mozzarella": "Ost", "cheddar": "Ost",
    "boller": "Boller", "rundstykker": "Rundstykker", "kiks": "Kiks",
    "knækbrød": "Knækbrød", "is": "Is", "flødeis": "Is", "sorbet": "Sorbet",
    "øl": "Øl", "ipa": "Øl", "ale": "Øl", "pilsner": "Øl",
    "akvavit": "Snaps", "snaps": "Snaps", "gin": "Gin", "vodka": "Vodka",
    "tequila": "Tequila", "whisky": "Whisky", "prosecco": "Mousserende vin",
    "cava": "Mousserende vin", "champagne": "Champagne", "kaffe": "Kaffe",
    "espresso": "Kaffe", "latte": "Kaffe", "te": "Te", "chai": "Te",
    "juice": "Juice", "saft": "Saft", "nektar": "Saft", "sodavand": "Sodavand",
    "cola": "Sodavand", "lemonade": "Sodavand", "tonic": "Tonic",
    "yoghurt": "Yoghurt", "yogurt": "Yoghurt", "skyr": "Skyr", "smør": "Smør",
    "æg": "Æg", "mayo": "Mayonnaise", "mayonnaise": "Mayonnaise", "ketchup": "Ketchup",
    "sennep": "Sennep", "dressing": "Dressing", "remoulade": "Remoulade",
    "pesto": "Pesto", "olie": "Olie", "olivenolie": "Olivenolie", "eddike": "Eddike",
    "honning": "Honning", "marmelade": "Marmelade", "syltetøj": "Syltetøj",
    "nødder": "Nødder", "mandler": "Mandler", "peanuts": "Jordnødder",
    "popcorn": "Popcorn", "vingummi": "Vingummi", "bolcher": "Bolcher",
    "karameller": "Karameller", "tyggegummi": "Tyggegummi", "nougat": "Nougat",
    "marcipan": "Marcipan", "kage": "Kage", "brownie": "Brownie",
    "croissant": "Croissant", "franskbrød": "Franskbrød", "rugbrød": "Rugbrød",
    "hvedemel": "Mel", "mel": "Mel", "havregryn": "Havregryn", "müsli": "Müsli",
    "musli": "Müsli", "cornflakes": "Cornflakes", "cheerios": "Cornflakes",
    "kylling": "Kylling", "oksekød": "Oksekød", "hakket": "Hakket kød",
    "medister": "Medister", "spegepølse": "Spegepølse", "leverpostej": "Leverpostej",
    "bacon": "Bacon", "skinke": "Skinke", "pølser": "Pølser", "hotdog": "Hotdog",
    "laks": "Laks", "tun": "Tun", "sild": "Sild", "rejer": "Rejer",
    "kartofler": "Kartofler", "løg": "Løg", "hvidløg": "Hvidløg",
    "tomater": "Tomater", "tomat": "Tomater", "agurk": "Agurk", "salat": "Salat",
    "æble": "Æble", "banan": "Banan", "appelsin": "Appelsin", "citron": "Citron",
    "lime": "Lime", "ananas": "Ananas", "mango": "Mango", "vandmelon": "Vandmelon",
    "melon": "Melon", "pære": "Pære", "kirsebær": "Kirsebær", "solbær": "Solbær",
    "blåbær": "Blåbær", "jordbær": "Jordbær", "hindbær": "Hindbær", "druer": "Druer",
    "bønner": "Bønner", "majs": "Majs", "spinat": "Spinat", "asparges": "Asparges",
    "gulerødder": "Gulerødder", "broccoli": "Broccoli", "chokolade": "Chokolade",
    "slik": "Slik", "sirup": "Sirup", "sukker": "Sukker", "suppe": "Suppe",
    "vin": "Vin", "salami": "Salami", "rullepølse": "Rullepølse",
    "kammerjunkere": "Kammerjunkere", "vafler": "Vafler", "nudler": "Nudler",
    "havarti": "Ost", "fromage": "Flødeost", "cremefraiche": "Creme fraiche",
    "creme": "Creme fraiche",
    # vin og spiritus
    "rom": "Rom", "cognac": "Cognac", "likør": "Likør", "cider": "Cider",
    "portvin": "Portvin", "sherry": "Sherry", "chardonnay": "Hvidvin",
    "riesling": "Hvidvin", "sauvignon": "Hvidvin", "grigio": "Hvidvin",
    "cabernet": "Rødvin", "cab": "Rødvin", "merlot": "Rødvin", "shiraz": "Rødvin",
    "zinfandel": "Rødvin", "primitivo": "Rødvin", "malbec": "Rødvin",
    "noir": "Rødvin", "rosé": "Rosévin", "rose": "Rosévin", "rosevin": "Rosévin",
    "whisky": "Whisky", "whiskey": "Whisky", "bitter": "Bitter",
    # kiosk/slik/snack
    "cookies": "Cookies", "cookie": "Cookies", "pastiller": "Pastiller",
    "skumfiduser": "Skumfiduser", "marshmallows": "Skumfiduser",
    "energidrik": "Energidrik", "energy": "Energidrik", "proteinbar": "Proteinbar",
    # brød og kager
    "lagkage": "Lagkage", "wienerbrød": "Wienerbrød", "ciabatta": "Ciabatta",
    "toast": "Toastbrød", "muffins": "Muffins", "donut": "Donut",
    "marcipanbrød": "Marcipanbrød", "flute": "Flute", "pandekager": "Pandekager",
    "kagebland": "Kagemix",
    # mejeri/frost
    "koldskål": "Koldskål", "fromageblanc": "Fromage blanc", "ymer": "Ymer",
    "skyr.is": "Skyr", "kefir": "Kefir",
    # pålæg/kød tilføjelser
    "frikadeller": "Frikadeller", "kødboller": "Kødboller", "schnitzel": "Schnitzel",
    "pølse": "Pølser", "roastbeef": "Roastbeef", "spegepølser": "Spegepølse",
}


def norm_pack(n):
    return f"{n}-pak"


def match_vare_dict(tokens):
    joined = " ".join(t.lower() for t in tokens)
    for phrase, canon in VARE_PHRASES:
        if phrase in joined:
            return canon
    for t in tokens:
        key = t.lower().strip(".,")
        if key in VARE_WORDS:
            return VARE_WORDS[key]
    return None


def extract(orig_title, category):
    if not orig_title:
        return {"brand": None, "subbrand": None, "vare": None, "mgd": None,
                "antal": None, "oko": None, "keywords": [], "needs_review": False}

    text = orig_title
    oko = None
    if OKO_RE.search(text):
        oko = "Økologisk"
        text = OKO_RE.sub(" ", text)

    text = PERCENT_RE.sub(" ", text)

    antal = None
    m = PACK_RE.search(text)
    if m:
        n = m.group(1) or m.group(2)
        antal = norm_pack(n)
        text = PACK_RE.sub(" ", text)

    mgd = None
    m2 = QTY_RE.search(text)
    if m2:
        num, unit = m2.group(1), m2.group(2).lower()
        unit = UNIT_MAP.get(unit, unit)
        mgd = f"{num}{unit}"
        text = QTY_RE.sub(" ", text)

    text = re.sub(r"[.,;:]+$", "", text)
    tokens = [t for t in re.split(r"\s+", text.strip()) if t]

    brand = None
    remaining_tokens = tokens[:]

    lower_orig = orig_title.lower()
    for mb in BRANDS_MULTI_SORTED:
        if mb.lower() in lower_orig:
            brand = mb
            mb_words = set(mb.lower().split())
            remaining_tokens = [t for t in remaining_tokens if t.lower() not in mb_words]
            break

    brand_pos = None
    if brand is None:
        for i, t in enumerate(remaining_tokens):
            key = t.lower().strip(".,")
            if key in BRANDS_SINGLE_SET:
                brand = BRANDS_SINGLE_SET[key]
                brand_pos = i
                remaining_tokens = remaining_tokens[:i] + remaining_tokens[i+1:]
                break

    leftover = [t for t in remaining_tokens if t.lower().strip(".,") not in NOISE_CODES and t.strip()]

    subbrand = None
    if brand and category in BEVERAGE_CATEGORIES and leftover:
        cand = leftover[0].lower().strip(".,")
        if cand not in GENERIC_STOP_SUBBRAND and len(leftover[0]) > 1 and cand not in VARE_WORDS:
            subbrand = leftover[0]
            leftover = leftover[1:]

    vare = match_vare_dict(leftover)
    needs_review = False
    if vare is None:
        if category in DIRECT_FALLBACK_CATEGORIES and leftover:
            vare = " ".join(leftover)
            leftover = []
        else:
            needs_review = True

    keywords = leftover if vare and " ".join(leftover) != vare else (leftover if not vare else [])
    # remove tokens already used as vare text from keywords
    if vare:
        vare_tokens_lower = set(vare.lower().split())
        keywords = [t for t in leftover if t.lower() not in vare_tokens_lower]
    else:
        keywords = leftover

    return {"brand": brand, "subbrand": subbrand, "vare": vare, "mgd": mgd,
            "antal": antal, "oko": oko, "keywords": keywords, "needs_review": needs_review}


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Products"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    idx = {h: i for i, h in enumerate(header)}

    orig_i = idx["Original Title"]
    cat_i = idx["Category"]

    results = []
    cat_stats = Counter()
    cat_filled = Counter()
    review_rows = []
    for r in rows[1:]:
        title = r[orig_i]
        cat = r[cat_i]
        res = extract(title, cat)
        results.append(res)
        cat_stats[cat] += 1
        if res["vare"]:
            cat_filled[cat] += 1
        else:
            review_rows.append((cat, title))

    total = len(results)
    filled_total = sum(1 for res in results if res["vare"])
    print(f"Total rows: {total}")
    print(f"Vare filled: {filled_total} ({filled_total/total:.1%})")
    print()
    print("Per category coverage:")
    for cat in cat_stats:
        n = cat_stats[cat]
        f = cat_filled[cat]
        print(f"  {cat:30s} {f:5d}/{n:5d}  ({f/n:.1%})")

    with open(REVIEW_OUT, "w", encoding="utf-8") as f:
        for cat, title in review_rows:
            f.write(f"{cat}\t{title}\n")
    print(f"\nWrote {len(review_rows)} needs-review rows to {REVIEW_OUT}")

    max_kw = 6
    new_header = ["HelloCal_Title", "Product Name", "Original Title", "Quantity", "Brand", "Subbrand",
                  "EAN", "Pack Count", "Category", "Subcategory", "Vare", "Øko", "Mængde", "Antal"]
    new_header += [f"Keyword {i}" for i in range(1, max_kw + 1)]
    new_header += ["Source URL", "Manufacturer", "Image File", "Parse Status"]

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Products"
    ws_out.append(new_header)

    for r, res in zip(rows[1:], results):
        row_dict = dict(zip(header, r))
        kw = res["keywords"][:max_kw]
        kw += [None] * (max_kw - len(kw))
        out_row = [
            row_dict.get("HelloCal_Title"), row_dict.get("Product Name"), row_dict.get("Original Title"),
            row_dict.get("Quantity"), res["brand"] or row_dict.get("Brand"), res["subbrand"],
            row_dict.get("EAN"), row_dict.get("Pack Count"), row_dict.get("Category"),
            row_dict.get("Subcategory"), res["vare"], res["oko"], res["mgd"], res["antal"],
        ] + kw + [row_dict.get("Source URL"), row_dict.get("Manufacturer"), row_dict.get("Image File"),
                  row_dict.get("Parse Status")]
        ws_out.append(out_row)

    wb_out.save(OUT)
    print("Saved preview to", OUT)


if __name__ == "__main__":
    main()
