# -*- coding: utf-8 -*-
"""SPAR (Produkter/SPAR/spar.xlsx): brugerens regler 2026-10-10.

Uden --alle springes vin-rækkerne (Category = 'Vin og spiritus') over; med --alle rettes kun de
ord brugeren har navngivet (anb/anbrud, zinf) i deres Keyword-celler — alt andet på vin
(Product Type/Brand/Subbrand/Variation, lande, delte ord) renses af vin-sessionen.
Skriver kun i egne kolonner (Keyword 1-5, Product Type, Variation, Subbrand,
Product type plural, _is_country_of_origen, _is_organic, _is_glutenfree,
_is_lactose_free) — aldrig i A/B (formler) eller originalkolonnerne.

Regler:
- 'anb' / 'anbrud' (= anbrud) slettes.
- 'Nbdt' (fx 14bdt) = 'N per bundt'.
- Alt med Øko: _is_organic = 'økologisk', selve øko-ordet slettes fra de øvrige egne kolonner.
- Lande i _is_country_of_origen skrives med stort begyndelsesbogstav.
- Keyword 1-5: ord der allerede står i brand/subbrand/produkttype/variant/_is_-felter
  (eller er afkortede/stavefejl af dem) slettes (intet ord to gange i samme række);
  delte fraser samles; huller i Keyword 1-5 lukkes.

Kør uden argument = tørløb (skriver kun rapport). Kør med --apply for at skrive.
"""
import os
import re
import sys
import shutil
import datetime
import collections
import unicodedata

import openpyxl

SRC = r"C:\Users\Peter\Desktop\Hello Cal\Produkter\SPAR\spar.xlsx"
BACKUP_DIR = r"C:\Users\Peter\Desktop\Hello Cal\Backup"
REPORT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "spar_ikke_vin_rapport.txt")
WINE_CATEGORY = "Vin og spiritus"

# Dokumenterede keywords (NAVNEREGLER.md): fjernes kun ved eksakt dobbeltord, aldrig som "afkortet".
KEEP_STANDARD = {"revet", "i blok", "sød", "sur", "søde", "groft", "grove", "fint", "fine",
                 "dobbelte", "instant", "langtidsholdbar", "zero"}
# Frasen står som ét keyword (aldrig delt midt i en frase).
PHRASES = {
    ("ice", "cream"): "ice cream",
    ("ice", "tea"): "ice tea",
    ("sour", "cream"): "sour cream",
    ("crushed", "ice"): "crushed ice",
    ("high", "protein"): "high protein",
    ("tutti", "frutti"): "tutti frutti",
    ("hubba", "bubba"): "hubba bubba",
    ("caf", "snitter"): "cafésnitter",
    ("san", "pellegrino"): "san pellegrino",
}
# Sikre stavefejl/forkortelser i keywords (rettes FØR dobbeltord-tjekket, så de kan blive slettet
# hvis ordet allerede står i rækken).
REPLACE = {
    "bluberry": "blueberry", "ceasar": "caesar", "cremefraihce": "creme fraiche",
    "sourcream": "sour cream", "icecream": "ice cream", "sourc": "sour cream",
    "parboild": "parboiled", "parboil": "parboiled", "parb": "parboiled",
    "varmrg": "varmrøget", "citronpeb": "citronpeber", "pebe": "peber",
    "grov": "grove", "grovv": "grove", "noäl": "noël", "pavä": "pavé",
    "sukkker": "sukker", "pelleg": "san pellegrino", "pellegr": "san pellegrino",
    "sanpellegrino": "san pellegrino", "tipo00": "tipo 00", "tipo0": "tipo 0",
    "6mdr": "6 mdr", "kl1": "klasse 1",
}
# Zinf = zinfandel (brugerens regel 2026-10-10); gælder også vin-rækkernes keywords.
ZINF = {"zinf", "zinfa", "zinfandl", "zinfandel14"}
FUNCTION_WORDS = {"og", "i", "med", "uden", "af", "på", "til", "fra", "the", "de", "la", "le", "du", "&"}
FLAG_FROM_KEYWORD = {
    "glf": ("_is_glutenfree", "glutenfri"),
    "lkf": ("_is_lactose_free", "laktosefri"),
}
DROP_WORDS = {"anb", "anbrud"}


def fold(s):
    s = re.sub(r"[äöüåæø]", "#", str(s).lower())
    # accenter væk (viña = vina, château = chateau); # (æøå/äöü) bevares som jokertegn
    s = unicodedata.normalize("NFKD", s)
    return "".join(ch for ch in s if not unicodedata.combining(ch))


def words(s):
    return [w for w in re.split(r"[^\w#\-]+", fold(s)) if w]


def lev(a, b):
    if abs(len(a) - len(b)) > 2:
        return 9
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def is_subsequence(t, w):
    it = iter(w)
    return all(ch in it for ch in t)


def cap_country(v):
    parts = [p.strip() for p in str(v).split("/")]
    out = []
    for p in parts:
        out.append(" ".join(x[:1].upper() + x[1:] if x else x for x in p.split(" ")))
    return " / ".join(out)


def strip_words(text, drop):
    keep = [w for w in str(text).split() if w.strip(".,;:").lower() not in drop]
    return " ".join(keep)


def strip_oko(text):
    """Fjerner øko-ord; sammensætninger ('Økohavregryn') mister kun øko-delen."""
    out = []
    removed = False
    for tok in str(text).split():
        low = tok.lower().strip(".,;:")
        if re.fullmatch(r"øko\.?|økol\.?|økologisk|økoo|øko\d+\w*", low):
            removed = True
            continue
        m = re.fullmatch(r"(øko)([a-zæøåäö]{4,})", low)
        if m:
            rest = m.group(2)
            if tok[:1].isupper():
                rest = rest[:1].upper() + rest[1:]
            out.append(rest)
            removed = True
            continue
        out.append(tok)
    return " ".join(out), removed


def main():
    apply = "--apply" in sys.argv
    mtime0 = os.path.getmtime(SRC)
    wb = openpyxl.load_workbook(SRC)
    ws = wb["Products"]
    hdr = [c.value for c in ws[1]]
    col = {}
    for i, h in enumerate(hdr, 1):
        col.setdefault(h, i)
    i_kw = [col["Keyword %d" % n] for n in range(1, 6)]
    i_type, i_var, i_sub, i_brand = col["Product Type"], col["Variation"], col["Subbrand"], col["Brand"]
    i_plural, i_cat, i_pack = col["Product type plural"], col["Category"], col["Packaging"]
    i_country, i_org = col["_is_country_of_origen"], col["_is_organic"]
    i_glf, i_lkf = col["_is_glutenfree"], col["_is_lactose_free"]
    i_orig = col["Original Title"]
    ctx_cols = [i_brand, i_sub, i_type, i_var, i_pack, i_country, i_org, i_glf, i_lkf,
                col["_is_sweeteners"], col["_is_vegan"], col["_is_sugar_free"], col["_is_fat"],
                col["_is_whole_grain"], col["_is_animal_welfare"], col["_is_frozen"],
                col["_is_raw"], col["_is_cooked"], col["_is_decaf"], col["_is_alcohol_free"]]
    editable = set(i_kw) | {i_type, i_var, i_sub, i_plural, i_country, i_org, i_glf, i_lkf}

    changes = []   # (row, header, old, new, reason)
    reasons = collections.Counter()
    kept = collections.Counter()
    kept_sample = {}
    merged_pairs = collections.Counter()
    deleted = []   # (row, token, grund, match)

    def put(r, ci, new, reason):
        old = ws.cell(r, ci).value
        if old == new:
            return
        assert ci in editable, hdr[ci - 1]
        ws.cell(r, ci).value = new
        changes.append((r, hdr[ci - 1], old, new, reason))
        reasons[reason] += 1

    include_wine = "--alle" in sys.argv

    for r in range(2, ws.max_row + 1):
        is_wine = ws.cell(r, i_cat).value == WINE_CATEGORY
        if is_wine and not include_wine:
            continue
        organic = False

        # vin-rækker: KUN Keyword 1-5 (brugerens afgrænsning 2026-10-10 "kun under keywords");
        # Product Type/Brand/Subbrand/Variation/lande/flag på vin ejes af vin-sessionen.
        # 1) Øko / anbrud i produkttype, variant, subbrand (aldrig brand)
        for ci in (() if is_wine else (i_type, i_var, i_sub)):
            v = ws.cell(r, ci).value
            if not isinstance(v, str):
                continue
            nv, removed = strip_oko(v)
            if removed:
                organic = True
            nv2 = strip_words(nv, DROP_WORDS)
            if nv2 != v:
                put(r, ci, nv2 or None, "øko/anbrud i " + hdr[ci - 1])
                if ci == i_type and ws.cell(r, i_plural).value:
                    pv = ws.cell(r, i_plural).value
                    pv2 = strip_words(strip_oko(pv)[0], DROP_WORDS)
                    if pv2 != pv:
                        put(r, i_plural, pv2 or None, "øko/anbrud i plural")
        brand = ws.cell(r, i_brand).value
        if isinstance(brand, str) and brand.lower().startswith("øko"):
            organic = True

        # 2) land med stort begyndelsesbogstav
        v = ws.cell(r, i_country).value
        if not is_wine and isinstance(v, str) and cap_country(v) != v:
            put(r, i_country, cap_country(v), "land med stort")

        # 3) keywords
        kws = [ws.cell(r, ci).value for ci in i_kw]
        ctx = set()
        for ci in ctx_cols:
            v = ws.cell(r, ci).value
            if v:
                ctx.update(words(v))
        orig = fold(ws.cell(r, i_orig).value or "")

        if is_wine:
            # Vin: KUN de ord brugeren har navngivet, og kun i Keyword-cellerne (anb/anbrud slettes,
            # zinf → zinfandel; står zinfandel allerede i rækken, slettes keywordet). Delte vinord,
            # druer, brand/variant m.m. renses af vin-sessionen — her røres intet andet.
            wine_kws = []
            for k in kws:
                if not k:
                    continue
                kl = str(k).strip()
                low = kl.lower()
                if low in DROP_WORDS:
                    reasons["vin: anb slettet"] += 1
                    continue
                if low in ZINF:
                    if "zinfandel" in ctx:
                        reasons["vin: zinf slettet (zinfandel står allerede)"] += 1
                        continue
                    kl = "zinfandel"
                    reasons["vin: zinf -> zinfandel"] += 1
                wine_kws.append(kl)
            for ci, old, new in zip(i_kw, kws, (wine_kws + [None] * 5)[:5]):
                if old != new:
                    put(r, ci, new, "keyword-felt rettet (vin)")
            continue

        # delte fraser samles først (så "ice cream" ikke kun halveres af dobbeltord-tjekket)
        flat = [str(k).strip() for k in kws if k]
        pre = []
        i = 0
        while i < len(flat):
            if i + 1 < len(flat) and (flat[i].lower(), flat[i + 1].lower()) in PHRASES:
                pre.append(PHRASES[(flat[i].lower(), flat[i + 1].lower())])
                merged_pairs[(flat[i].lower(), flat[i + 1].lower())] += 1
                i += 2
            else:
                pre.append(flat[i])
                i += 1
        new_kws = []
        for k in pre:
            if not k:
                continue
            kl = str(k).strip()
            low = kl.lower()
            if low in ZINF:
                reasons["stavefejl/forkortelse rettet"] += 1
                kl = low = "zinfandel"
            elif low in REPLACE:
                reasons["stavefejl/forkortelse rettet"] += 1
                kl = REPLACE[low]
                low = kl
            # anbrud
            if low in DROP_WORDS:
                reasons["anb slettet"] += 1
                continue
            # øko
            if re.fullmatch(r"øko\.?|økol\.?|økologisk|økoo|øko\d+\w*", low):
                organic = True
                reasons["øko-keyword slettet"] += 1
                continue
            # Nbdt = N per bundt
            m = re.fullmatch(r"(\d+)\s*bdt\.?", low)
            if m:
                new_kws.append("%s per bundt" % m.group(1))
                reasons["Nbdt -> N per bundt"] += 1
                continue
            if low == "per bundt" and any(re.fullmatch(r"\d+\s*bdt\.?", str(x or "").lower()) for x in kws):
                reasons["per bundt slettet (står i N per bundt)"] += 1
                continue
            # flag-forkortelser
            if low in FLAG_FROM_KEYWORD:
                col_i = i_glf if low == "glf" else i_lkf
                flag_val = FLAG_FROM_KEYWORD[low][1]
                if not ws.cell(r, col_i).value:
                    put(r, col_i, flag_val, "flag fra keyword " + low)
                reasons["flag-keyword slettet"] += 1
                continue
            tw = words(kl)
            sig = [w for w in tw if w not in FUNCTION_WORDS]
            if sig:
                # fuldt dobbeltord
                if all(w in ctx for w in sig):
                    reasons["dobbeltord slettet"] += 1
                    deleted.append((r, kl, "dobbeltord", ""))
                    continue
                ok = low not in KEEP_STANDARD
                how = []
                for w in sig:
                    if w in ctx:
                        continue
                    hit = None
                    if len(w) >= 4 or (len(w) == 3 and re.search(r"\b" + re.escape(w) + r"(?:[./]|$)", orig)):
                        hit = next((c for c in sorted(ctx) if c.startswith(w) and len(c) > len(w)), None)
                        if hit:
                            how.append("præfiks:" + hit)
                    if not hit and len(w) >= 5:
                        hit = next((c for c in sorted(ctx) if len(c) >= 5 and lev(w, c) <= (2 if len(w) >= 9 else 1)), None)
                        if hit:
                            how.append("stavning:" + hit)
                    if not hit and len(w) >= 4:
                        hit = next((c for c in sorted(ctx) if c[:1] == w[:1] and 0 < len(c) - len(w) <= 3 and is_subsequence(w, c)), None)
                        if hit:
                            how.append("delfølge:" + hit)
                    if not hit:
                        ok = False
                        break
                if ok:
                    reasons["afkortet/stavefejl-dobbeltord slettet"] += 1
                    deleted.append((r, kl, "afkortet", ",".join(how)))
                    continue
            new_kws.append(kl)

        for k in new_kws:
            kept[k] += 1
            kept_sample.setdefault(k, (r, ws.cell(r, i_orig).value))

        padded = (new_kws + [None] * 5)[:5]
        for ci, old, new in zip(i_kw, kws, padded):
            if old != new:
                put(r, ci, new, "keyword-felt rettet")

        # 4) økologisk-flag
        if organic and ws.cell(r, i_org).value != "økologisk":
            put(r, i_org, "økologisk", "_is_organic udfyldt")

    # rapport
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write("SPAR ikke-vin rettelser (%s)\n" % ("ANVENDT" if apply else "TØRLØB"))
        f.write("Ændrede celler: %d\n" % len(changes))
        for k, v in reasons.most_common():
            f.write("  %5d  %s\n" % (v, k))
        f.write("\nSamlede fraser: %r\n" % dict(merged_pairs))
        f.write("\nTilbageblevne keywords (alle):\n")
        for k, v in kept.most_common():
            f.write("  %4d %s\t%s\n" % (v, k, kept_sample[k]))
        f.write("\nSlettede (afkortet/stavning):\n")
        for d in deleted:
            if d[2] == "afkortet":
                f.write("%d\t%s\t%s\t| %s\n" % (d[0], d[1], d[3], ws.cell(d[0], i_orig).value))
        f.write("\nÆndringslog:\n")
        for ch in changes:
            f.write("%d\t%s\t%r\t%r\t%s\n" % ch)
    print("changes", len(changes), dict(reasons))

    if not apply:
        return

    # skrivebeskyttelse: filen må ikke være åben i Excel, og må ikke være ændret siden vi læste den
    lock = os.path.join(os.path.dirname(SRC), "~$" + os.path.basename(SRC))
    if os.path.exists(lock):
        print("LÅST: spar.xlsx er åben i Excel — intet skrevet")
        sys.exit(2)
    if os.path.getmtime(SRC) != mtime0:
        print("ÆNDRET af andre under kørsel — intet skrevet, kør igen")
        sys.exit(3)
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
    backup = os.path.join(BACKUP_DIR, "spar_%s_foer_ikke_vin_rettelser.xlsx" % stamp)
    shutil.copy2(SRC, backup)
    tmp = SRC + ".tmp_ikke_vin"
    wb.save(tmp)
    os.replace(tmp, SRC)
    print("skrevet; backup:", backup)


if __name__ == "__main__":
    main()
