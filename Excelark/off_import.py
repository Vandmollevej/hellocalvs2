# -*- coding: utf-8 -*-
"""Open Food Facts -> butiksark (Lidl, Aldi, Kvickly/Coop) i Bilkas 45 kolonner + kvalitetskontrol.

Brug:  py -3 Excelark\\off_import.py [lidl|aldi|kvickly ...]     (uden argument: alle)
Læser rå OFF-JSON fra Backup\\off_raa (hentet med off_hent.py), vælger kun varer solgt i Danmark med et
af butikkens egne mærker, lægger dem i Bilkas kolonner og krydstjekker felterne mod hinanden
(label / navn / ingredienser / næringsindhold). Rettelser i hånden: Excelark\\off_rettelser.json
({"<EAN>": {"Product Type": "...", "Variation": "...", "_drop": true, ...}}), anvendes efter reglerne.
Alt der er gættet, mistænkeligt eller modstridende står i Excelark\\off_kvalitetskontrol.xlsx.
"""
import collections, datetime, glob, json, os, re, shutil, sys, tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, 'Backup', 'off_raa')
BILKA = os.path.join(HERE, 'bilka.xlsx')
FIXES = os.path.join(HERE, 'off_rettelser.json')

# ---------------------------------------------------------------- butikkernes egne mærker
STORES = {
    'lidl': dict(
        file='lidl.xlsx', chain='Lidl', store_tags=['lidl'],
        brands=['Lidl', 'Vilstrupgård', "Butcher's", 'Engvang', 'Milbona', 'Sweet Corner', 'Snack Day', 'Favorina',
                'Karlens', 'Solevita', 'Crownfield', 'Vemondo', 'Alesto', 'Bon Gelati', 'Dulano', 'Culinea',
                'Freeway', 'Pilos', 'Pikok', 'Chef Select', 'Deluxe', 'Eridanous', 'Baresa', 'Belbake', 'Combino',
                'McEnnedy', 'Fin Carré', 'Kania', "Vita d'Or", 'Gelatelli', 'Pastella', 'Envia', 'Salina',
                'Tastino', 'Perlenbacher', 'Rivercote', 'Lovilio', 'Maribel', 'Kena', 'Mondo Italiano', 'Italiamo',
                'Sol & Mar', 'Duc De Coeur', 'Saskia', 'Freshona', 'Sondey', 'Vitasia', 'Bellarom', "Pomo d'Oro",
                'Oceansea', 'Trattoria Alfredo', "Rio d'Oro", 'Nutrisso', 'Paesano']),
    'aldi': dict(
        file='aldi.xlsx', chain='Aldi', store_tags=['aldi', 'aldi-nord', 'aldi-danmark'],
        brands=['Aldi', 'Milsani', 'Biscotto', 'Økolivet', 'Mamia', 'Tandil', 'Nur Nur', 'Cucina Nobile',
                'Moser Roth', 'Choceur', 'GutBio', 'Casa Morando', 'Golden Bridge', 'Primo Gusto', 'Fairglobe',
                'Karlsquell', 'Meine Metzgerei', "Mike Mitchell's", 'Alpenfest', 'Lord Nelson', 'Mühlenhof']),
    'kvickly': dict(
        file='kvickly.xlsx', chain='Coop', store_tags=['kvickly', 'coop', 'superbrugsen', 'brugsen', '365discount'],
        brands=['Coop', 'Änglamark', 'Irma', 'Xtra', 'Cirkel', '365Minirisk', 'Coop 365', 'Coop Budget',
                'Coop Naturlig']),
}
# alternative stavemåder -> kanonisk (efter slug)
BRAND_ALIAS = {'oko-livet': 'Økolivet', 'anglamark': 'Änglamark', 'x-tra': 'Xtra', 'cirkel-kaffe': 'Cirkel',
               '365-minirisk': '365Minirisk', 'vita-d-or': "Vita d'Or", 'butcher-s': "Butcher's"}
CHAIN_ONLY = {'lidl', 'aldi', 'aldi-nord'}  # kædenavn alene = ikke et specifikt mærke ("Coop" er selv et egetmærke)
DROP_BRAND_TOKENS = {'lidl', 'aldi', 'coop', 'coop-365', 'kvickly', 'superbrugsen', 'brugsen'}

# ---------------------------------------------------------------- små værktøjer


def slug(s):
    s = (s or '').lower().replace('ø', 'o').replace('å', 'a').replace('æ', 'ae').replace('ä', 'a').replace('ö', 'o')
    s = s.replace('ü', 'u').replace('é', 'e').replace('’', "'")
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')


def row_formula(tpl, row):
    return re.sub(r'(?<![A-Za-z_.])([A-Z]{1,2})2(?![0-9A-Za-z])', lambda m: f'{m.group(1)}{row}', tpl)


def clean(s):
    if s is None:
        return ''
    return re.sub(r'\s+', ' ', str(s)).strip()


def lower1(s):
    s = clean(s)
    if not s or s.isupper() and len(s) <= 5:
        return s
    return s[0].lower() + s[1:]


def norm_qty(q, pq=None, pu=None):
    """-> (Quantity i Bilka-stil, Pack Count eller None, mistanke)"""
    q = clean(q).lower().replace(',', '.')
    if q:
        q = q.replace('liter', 'l').replace('ltr', 'l').replace('pcs', 'stk').replace('pieces', 'stk').replace('piece', 'stk')
        q = q.replace('stk.', 'stk').replace('gram', 'g').replace('stück', 'stk')
        m = re.search(r'(\d+)\s*[x×]\s*(\d+(?:\.\d+)?)\s*(kg|g|ml|cl|l|stk)\b', q)
        if m:
            return f'{m.group(1)} x {float(m.group(2)):g} {m.group(3)}', int(m.group(1)), ''
        m = re.search(r'(\d+(?:\.\d+)?)\s*(kg|g|ml|cl|l|stk)\b', q)
        if m:
            num = float(m.group(1))
            unit = m.group(2)
            pack = int(num) if unit == 'stk' else None
            return f'{num:g} {unit}', pack, ''
    if pq:
        try:
            num = float(pq)
            unit = clean(pu).lower() or 'g'
            if unit in ('l', 'liter'):
                unit = 'l'
            return f'{num:g} {unit}', None, 'mængde fra product_quantity'
        except ValueError:
            pass
    return '', None, 'mangler mængde'


COUNTRY = {'denmark': 'Danmark', 'germany': 'Tyskland', 'sweden': 'Sverige', 'norway': 'Norge', 'netherlands': 'Holland',
           'italy': 'Italien', 'spain': 'Spanien', 'france': 'Frankrig', 'poland': 'Polen', 'ireland': 'Irland',
           'united-kingdom': 'England', 'belgium': 'Belgien', 'portugal': 'Portugal', 'greece': 'Grækenland',
           'austria': 'Østrig', 'switzerland': 'Schweiz', 'finland': 'Finland', 'turkey': 'Tyrkiet', 'thailand': 'Thailand',
           'morocco': 'Marokko', 'china': 'Kina', 'india': 'Indien', 'brazil': 'Brasilien', 'chile': 'Chile',
           'argentina': 'Argentina', 'new-zealand': 'New Zealand', 'south-africa': 'Sydafrika', 'usa': 'USA',
           'united-states': 'USA', 'peru': 'Peru', 'ecuador': 'Ecuador', 'israel': 'Israel', 'egypt': 'Egypten',
           'lithuania': 'Litauen', 'latvia': 'Letland', 'estonia': 'Estland', 'czech-republic': 'Tjekkiet',
           'hungary': 'Ungarn', 'romania': 'Rumænien', 'bulgaria': 'Bulgarien', 'croatia': 'Kroatien'}

# nøgleord -> Bilkas 9 kategorier (første træffer, rækkefølgen er prioriteten)
CAT_RULES = [
    ('Frost', r'\b(frost|dybfrost|frossen|frosne|frozen|tiefk[uü]hl)\b'),
    ('Drikkevarer', r'\b(sodavand|cola|juice|saft|vand|mineralvand|øl|pilsner|lager|vin|rødvin|hvidvin|rosé|cider|energidrik|læskedrik|iste|ice tea|smoothie|drik|kaffe|kaffekapsler|te|kakaodrik|tonic|sirup|spiritus|likør|whisky|vodka|gin|rom)\b'),
    ('Kød & fisk', r'\b(hakket|oksekød|svinekød|kylling|kalkun|and|lam|kalv|bøf|steak|filet|medister|bacon|pølse|frikadelle|laks|torsk|rejer|fisk|tun|sild|makrel|skinke|flæsk|kotelet|rib|gulerod?)\b'),
    ('Mejeri & køl', r'\b(mælk|ost|yoghurt|yogurt|skyr|smør|fløde|creme fraiche|æg|kvark|kefir|tykmælk|cheese|pålæg|salami|leverpostej|hummus|margarine|piskefløde|mozzarella|feta|cottage)\b'),
    ('Brød & kager', r'\b(brød|rugbrød|boller|bolle|kage|kager|tærte|croissant|wienerbrød|toast|knækbrød|kiks|rundstykke|muffin|brownie|tortilla|wrap)\b'),
    ('Slik & snacks', r'\b(slik|chokolade|chips|snack|vingummi|lakrids|bolsjer|popcorn|nødder|peanuts|bar|riskiks|marcipan|karamel|vafler|is|flødeboller|dragé|tyggegummi|pastiller)\b'),
    ('Frugt & grønt', r'\b(frugt|grønt|salat|tomat|agurk|banan|æble|æbler|kartoffel|løg|gulerod|gulerødder|peberfrugt|broccoli|svampe|avocado|citron|appelsin|jordbær|blåbær|bær)\b'),
    ('Mad fra hele verden', r'\b(sushi|nudler|noodle|ramen|curry|taco|tortilla|salsa|kokosmælk|soja|wok|thai|asiatisk|indisk|mexicansk|pasta|ris)\b'),
]

KW_MEAT = [('kylling', r'kylling|chicken|hähnchen'), ('kalkun', r'kalkun|turkey'), ('okse', r'okse|rind|beef|nötfärs|bøf'),
           ('gris', r'\bgris|svin|flæsk|medister|bacon|skinke|schwein|pork'), ('kalv', r'\bkalv'), ('lam', r'\blam\b|lamm'),
           ('and', r'\band\b|ente'), ('laks', r'laks'), ('tun', r'\btun'), ('rejer', r'rejer|reje|shrimp'), ('sild', r'\bsild'),
           ('fisk', r'fisk|torsk|makrel|sej|rødspætte|fish')]

# negative udtryk der ophæver et label ("ikke økologisk", "uden gluten" er derimod positivt for glutenfri)
FLAG_NEG = re.compile(r'\b(ikke|non|not|ohne)\s+(økologisk|øko|organic|bio)\b', re.I)


# ---------------------------------------------------------------- vokabular fra de eksisterende ark
def load_vocab():
    """Product Type (lille bogstav) -> (Category, plural, antal) fra Bilka m.fl. — arkenes egen standard."""
    files = [BILKA] + [os.path.join(HERE, n) for n in ('spar.xlsx', 'rema1000_version 2.xlsx', 'wolt.xlsx')]
    types = collections.defaultdict(lambda: [collections.Counter(), collections.Counter(), collections.Counter(), ''])
    for f in files:
        if not os.path.exists(f):
            continue
        t = os.path.join(tempfile.gettempdir(), 'vocab_' + os.path.basename(f))
        try:
            shutil.copy2(f, t)
            wb = openpyxl.load_workbook(t, read_only=True)
        except Exception as e:  # noqa: BLE001
            print('  vokabular: springer', os.path.basename(f), 'over:', e)
            continue
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = [clean(h) for h in next(it)]
        alias = {'productType': 'Product Type', 'category': 'Category', 'Product type plural': 'Product type plural'}
        hdr = [alias.get(h, h) for h in hdr]
        try:
            ti, ci = hdr.index('Product Type'), hdr.index('Category')
        except ValueError:
            continue
        pi = hdr.index('Product type plural') if 'Product type plural' in hdr else None
        for r in it:
            ty = clean(r[ti]) if ti < len(r) else ''
            if not ty or len(ty) < 3:
                continue
            e = types[ty.lower()]
            e[2][ty] += 1
            if ci < len(r) and r[ci] and f == BILKA:  # kun Bilkas 9 kategorier er standarden (de andre ark har egne)
                e[0][clean(r[ci])] += 1
            if pi is not None and pi < len(r):
                e[1][clean(r[pi])] += 1
        wb.close()
    vocab = {}
    for k, (cats, plurals, forms, _) in types.items():
        vocab[k] = dict(form=forms.most_common(1)[0][0], cat=(cats.most_common(1)[0][0] if cats else ''),
                        plural=plurals.most_common(1)[0][0] if plurals else '', n=sum(forms.values()))
    return vocab


# ---------------------------------------------------------------- afledning af felterne
class Row(dict):
    pass


_TOK = re.compile(r'[a-zæøåäöüéèàêçñ]+')
_IDX = {}


def _index(vocab):
    """ordsekvens (små bogstaver, uden tegn) -> vokabularnøgle; bygges én gang pr. vokabular"""
    key = id(vocab)
    if key not in _IDX:
        idx = {}
        for k, v in vocab.items():
            if len(k) < 4 or v['n'] < 2:
                continue
            toks = ' '.join(_TOK.findall(k))
            if toks and (toks not in idx or len(k) > len(idx[toks])):
                idx[toks] = k
        _IDX[key] = (idx, max((len(t.split()) for t in idx), default=1))
    return _IDX[key]


def find_type(name, vocab):
    """Længste kendte produkttype der står som helt ord(e) i navnet -> (nøgle, None) eller None."""
    idx, maxn = _index(vocab)
    toks = _TOK.findall(name.lower())
    best = None
    for i in range(len(toks)):
        for n in range(1, min(maxn, len(toks) - i) + 1):
            k = idx.get(' '.join(toks[i:i + n]))
            if k and (best is None or len(k) > len(best[0])):
                best = (k, None)
    return best


def derive(p, store, cfg, vocab):
    """-> (row dict med Bilkas kolonnenavne, liste af (felt, problem, handling))"""
    qa = []
    code = clean(p.get('code'))
    name_raw = clean(p.get('product_name'))
    name_da = clean(p.get('product_name_da'))
    name = name_da or name_raw
    generic = clean(p.get('generic_name_da') or p.get('generic_name'))
    cats_txt = clean(p.get('categories'))
    cats_tags = p.get('categories_tags') or []
    labels = p.get('labels_tags') or []
    allerg = p.get('allergens_tags') or []
    ana = p.get('ingredients_analysis_tags') or []
    ingr = clean(p.get('ingredients_text_da') or p.get('ingredients_text')).lower()
    nut = p.get('nutriments') or {}
    hay = ' '.join([name, generic, cats_txt]).lower()
    r = Row()

    # ---- brand
    toks = [clean(t) for t in re.split(r',', p.get('brands') or '') if clean(t)]
    canon = canon_map(cfg)
    own = [canon[slug(t)] for t in toks if slug(t) in canon]
    own = list(dict.fromkeys(own))  # samme mærke i to stavemåder -> ét
    specific = [t for t in own if slug(t) not in CHAIN_ONLY]
    brand = specific[0] if specific else (own[0] if own else '')
    others = [t for t in toks if slug(t) not in canon and slug(t) not in DROP_BRAND_TOKENS]
    if len(specific) > 1:
        qa.append(('Brand', f'flere egne mærker: {", ".join(specific)}', 'første valgt, resten i Subbrand'))
    r['Brand'] = brand
    r['Subbrand'] = ', '.join(specific[1:]) if len(specific) > 1 else ''
    r['Manufacturer'] = ', '.join(others)
    if others:
        qa.append(('Brand', f'andre mærker på varen: {", ".join(others)}', 'lagt i Manufacturer, ikke i Brand'))

    # ---- navn -> produkttype / variation
    base = name
    for t in toks:
        base = re.sub(r'(?i)^\s*' + re.escape(t) + r'\b[\s:,-]*', '', base)
    base = re.sub(r'(?i)\b\d+([.,]\d+)?\s*(kg|g|ml|cl|l|liter|stk|pcs)\b', '', base)
    pct = re.findall(r'(\d+(?:[.,]\d+)?(?:\s*[-–]\s*\d+(?:[.,]\d+)?)?)\s*%', base)
    base_nopct = re.sub(r'\s*\d+(?:[.,]\d+)?(?:\s*[-–]\s*\d+(?:[.,]\d+)?)?\s*%\s*(fedt|alk\.?|vol\.?)?', ' ', base, flags=re.I)

    flags = {}
    wl = ' '.join(x for x in (name_raw, name_da, generic) if x).lower()  # begge navne: den danske mangler ofte ordet

    def has(rx, text=wl):
        return re.search(rx, text, re.I) is not None

    # økologi: label + navn (krydstjek)
    lab_org = any(re.search(r'organic|øko|oko-|ø-mærk|eu-organic|bio$', t) and 'non-eu' not in t and 'agriculture' not in t
                  for t in labels)
    name_org = has(r'\b(økologisk|økologiske|øko|organic|bio)\b') and not FLAG_NEG.search(wl)
    if lab_org and name_org:
        flags['_is_organic'] = 'økologisk'
    elif lab_org:
        flags['_is_organic'] = 'økologisk'
        qa.append(('_is_organic', 'økologisk kun ifølge label (ikke i navnet)', 'sat til økologisk'))
    elif name_org:
        flags['_is_organic'] = 'økologisk'
        qa.append(('_is_organic', 'økologisk kun ifølge navnet (ingen label i OFF)', 'sat til økologisk'))
    if FLAG_NEG.search(wl) and lab_org:
        qa.append(('_is_organic', 'label siger økologisk, navnet siger "ikke økologisk"', 'ikke sat — tjek manuelt'))
        flags.pop('_is_organic', None)

    # glutenfri
    gf = any(t in ('en:gluten-free', 'da:glutenfri') or t.endswith('gluten-free') for t in labels) or has(r'glutenfri|uden gluten|gluten-free|glutenfrei')
    if gf and 'en:gluten' in allerg:
        qa.append(('_is_glutenfree', 'glutenfri, men allergener indeholder gluten', 'ikke sat — tjek manuelt'))
        gf = False
    if gf:
        flags['_is_glutenfree'] = 'glutenfri'
    # laktosefri
    lf = any('lactose' in t and ('no' in t or 'free' in t) for t in labels) or has(r'laktosefri|lactofree|lactose.?free|uden laktose|laktosefri')
    if lf:
        flags['_is_lactose_free'] = 'laktosefri'
    # vegansk / plantebaseret
    vlab = any(t in ('en:vegan', 'en:european-vegetarian-union-vegan', 'en:the-vegan-society') for t in labels)
    if vlab or has(r'\bvegansk|\bvegan\b|plantebaseret'):
        if 'en:milk' in allerg or 'en:eggs' in allerg:
            qa.append(('_is_vegan', 'vegansk i navn/label, men allergener har mælk/æg', 'ikke sat — tjek manuelt'))
        else:
            flags['_is_vegan'] = 'plantebaseret' if has(r'plantebaseret') and not vlab else 'vegansk'
    elif has(r'vegetar'):
        flags['_is_vegan'] = 'vegetarisk'
    # sukker
    sug = nut.get('sugars_100g')
    claim_sf = has(r'sukkerfri|zero sugar|0 ?% sukker|sugar.?free|no sugar|uden sukker\b') or any(t in ('en:no-sugar', 'en:no-sugars') for t in labels)
    claim_nas = has(r'uden tilsat sukker|ikke tilsat sukker|no added sugar|ingen tilsat sukker') or 'en:no-added-sugar' in labels
    if claim_sf:
        if isinstance(sug, (int, float)) and sug > 2.5:
            qa.append(('_is_sugar_free', f'påstår sukkerfri, men sukker pr. 100 g er {sug}', 'ikke sat — tjek manuelt'))
        else:
            flags['_is_sugar_free'] = 'sukkerfri'
    elif claim_nas:
        flags['_is_sugar_free'] = 'uden tilsat sukker'
    if re.search(r'sødestof|acesulfam|aspartam|sucralose|steviol|sweetener|süßungsmittel|cyklamat|saccharin|erythritol|xylitol', ingr):
        flags['_is_sweeteners'] = 'sødestoffer'
    # fuldkorn / dyrevelfærd / frost / koffeinfri
    if has(r'fuldkorn|wholegrain|whole grain'):
        flags['_is_whole_grain'] = 'fuldkorn'
    welfare = [t for t in labels if 'dyrevelf' in t or 'animal-welfare' in t]
    if has(r'skrabeæg|skrabe'):
        flags['_is_animal_welfare'] = 'skrabeæg'
    elif has(r'fritgående|frilands'):
        flags['_is_animal_welfare'] = 'fritgående'
    elif welfare:
        n = re.search(r'(\d)-?hjerte', ' '.join(welfare))
        flags['_is_animal_welfare'] = {'1': '1 hjerte', '2': '2 hjerter', '3': '3 hjerter'}.get(n.group(1), 'bedre dyrevelfærd') if n else 'bedre dyrevelfærd'
    frozen = any('frozen' in t or 'dybfrost' in t or 'frost' in t for t in cats_tags + labels) or has(r'\b(frost|dybfrost|frossen|frosne|frozen)\b')
    if frozen:
        flags['_is_frozen'] = 'frozen'
    if has(r'koffeinfri|decaf|entkoffeiniert'):
        flags['_is_decaf'] = 'koffeinfri'
    # alkohol
    alc_val = nut.get('alcohol_100g')
    alc_tags = [t for t in cats_tags if re.search(r'beer|wine|spirits|alcohol|øl|vin\b|cider|likør', t)
                and not re.search(r'non-alco|alcohol-free|alkoholfri|without-alcohol|soft-drink|sodavand|0-0', t)]
    is_alc_cat = bool(alc_tags) or has(r'\b(øl|pilsner|vin|rødvin|hvidvin|cider|whisky|vodka|gin|rom|likør|lager)\b')
    if any(re.search(r'non-alco|alcohol-free|alkoholfri|without-alcohol', t) for t in cats_tags + labels):
        is_alc_cat = False
    if has(r'alkoholfri|0[.,]0 ?%|alcohol.?free|uden alkohol|alkoholfrei'):
        flags['_is_alcohol_free'] = 'alkoholfri'
    elif is_alc_cat or (isinstance(alc_val, (int, float)) and alc_val > 0.5):
        flags['_is_alcohol'] = 'indeholder alkohol'
        a = None
        if isinstance(alc_val, (int, float)) and alc_val > 0.5:
            a = f'{alc_val:g}%'
        elif pct:
            a = pct[0].replace(',', '.').replace(' ', '') + '%'
        if a:
            flags['alk_pct'] = a
    # kød / fisk-art
    meatish = any(re.search(r'meat|fish|poultry|sausage|kød|fisk|pålæg|bacon', t) for t in cats_tags) or has(r'kød|fisk|kylling|bøf|pølse|medister|bacon|skinke|laks|rejer|filet')
    if meatish and not flags.get('_is_vegan') in ('vegansk', 'plantebaseret'):
        kinds = [k for k, rx in KW_MEAT if re.search(rx, wl)]
        if kinds:
            flags['_is_meat'] = ' / '.join(kinds[:2])
    # tilstand
    for w, v in (('røget', 'røget'), ('kogt', 'kogt'), ('stegt', 'stegt'), ('tørret', 'tørret'), ('syltet', 'syltet'),
                 ('ristet', 'ristet'), ('bagt', 'bagt'), ('grillet', 'grillet')):
        if re.search(r'\b' + w + r'(e|de)?\b', wl):
            flags['_is_cooked'] = v
            break
    # fedt-% (kun fødevarer hvor % er fedt: mejeri, kød, fedtstoffer)
    fat_ok = has(r'mælk|fløde|creme|yoghurt|skyr|smør|margarine|ost\b|hakket|kød|oksekød|svinekød|fars|medister|pålæg|kvark|tykmælk')
    if pct and fat_ok and not flags.get('alk_pct'):
        v = pct[0].replace(',', '.').replace(' ', '')
        flags['_is_fat'] = v + '%'
        fat100 = nut.get('fat_100g')
        try:
            lo = float(re.split(r'[-–]', v)[0])
            if isinstance(fat100, (int, float)) and abs(lo - fat100) > 4 and not re.search(r'[-–]', v):
                qa.append(('_is_fat', f'navn siger {v}%, næringsindhold fedt er {fat100:g} g/100 g', 'navnet bevaret — tjek'))
        except ValueError:
            pass

    # produkttype: først Bilkas vokabular, ellers OFF's mest specifikke kategori, ellers navnet
    ptype, conf = '', ''
    cand = find_type(base_nopct, vocab) or find_type(generic, vocab)
    if cand:
        k, m = cand
        ptype = vocab[k]['form']
        conf = 'vokabular'
        var_txt = re.sub(r'(?i)(?<![a-zæøåäöü])' + re.escape(k) + r'(?![a-zæøåäöü])', ' ', base_nopct, count=1)
    else:
        var_txt = base_nopct
        danish_cats = [c for c in cats_txt.split(',') if c.strip() and not re.match(r'\s*[a-z]{2}:', c)]
        if danish_cats:
            ptype = clean(danish_cats[-1])
            conf = 'OFF-kategori'
        elif clean(base_nopct):
            ptype = clean(base_nopct)
            conf = 'navn'
        else:
            conf = 'ingen'
        qa.append(('Product Type', f'ikke i vokabularet ({conf}): "{ptype}"', 'gennemgå'))
    # ord der er flag/mængde fjernes fra variationen
    junk = r'\b(økologisk|økologiske|øko|organic|bio|glutenfri|laktosefri|vegansk|vegan|sukkerfri|frost|frossen|frosne|dybfrost|alkoholfri|koffeinfri|fuldkorn|skrabeæg|røget|kogt|stegt|tørret)\b'
    var = re.sub(junk, ' ', var_txt, flags=re.I)
    var = re.sub(r'[()\[\]]', ' ', var)
    var = re.sub(r'\s+', ' ', var).strip(' -,.:;/')
    # ord fra produkttypen/brandet må ikke stå to gange
    for w in {w for w in re.findall(r'\w+', ptype.lower() + ' ' + brand.lower()) if len(w) > 2}:
        var = re.sub(r'(?i)\b' + re.escape(w) + r'\b', ' ', var)
    var = re.sub(r'\s+', ' ', var).strip(' -,.:;/')
    r['Product Type'] = ptype[:1].upper() + ptype[1:] if ptype else ''
    r['Variation'] = lower1(var)
    r['_typekilde'] = conf  # vokabular / OFF-kategori / navn / ingen (kun til statistik og gennemgang)

    # kategori
    cat = vocab.get(ptype.lower(), {}).get('cat', '')
    if not cat:
        for c, rx in CAT_RULES:
            if re.search(rx, hay, re.I):
                cat = c
                break
    if frozen:
        cat = 'Frost'
    if not cat:
        cat = 'Kolonial'
        qa.append(('Category', 'kunne ikke afgøres — sat til Kolonial', 'gennemgå'))
    r['Category'] = cat

    # mængde
    qty, pack, qnote = norm_qty(p.get('quantity'), p.get('product_quantity'), p.get('product_quantity_unit'))
    r['Quantity'] = qty
    r['Pack Count'] = pack
    if qnote:
        qa.append(('Quantity', qnote, '' if qty else 'tom'))
    # oprindelse
    orig = [t.split(':')[-1] for t in (p.get('origins_tags') or [])]
    names = []
    for o in orig:
        c = COUNTRY.get(o)
        if c and c not in names:
            names.append(c)
    if not names and has(r'^dansk\b|\bdansk\b|made in denmark|danish') or any('made-in-denmark' in t for t in labels):
        names = ['Danmark']
    r['_is_country_of_origen'] = ' / '.join(names[:3])

    # felter fra flags
    for k in ('_is_organic', '_is_glutenfree', '_is_lactose_free', '_is_vegan', '_is_sugar_free', '_is_sweeteners', '_is_meat',
              '_is_whole_grain', '_is_animal_welfare', '_is_frozen', '_is_cooked', '_is_decaf', '_is_alcohol_free', '_is_fat'):
        r[k] = flags.get(k, '')
    r['_is_alcohol#1'] = flags.get('_is_alcohol', '')
    r['_is_alcohol#2'] = flags.get('alk_pct', '')
    r['_is_raw'] = ''
    r['_is_allergy'] = ''

    # originalkolonnerne = rå OFF (aldrig rettet)
    r['EAN'] = code
    r['Original Title'] = name_raw
    r['Product Name'] = name_da or name_raw
    r['Subtitle'] = generic
    r['Source URL'] = f'https://world.openfoodfacts.org/product/{code}'
    r['Image File'] = clean(p.get('image_front_url') or p.get('image_url'))
    r['Parse Status'] = ''
    r['Servings'] = clean(p.get('serving_size'))
    r['Clean Subtitle'] = ' | '.join(x for x in (brand, qty) if x)

    # dansk dokumentation: ellers kan "solgt i Danmark" være en flueben-fejl i en flersproget OFF-post
    dk_stores = {'lidl', 'netto', 'rema-1000', 'bilka', 'foetex', 'fotex', 'spar', 'meny', 'coop-365', '365discount',
                 'superbrugsen', 'kvickly', 'brugsen', 'nemlig-com', 'nemlig', 'wolt-marked', 'aldi', 'irma', 'salling',
                 'dagli-brugsen', 'min-kobmand', 'lokalbrugsen', 'coop', 'fakta', 'abc-lavpris'}
    stores_s = {slug(s) for s in (p.get('stores_tags') or [])} | {slug(s) for s in clean(p.get('stores')).split(',')}
    dk_ev = []
    if name_da:
        dk_ev.append('navn_da')
    if p.get('lang') == 'da':
        dk_ev.append('lang=da')
    if p.get('ingredients_text_da'):
        dk_ev.append('ingredienser_da')
    if re.match(r'57[0-9]', code):
        dk_ev.append('dansk stregkode')
    if stores_s & dk_stores:
        dk_ev.append('butik')
    if 'Danmark' in r['_is_country_of_origen']:
        dk_ev.append('oprindelse')
    r['_dk_evidens'] = ', '.join(dk_ev)
    if not dk_ev:
        qa.append(('Land', 'ingen dansk dokumentation i OFF (kun flueben "Danmark")', 'tjek om varen findes i Danmark'))
    # datakvalitet pr. vare
    if not name_da and re.search(r'\b(ohne|mit|und|für|von|Käse|Fleisch|Milch|Brot|Soße|Sauce)\b', name_raw):
        qa.append(('Product Name', f'tysk/ikke-dansk navn "{name_raw}"', 'bevaret som i OFF — oversæt manuelt'))
    if not cats_txt:
        qa.append(('Category', 'OFF har ingen kategorier', 'afledt af navn/vokabular'))
    return r, qa


# ---------------------------------------------------------------- læsning af råfiler
def load_raw(cfg):
    prods = {}
    pats = ['brand_*.json', 'store_*.json']
    for pat in pats:
        for f in glob.glob(os.path.join(RAW, pat)):
            try:
                d = json.load(open(f, encoding='utf-8'))
                if isinstance(d, dict):  # .partial.json fra en afbrudt hentning: {"page": n, "items": [...]}
                    d = d.get('items', [])
                for p in d:
                    prods[p['code']] = p
            except Exception as e:  # noqa: BLE001
                print('  kunne ikke læse', f, e)
    return prods


def canon_map(cfg):
    """slug -> kanonisk mærkenavn (første stavemåde i listen vinder; alias lægges til)"""
    m = {}
    for b in cfg['brands']:
        m.setdefault(slug(b), b)
    for a, c in BRAND_ALIAS.items():
        if slug(c) in m:
            m.setdefault(a, m[slug(c)])
    return m


def selected(prods, cfg):
    own = set(canon_map(cfg))
    keep, drop = [], collections.Counter()
    for code, p in prods.items():
        tags = {slug(t) for t in (p.get('brands_tags') or [])} | {slug(t) for t in (p.get('brands') or '').split(',')}
        stores = {slug(s) for s in (p.get('stores_tags') or [])}
        if 'en:denmark' not in (p.get('countries_tags') or []):
            if tags & own:
                drop['ikke solgt i Danmark (OFF countries)'] += 1
            continue
        if not tags & own:
            continue
        # kædenavnet alene (fx Lidl) tæller kun, hvis varen faktisk er solgt i den kæde
        specific = {t for t in tags & own if t not in CHAIN_ONLY}
        if not specific and not (stores & {slug(s) for s in cfg['store_tags']}):
            drop['kun kædenavn som brand, men ikke solgt i kæden'] += 1
            continue
        if not clean(p.get('product_name')) and not clean(p.get('product_name_da')):
            drop['mangler produktnavn i OFF'] += 1
            continue
        keep.append(p)
    return keep, drop


# ---------------------------------------------------------------- skrivning
def bilka_template():
    wb = openpyxl.load_workbook(BILKA)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    a2, b2 = ws['A2'].value, ws['B2'].value
    b2 = b2.text if isinstance(b2, ArrayFormula) else b2
    wb.close()
    return hdr, a2, b2


def write_sheet(path, hdr, a2, b2, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(hdr)
    for rn, r in enumerate(rows, 2):
        out = []
        seen_alc = 0
        for i, h in enumerate(hdr):
            if h == 'Product title singular':
                out.append(row_formula(a2, rn))
            elif h == 'Product title plural':
                out.append(ArrayFormula(f'B{rn}', row_formula(b2, rn)))
            elif h == '_is_alcohol':
                seen_alc += 1
                out.append(r.get(f'_is_alcohol#{seen_alc}', ''))
            elif h == 'Product type plural':
                out.append(r.get('Product type plural', ''))
            else:
                out.append(r.get(h, ''))
        ws.append([None if v == '' else v for v in out])
    ei = hdr.index('EAN') + 1
    for rr in range(2, len(rows) + 2):
        ws.cell(rr, ei).number_format = '0'
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{L(len(hdr))}{len(rows) + 1}'
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 40
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


FIX_TXT = os.path.join(HERE, 'off_rettelser.txt')
FIX_KEYS = {  # kort nøgle i off_rettelser.txt -> kolonne
    'org': '_is_organic', 'gf': '_is_glutenfree', 'lf': '_is_lactose_free', 'veg': '_is_vegan', 'sf': '_is_sugar_free',
    'sw': '_is_sweeteners', 'meat': '_is_meat', 'wg': '_is_whole_grain', 'aw': '_is_animal_welfare', 'frozen': '_is_frozen',
    'cook': '_is_cooked', 'dec': '_is_decaf', 'af': '_is_alcohol_free', 'fat': '_is_fat', 'alc': '_is_alcohol#1',
    'alcp': '_is_alcohol#2', 'country': '_is_country_of_origen', 'qty': 'Quantity', 'pack': 'Pack Count', 'brand': 'Brand',
    'sub': 'Subbrand', 'kw1': 'Keyword 1', 'kw2': 'Keyword 2', 'kw3': 'Keyword 3', 'kw4': 'Keyword 4', 'kw5': 'Keyword 5',
    'pl': 'Product type plural', 'raw': '_is_raw',
}


def load_fix_txt():
    """off_rettelser.txt: EAN|Product Type|Variation|Category|nøgle=værdi;nøgle=værdi  ('~' = tøm, tom = uændret,
    'drop' = fjern varen, 'qa=tekst' = note i kvalitetskontrollen). Linjer der starter med # er kommentarer."""
    fx = {}
    if not os.path.exists(FIX_TXT):
        return fx
    for ln in open(FIX_TXT, encoding='utf-8'):
        ln = ln.rstrip('\n')
        if not ln.strip() or ln.lstrip().startswith('#'):
            continue
        parts = [x.strip() for x in ln.split('|')]
        parts += [''] * (5 - len(parts))
        ean, ty, va, ca, ex = parts[:5]
        d = {}
        if ty:
            d['Product Type'] = ty
        if va:
            d['Variation'] = va
        if ca:
            d['Category'] = ca
        for tok in [t.strip() for t in ex.split(';') if t.strip()]:
            if tok == 'drop':
                d['_drop'] = True
            elif tok.startswith('qa='):
                d['_qa'] = tok[3:]
            elif '=' in tok:
                k, v = tok.split('=', 1)
                col = FIX_KEYS.get(k.strip())
                if col is None:
                    print('  ukendt nøgle i off_rettelser.txt:', k, 'EAN', ean)
                else:
                    d[col] = v.strip()
        fx[ean] = d
    return fx


def apply_fix(r, fx, store, qa_rows, manual_drop):
    """lægger håndrettelserne ovenpå reglerne og logger før/efter; returnerer False hvis varen skal fjernes"""
    ean = r['EAN']
    if fx.get('_drop'):
        manual_drop.append((store, ean, r['Original Title'], fx.get('_qa', '')))
        return False
    for col, v in fx.items():
        if col in ('_drop', '_qa'):
            continue
        new = '' if v == '~' else v
        if col == 'Pack Count' and new != '':
            new = int(new)
        old = r.get(col, '')
        if col in ('Product Type', 'Variation', 'Category') and old != new:
            qa_rows.append((store, ean, r['Original Title'], col, f'regel gav "{old}"', f'rettet manuelt til "{new}"'))
        r[col] = new
    if fx.get('_qa'):
        qa_rows.append((store, ean, r['Original Title'], 'Gennemgang', fx['_qa'], 'noteret'))
    # konsistens: Frost = kategorien Frost (brugerens regel)
    if r.get('Category') == 'Frost':
        r['_is_frozen'] = 'frozen'
    elif r.get('_is_frozen') == 'frozen' and r.get('Category') != 'Frost':
        qa_rows.append((store, ean, r['Original Title'], '_is_frozen', 'frost-flag men kategori ikke Frost', 'flag fjernet'))
        r['_is_frozen'] = ''
    return True


def main(argv):
    names = argv or list(STORES)
    fixes = json.load(open(FIXES, encoding='utf-8')) if os.path.exists(FIXES) else {}
    for k, v in load_fix_txt().items():
        fixes.setdefault(k, {}).update(v)
    manual_drop = []
    used_fix = set()
    print('henter vokabular fra Bilka/SPAR/REMA/Wolt ...')
    vocab = load_vocab()
    hdr, a2, b2 = bilka_template()
    qa_rows, summary = [], []
    for s in names:
        cfg = STORES[s]
        prods = load_raw(cfg)
        keep, drop = selected(prods, cfg)
        rows = []
        for p in keep:
            r, qa = derive(p, s, cfg, vocab)
            used_fix.add(r['EAN'])
            fx = dict(fixes.get(r['EAN'], {}))
            # flertal fra vokabularet, medmindre håndrettelsen angiver det
            if 'Product type plural' not in fx:
                r['Product type plural'] = vocab.get((r.get('Product Type') or '').lower(), {}).get('plural', '')
            if not apply_fix(r, fx, s, qa_rows, manual_drop):
                continue
            rows.append(r)
            for field, problem, action in qa:
                qa_rows.append((s, r['EAN'], r['Original Title'], field, problem, action))
        rows.sort(key=lambda x: (x.get('Brand', ''), x.get('Category', ''), x.get('Product Type', ''), x.get('Variation', '')))
        out = os.path.join(HERE, cfg['file'])
        if os.path.exists(out):
            os.makedirs(os.path.join(ROOT, 'Backup'), exist_ok=True)
            shutil.copy2(out, os.path.join(ROOT, 'Backup', f'{s}_{datetime.datetime.now():%Y-%m-%d_%H%M}_foer_off-import.xlsx'))
        write_sheet(out, hdr, a2, b2, rows)
        bc = collections.Counter(r['Brand'] for r in rows)
        summary.append((s, len(prods), len(rows), dict(drop), dict(bc)))
        print(f'{s}: {len(prods)} hentet, {len(rows)} i arket -> {out}')
        for k, v in sorted(bc.items(), key=lambda kv: -kv[1]):
            print('   ', k, v)
        for k, v in drop.items():
            print('   fravalgt:', k, v)
    if len(names) == len(STORES):
        unused = [e for e in fixes if e not in used_fix]
        if unused:
            print('håndrettelser uden vare i arkene (forkert EAN eller fravalgt):', ', '.join(unused))
    # kvalitetskontrol
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Overblik'
    ws.append(['Butik', 'Hentet fra OFF (rå)', 'I arket', 'Fravalgt', 'Varer pr. brand'])
    for s, n, k, d, bc in summary:
        ws.append([s, n, k, '; '.join(f'{a}: {b}' for a, b in d.items()), '; '.join(f'{a}: {b}' for a, b in bc.items())])
    w2 = wb.create_sheet('Mistanke og rettelser')
    w2.append(['Butik', 'EAN', 'OFF-navn', 'Felt', 'Problem', 'Handling'])
    for row in qa_rows:
        w2.append([str(row[0]), str(row[1])] + list(row[2:]))
    w3 = wb.create_sheet('Fravalgt manuelt')
    w3.append(['Butik', 'EAN', 'OFF-navn', 'Begrundelse'])
    for row in manual_drop:
        w3.append([str(row[0]), str(row[1]), row[2], row[3]])
    for w in (ws, w2, w3):
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = 'A2'
    qp = os.path.join(HERE, 'off_kvalitetskontrol.xlsx')
    wb.save(qp)
    print('kvalitetskontrol ->', qp, len(qa_rows), 'punkter')


if __name__ == '__main__':
    main(sys.argv[1:])
