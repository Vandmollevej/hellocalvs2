"""Omlaeg SPAR, Nemlig, Wolt, DRK, Aarstiderne og de tyske ark (dm, EDEKA, REWE) til Bilka-kolonnerne (2026-10-10).
Brug: python omlaeg_til_bilka_kolonner.py [butik ...] [--dry]   (uden butik: alle; --dry skriver kun til TEMP)

- Kolonne 1-43 er praecis som Excelark/bilka.xlsx: samme navne, samme raekkefoelge og Bilkas egne titelformler.
- ORIGINALKOLONNERNE (Original Title, Product Name, Subtitle, Source URL, Image File, Parse Status + butikkens
  oevrige skrabede kolonner som Manufacturer, Servings, Price, Venue, Subcategory) kopieres 1:1 fra skrabningen
  og rettes ALDRIG. Har en tidligere omlaegning (omlaeg_ark.py) overskrevet dem, hentes de tilbage fra det raa ark.
- Danske ark faar Bilka-reglerne (titel_ental_flertal.py + procent_ost_regler.py + reglerne fra 2026-10-10).
  Tyske ark faar kun kolonnerne, tyske ord i stedet for "Yes" og tyske frost-ord i titelformlen.
- Alle ark: decimaler med punktum ("1.5 l", "12.5%") og intet tusindtalspunktum - kun i vores egne kolonner.
- Kan koeres igen paa et allerede omlagt ark (Bilka-navnene laeses foerst).
- Arket gemmes oven i sig selv; foerst laegges en kopi i Excelark/backup.
"""
import collections, datetime, os, re, shutil, sys, tempfile
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import titel_ental_flertal as T  # noqa: E402
import procent_ost_regler as P  # noqa: E402

BILKA = os.path.join(HERE, 'bilka.xlsx')
BACKUP = os.path.join(HERE, 'backup')
P_ = lambda *a: os.path.join(ROOT, *a)
STORES = {  # raw = tidligste raa skrabning; originalkolonnerne hentes derfra
    # SPAR: den gamle Packaging "Frozen" staar ogsaa paa vin, sodavand osv. -> frost foelger kategorien Frost (som Bilka)
    'spar': dict(src=P_('Produkter', 'SPAR', 'spar.xlsx'), raw=P_('Produkter', 'SPAR', 'Snapshots', 'spar.xlsx'), lang='da',
                 frozen_packaging=False),
    'nemlig': dict(src=P_('Produkter', 'Nemlig', 'nemlig.xlsx'), lang='da'),
    'wolt': dict(src=P_('Excelark', 'wolt_ny.xlsx'), raw=P_('Excelark', 'wolt.xlsx'), lang='da'),
    'drk': dict(src=P_('Excelark', 'drk.xlsx'), lang='da'),
    'aarstiderne': dict(src=P_('Excelark', 'aarstiderne.xlsx'), lang='da'),
    'dm': dict(src=P_('Excelark', 'dm_ny.xlsx'), raw=P_('Excelark', 'dm.xlsx'), lang='de'),
    'edeka': dict(src=P_('Excelark', 'edeka_ny.xlsx'), raw=P_('Excelark', 'edeka.xlsx'), lang='de'),
    'rewe': dict(src=P_('Excelark', 'rewe_ny.xlsx'), raw=P_('Excelark', 'rewe.xlsx'), lang='de'),
    # REMA omlaegges af Excelark/omlaeg_rema.py (bygger paa dette script; raa skrabning 'rema1000 - To be compaired.xlsx')
}
HEADER_ALIAS = {'Source url': 'Source URL', 'Image file': 'Image File', 'Parse status': 'Parse Status', 'Original title': 'Original Title',
                'Hellocal_title': 'Product Name'}  # REMA: skraberens sammensatte navn svarer til Bilkas Product Name

ORIG = ['Original Title', 'Product Name', 'Subtitle', 'Source URL', 'Image File', 'Parse Status']
ORIG_EXTRA = ['Manufacturer', 'Servings', 'Price', 'Venue', 'Subcategory', 'Vare', 'Variant', 'quantity', 'Kategori', 'Sort']
SRC = {  # Bilka-kolonne -> kildekolonner (foerste udfyldte vinder); is_*-navnene er REMA's
    'packageSizeText': ['packageSizeText', 'Quantity'], 'brand': ['brand', 'Brand'], 'subbrand': ['subbrand', 'Subbrand'],
    'barcode': ['barcode', 'EAN'], 'packCount': ['packCount', 'Pack Count'],
    'productType': ['productType', 'Product Type', 'Vare'], 'variant': ['variant', 'Variation', 'Variant'],
    'category': ['category', 'Category'], 'packaging': ['packaging', 'Packaging', 'is_Packaging'],
    '_is_animal_welfare': ['_is_animal_welfare', '_is_animal_wellfare', 'is_animal_wellfare'],
    '_is_country_of_origen': ['_is_country_of_origen', 'is_country_of_origin'], '_is_organic': ['_is_organic', 'is_biological'],
    '_is_glutenfree': ['_is_glutenfree', 'is_gluten_free'], '_is_lactose_free': ['_is_lactose_free', 'is_lactose_free'],
    '_is_vegan': ['_is_vegan', 'is_vegan'], '_is_meat': ['_is_meat', 'is_meat'], '_is_allergy': ['_is_allergy', '_is_allergies'],
    '_is_sugar_free': ['_is_sugar_free', 'is_sugar_free'], '_is_whole_grain': ['_is_whole_grain', 'is_whole_grain'],
    '_is_alcohol_free': ['_is_alcohol_free', 'is_alcohol_free'],
    **{f'keyword{n}': [f'keyword{n}', f'Keyword {n}'] for n in range(1, 6)},
}
DROP = {'HelloCal_Title', 'Hellocal_title', 'Hello Cal product title', 'Key', 'Product Title (ny)'}  # titlen er nu formlerne i kolonne A/B
ALC_SRC = {'_is_alcohol', '_is_alcohol_pct'}
YES = {'yes', 'ja', 'true', 'x'}
WORD = {
    'da': {'_is_organic': 'økologisk', '_is_glutenfree': 'glutenfri', '_is_lactose_free': 'laktosefri', '_is_vegan': 'vegansk',
           '_is_sugar_free': 'sukkerfri', '_is_sweeteners': 'sødestoffer', '_is_whole_grain': 'fuldkorn', '_is_alcohol_free': 'alkoholfri',
           '_is_decaf': 'koffeinfri', '_is_frozen': 'frozen'},
    'de': {'_is_organic': 'bio', '_is_glutenfree': 'glutenfrei', '_is_lactose_free': 'laktosefrei', '_is_vegan': 'vegan',
           '_is_sugar_free': 'zuckerfrei', '_is_sweeteners': 'süßungsmittel', '_is_whole_grain': 'vollkorn', '_is_alcohol_free': 'alkoholfrei',
           '_is_decaf': 'entkoffeiniert', '_is_frozen': 'frozen'},
}
ALC_FREE = {'alkoholfri', 'uden alkohol', 'alkoholfrei', 'ohne alkohol'}
CHEESE = re.compile(r'\bost\b|\boste\b|ost$|mozzarella|parmesan|cheddar|gouda|emmentaler|\bfeta\b|pecorino|grana padano|mascarpone|ricotta'
                    r'|\bbrie\b|camembert|gorgonzola|danbo|havarti|maasdam|riberhus|castello|salatost|pizzatopping|pastaost|pizzaost|gruyere|manchego', re.I)
MASS_BLANK = {'kylling', 'is', 'kaffe', 'laks', 'skinke', 'sodavand', 'te', 'vand', 'mælk', 'juice', 'saft', 'olie', 'mel', 'sukker', 'salt',
              'smør', 'fløde', 'yoghurt', 'skyr', 'ris', 'pasta', 'honning', 'sirup', 'ketchup', 'sennep', 'mayonnaise', 'remoulade', 'eddike',
              'kakao', 'chokolade', 'lakrids', 'marcipan', 'nougat', 'slik', 'chips', 'müsli', 'mysli', 'granola', 'havregryn', 'pålæg', 'kød',
              'fisk', 'bouillon', 'fond', 'suppe', 'sovs', 'sauce', 'dressing', 'pesto', 'hummus', 'marmelade', 'syltetøj', 'gær', 'cider',
              'cola', 'smoothie', 'tofu', 'tempeh', 'dej', 'pulver', 'tun', 'torsk', 'spinat', 'kanel', 'peber', 'oregano', 'basilikum'}
FLAG_KW = {'økologisk': ('_is_organic', 'økologisk'), 'øko': ('_is_organic', 'økologisk'), 'glutenfri': ('_is_glutenfree', 'glutenfri'),
           'laktosefri': ('_is_lactose_free', 'laktosefri'), 'vegansk': ('_is_vegan', 'vegansk'), 'vegan': ('_is_vegan', 'vegansk'),
           'sødestoffer': ('_is_sweeteners', 'sødestoffer'), 'med sødestoffer': ('_is_sweeteners', 'sødestoffer'),
           'fuldkorn': ('_is_whole_grain', 'fuldkorn')}
FROST_CATS = {'frost', 'frostvarer', 'dybfrost', 'frys', 'frost & is'}
EGG_SIZE = re.compile(r'^(?:xs|s|m|l|xl)\s*/\s*(?:s|m|l|xl|xxl)$', re.I)
SAME_PL = {'løg', 'brød', 'spyd', 'æg', 'rejer', 'dumplings', 'boller', 'chips', 'nachos', 'pommes frites', 'ærter', 'bønner', 'nødder'}
REDUP = re.compile(r'mogu mogu|yum yum|chop chop|bora bora|can can|tutti frutti', re.I)


def fix_text(s):
    """kolonnenavne der er gemt dobbelt-kodet ('Alternativ vÃ¦gt' -> 'Alternativ vægt')"""
    try:
        return s.encode('latin-1').decode('utf-8') if re.search('[ÃÂ]', s) else s
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s


def col_fill(rows, i):
    return any(i < len(r) and r[i] not in (None, '') for r in rows)


def read(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    it = ws.iter_rows(values_only=True)
    hdr = [fix_text(h.strip()) if isinstance(h, str) else h for h in next(it)]
    hdr = [HEADER_ALIAS.get(h, h) if not (HEADER_ALIAS.get(h) in hdr) else h for h in hdr]
    rows = [list(r) + [None] * (len(hdr) - len(r)) for r in it if any(v not in (None, '') for v in r)]
    return hdr, rows


def bilka_master():
    wb = openpyxl.load_workbook(BILKA)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    a2 = ws['A2'].value
    b2 = ws['B2'].value
    b2 = b2.text if isinstance(b2, ArrayFormula) else b2
    a3 = ws['A3'].value
    assert row_formula(a2, 3) == a3, 'Bilkas formel kan ikke flyttes til en anden raekke'
    # flertal pr. produkttype (Bilka er master): flertallet af raekkerne afgoer, tom = intet flertal
    vote = collections.defaultdict(collections.Counter)
    ti, pi = hdr.index('productType'), hdr.index('Product type plural')
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[ti]:
            vote[str(r[ti]).strip().lower()][(str(r[pi]).strip() if r[pi] else '')] += 1
    plural = {k: c.most_common(1)[0][0] for k, c in vote.items()}
    wb.close()
    return hdr[:43], a2, b2, plural


def row_formula(tpl, row):
    return re.sub(r'(?<![A-Za-z_.])([A-Z]{1,2})2(?![0-9A-Za-z])', lambda m: f'{m.group(1)}{row}', tpl)


def german_formula(f):
    """samme formel, men tyske frost-ord og variant/produkttype beholder stort begyndelsesbogstav (tyske navneord)"""
    f = f.replace('"Frosne"', '"Tiefgekühlte"').replace('"Frosset"', '"Tiefgekühlt"')
    f = f.replace('LOWER(LEFT(_xlpm.vr0,1))&MID(_xlpm.vr0,2,300)', '_xlpm.vr0')
    f = f.replace('LOWER(LEFT(_xlpm.ty,1))&MID(_xlpm.ty,2,200)', '_xlpm.ty')
    return f


# ---------- danske regler ----------
def remove_phrase(text, phrase):
    out = re.sub(r'(?<![\wæøåÆØÅ])' + re.escape(phrase) + r'(?![\wæøåÆØÅ])', ' ', text, flags=re.I)
    return T.clean_phrase(re.sub(r'\s+', ' ', out).strip())


def cap(s):
    return s[0].upper() + s[1:] if s else s


def plural_for(ty, master, fr):
    key = ty.strip().lower()
    if key in master:
        return master[key] or None
    if key in fr:
        return fr[key]
    words = key.split(' ')
    head = words[0]
    for i, w in enumerate(words):
        if i > 0 and w in T.PREP:
            break
        head = w
    if head in MASS_BLANK:
        return None
    pl = T.plural_phrase(ty, fr)
    if pl != ty:
        return pl
    if head in SAME_PL or key in SAME_PL or head.endswith('bær'):
        return ty
    return None


def danish_rules(nr, nix, stats, master, fr, cats, frozen_packaging=True):
    ti, si = nix['productType'], nix['_is_sugar_free']
    vi = [nix['variant']]
    ki = [nix[f'keyword{n}'] for n in range(1, 6)]
    for i in [ti] + vi + ki:  # tal i tekstfelter (fx variant 1664) som tekst
        if isinstance(nr[i], (int, float)) and not isinstance(nr[i], bool):
            nr[i] = str(int(nr[i])) if float(nr[i]).is_integer() else str(nr[i]).replace('.', ',')
    is_cols = [i for h, i in nix.items() if h.startswith('_is_') and i < 43 and h != '_is_alcohol_pct']  # kun Bilkas _is_-kolonner
    for i in is_cols:
        if isinstance(nr[i], str):
            nv = T.lower_phrase(nr[i])  # alt smaat, forkortelser (USA, MSC) beholdes
            if nv != nr[i]:
                stats['is_smaat'] += 1
            nr[i] = nv
    # 'Uden' + 'Tilsat sukker' (fejlagtigt delt) -> 'uden tilsat sukker'
    for n in range(len(ki) - 1):
        x, y = nr[ki[n]], nr[ki[n + 1]]
        if isinstance(x, str) and isinstance(y, str) and x.strip(' ,/').lower() == 'uden' and y.strip().lower().startswith('tilsat'):
            nr[ki[n]], nr[ki[n + 1]] = 'uden ' + y.strip().lower(), None
    for i in vi + ki:
        if isinstance(nr[i], str):
            nr[i] = T.lower_phrase(nr[i])
    # keywords der bare er butikkens kategori-/menunavne fjernes; flag-ord flyttes til deres egen kolonne
    for i in ki:
        v = nr[i].strip().lower() if isinstance(nr[i], str) else None
        if not v:
            continue
        if v in cats:
            nr[i] = None
            stats['kategori_keyword_fjernet'] += 1
        elif v in FLAG_KW:
            h, w = FLAG_KW[v]
            nr[nix[h]] = nr[nix[h]] or w
            nr[i] = None
            stats['flag_keyword_til_kolonne'] += 1
    sizes = {i: nr[i] for i in vi + ki if isinstance(nr[i], str) and EGG_SIZE.match(nr[i].strip())}  # M/L er en aeggestoerrelse, ikke 'med l'
    cfg = dict(type='productType', fat='_is_fat', alcfree='_is_alcohol_free', pack='packaging' if frozen_packaging else None, alc='_is_alcohol')
    T.frida_rules(nr, vi, ki, nix, cfg, stats)
    for i, v in sizes.items():
        nr[i] = v.strip().upper().replace(' ', '')
    # procent og ost
    cat = str(nr[nix['category']] or '')
    pcat = 'Mejeri & køl' if re.search(r'mejeri|køl|\bost', cat, re.I) else 'Drikkevarer' if re.search(r'drik|vin|øl|spiritus', cat, re.I) else cat
    pct_i = nix['_is_alcohol_pct']
    pr = P.process(dict(type=nr[ti], variant=nr[vi[0]], kw=[nr[i] for i in ki], fat=nr[nix['_is_fat']], alc=nr[pct_i],
                        alcflag=bool(nr[pct_i]) or str(nr[nix['_is_alcohol']] or '') == 'indeholder alkohol' or bool(re.search(r'vin|spiritus|øl', cat, re.I)),
                        plural=None, orig=nr[nix['Original Title']], cat=pcat, ean=nr[nix['barcode']]))
    nr[ti], nr[nix['_is_fat']], nr[vi[0]], nr[pct_i] = pr['type'], pr['fat'], pr['variant'], pr['alc']
    for i, v in zip(ki, pr['kw']):
        nr[i] = v
    # sukker: keywords som i Frida, sukkerfri/uden/med tilsat sukker/usoedet -> sukkerkolonnen
    for i in ki:
        v = nr[i]
        if isinstance(v, str) and v.strip().lower() in T.SUGAR_KW:
            nr[i] = T.SUGAR_KW[v.strip().lower()]
            if 'sødestof' in nr[i] and not nr[nix['_is_sweeteners']]:
                nr[nix['_is_sweeteners']] = 'sødestoffer'
    if isinstance(nr[si], str) and nr[si].startswith('ikke tilsat'):
        nr[si] = 'uden tilsat sukker'
    for i in vi + ki:
        v = nr[i]
        if not isinstance(v, str):
            continue
        mo = T.SUGAR_RE.search(v) or re.search(r'\b(med tilsat sukker|usødet|usødede)\b', v, re.I)
        if mo:
            w = mo.group(1).lower()
            val = 'sukkerfri' if w.startswith('sukkerfri') else 'med tilsat sukker' if w == 'med tilsat sukker' else 'uden tilsat sukker'
            if not nr[si]:
                nr[si] = val
            nr[i] = T.clean_phrase(re.sub(re.escape(mo.group(0)), '', v, flags=re.I).strip()) or None
            stats['sukker_til_egen_kolonne'] += 1
    # 'tilsat kulsyre' hedder 'med kulsyre'
    for i in ki:
        if isinstance(nr[i], str) and re.fullmatch(r'(?:tilsat kulsyre|kulsyre tilsat|med tilsat kulsyre|kulsyreholdig)', nr[i].strip(), re.I):
            nr[i] = 'med kulsyre'
            stats['med_kulsyre'] += 1
    # brand/subbrand staar aldrig i produkttypen
    ty = nr[ti]
    for b in (nr[nix['brand']], nr[nix['subbrand']]):
        if isinstance(ty, str) and isinstance(b, str) and len(b.strip()) >= 3 and ty.strip().lower() != b.strip().lower():
            if re.search(r'(?<![\wæøå])' + re.escape(b.strip().lower()) + r'(?![\wæøå])', ty.lower()):
                ty = cap(remove_phrase(ty, b.strip())) or ty
                stats['brand_ud_af_produkttype'] += 1
    nr[ti] = ty
    # ost (selve varen er ost): revet / i blok; originaltitlen afgoer, hvis arket har mistet ordet
    orig = str(nr[nix['Original Title']] or '')
    own = ' '.join(str(nr[i] or '') for i in [ti] + vi)
    if isinstance(nr[ti], str) and CHEESE.search(nr[ti]) and not re.search(r'\b(?:med|og|i|på|uden)\s+\S*ost\b', nr[ti], re.I):
        if re.search(r'\brevet\b|\brevne\b', own + ' ' + orig, re.I):
            for i in [ti] + vi:
                if isinstance(nr[i], str):
                    nr[i] = remove_phrase(remove_phrase(nr[i], 'revet'), 'revne') or None
            nr[ti] = cap(nr[ti] or 'Ost')
            T.put_kw(nr, ki, 'revet')
            stats['ost_revet'] += 1
        elif not re.search(r'skiver|skive|\btern\b|halvfast|smøreost|flødeost|friskost', own + ' ' + orig, re.I):
            hit = bool(re.search(r'\bi (?:stykke|blok)\b', orig, re.I))
            for i in [ti] + vi:
                if isinstance(nr[i], str) and re.search(r'\bi (?:stykke|blok)\b', nr[i], re.I):
                    nr[i] = remove_phrase(remove_phrase(nr[i], 'i stykke'), 'i blok') or None
                    hit = True
            if hit or re.search(r'\bfast\b', nr[ti] or '', re.I):
                nr[ti] = cap(nr[ti] or 'Ost')
                T.put_kw(nr, ki, 'i blok')
                stats['ost_i_blok'] += 1
    # 'ben' hoerer til produkttypen
    v = nr[vi[0]]
    if isinstance(v, str) and isinstance(nr[ti], str):
        if v.strip().lower() == 'ben':
            nr[ti] = nr[ti].rstrip() + 'ben'
            nr[vi[0]] = None
            stats['ben_i_type'] += 1
        else:
            mo = re.search(r'(?<![\wæøå])(med|uden) ben(?![\wæøå])', v, re.I)
            if mo and not re.search(r'\bog ben\b', v):
                nr[ti] = nr[ti].rstrip() + ' ' + mo.group(0).lower()
                nr[vi[0]] = T.clean_phrase(v.replace(mo.group(0), '').strip()) or None
                stats['ben_i_type'] += 1
    # intet ord to gange i produkttypen
    if isinstance(nr[ti], str) and not REDUP.search(nr[ti]):
        nt = T.dedupe(nr[ti], set())
        if nt and nt != nr[ti]:
            nr[ti] = cap(nt)
            stats['dobbeltord_i_type'] += 1
    # dobbeltord paa tvaers af felterne
    seen = set(T.tokens(str(nr[ti] or '')))
    for i in is_cols:
        if isinstance(nr[i], str) and i not in (nix['_is_fat'], nix['_is_meat']) and not re.search(r'\d', nr[i]):
            seen.update(T.tokens(nr[i]))
    for i in vi + ki:
        if isinstance(nr[i], str):
            nv = T.dedupe(nr[i], seen) or None
            if nv != nr[i]:
                stats['dobbeltord_fjernet'] += 1
            nr[i] = nv
    kv = [nr[i] for i in ki if nr[i]]
    for j, i in enumerate(ki):
        nr[i] = kv[j] if j < len(kv) else None
    if nr[ti]:
        ty = str(nr[ti])
        pl = plural_for(ty, master, fr)
        if pl:  # samme begyndelsesbogstav som produkttypen
            pl = ty if pl.lower() == ty.lower() else (pl[0].upper() if ty[0].isupper() else pl[0].lower()) + pl[1:]
        nr[nix['Product type plural']] = pl


NUMERIC_TEXT = {'packageSizeText', 'Clean Subtitle'}


def decimal_points(nr, out_hdr, orig_ix):
    """decimaler med punktum i vores egne kolonner (brugerens regel 2026-10-10); originalkolonnerne roeres ikke"""
    n = 0
    for i, v in enumerate(nr):
        if i < 2 or i in orig_ix or not isinstance(v, str) or not re.search(r'\d[.,]\s*\d', v):
            continue
        h = out_hdr[i]
        s = re.sub(r'(?<![\d.,])(\d)\.(\d{3})(?![\d.,])', r'\1\2', v)  # tusindtalspunktum: 1.080 g -> 1080 g
        s = re.sub(r'(?<=\d),(?=\d)', '.', s)  # 1,5 l -> 1.5 l
        if h in NUMERIC_TEXT:
            s = re.sub(r'(?<=\d),\s+(?=\d)', '.', s)  # oedelagt: 0, 5 l -> 0.5 l
        elif h == 'productType':
            s = re.sub(r'(?<![\d])(\d),\s+(\d{1,2})(?=\s*(?:l|cl|ml|g|kg|%)\b)', r'\1.\2', s)
        if s != v:
            nr[i] = s
            n += 1
    return n


# ---------- omlaegning ----------
def convert(name, cfg, bhdr, a2, b2, master, fr, dry):
    src, lang = cfg['src'], cfg['lang']
    lock = os.path.join(os.path.dirname(src), '~$' + os.path.basename(src))
    if os.path.exists(lock) and not dry:
        print(f'{name}: SPRUNGET OVER - {os.path.basename(src)} er aaben i Excel')
        return None
    mtime = os.path.getmtime(src)
    hdr, rows = read(cfg.get('read', src))  # --fra=<fil>: laes fra en backup, skriv stadig til arket
    raw_hdr, raw_rows = read(cfg['raw']) if cfg.get('raw') else (None, None)
    stats = collections.Counter()
    pos = collections.defaultdict(list)
    for i, h in enumerate(hdr):
        pos[h].append(i)
    # raa raekker til originalkolonnerne (match paa Source URL; ved flere: den med samme titel)
    raw_q = collections.defaultdict(list)
    if raw_rows:
        rix = {h: i for i, h in enumerate(raw_hdr)}
        for r in raw_rows:
            raw_q[r[rix['Source URL']]].append(r)

    def raw_match(url, titles):
        cands = raw_q.get(url)
        if not cands:
            return None
        for k, c in enumerate(cands):
            if c[rix['Original Title']] in titles or c[rix.get('Product Name', rix['Original Title'])] in titles:
                return cands.pop(k)
        return cands.pop(0)
    cats = {str(r[pos[h][0]]).strip().lower() for h in ('Category', 'category', 'Subcategory') if h in pos for r in rows if r[pos[h][0]]}
    cats |= {p for c in list(cats) for p in re.split(r'\s*(?:&|,|/|\bog\b|-)\s*', c) if len(p) > 2}  # 'Vin & spiritus' -> vin, spiritus
    cats |= {'kolonial', 'kiosk', 'køl', 'drikke', 'drikkevarer', 'frost', 'frugt', 'grønt', 'frugt & grønt', 'mejeri', 'pålæg', 'slik', 'brød', 'kød', 'fisk'}
    cats -= set(FLAG_KW) | {'frost'}  # frost -> _is_frozen (frida_rules)
    # ekstra kolonner bagerst: foerst skrabede (originale), saa vores egne
    used = set(DROP) | ALC_SRC | set(bhdr) | {s for v in SRC.values() for s in v}
    orig_extra = [h for h in ORIG_EXTRA if (h in pos and col_fill(rows, pos[h][0])) or (raw_hdr and h in raw_hdr and col_fill(raw_rows, raw_hdr.index(h)))]
    own_extra = [h for h in dict.fromkeys(hdr) if h and h not in used and h not in ORIG_EXTRA and col_fill(rows, pos[h][0])]
    if cfg.get('identical'):  # brugeren 2026-10-10: kolonnerne skal vaere 100% som Bilkas, ingen ekstra bagerst
        orig_extra, own_extra = [], []
    out_hdr = list(bhdr) + orig_extra + own_extra
    nix = {}
    for i, h in enumerate(out_hdr):
        nix.setdefault(h, i)
    nix['_is_alcohol_pct'] = [i for i, h in enumerate(bhdr) if h == '_is_alcohol'][1]
    alc_src = [i for i, h in enumerate(hdr) if h in ALC_SRC]
    orig_ix = {nix[h] for h in ORIG + orig_extra}
    flags = []
    out_rows = []
    for r in rows:
        g = lambda h: r[pos[h][0]] if h in pos else None
        nr = [None] * len(out_hdr)
        for j, h in enumerate(bhdr):
            if h in ('Product title singular', 'Product title plural', '_is_alcohol') or h in ORIG:  # flertal genberegnes af reglerne
                continue
            for s in SRC.get(h, [h]):
                if g(s) not in (None, ''):
                    nr[j] = g(s)
                    break
        for i in alc_src:  # procent -> 2. _is_alcohol, ellers flag -> 1. _is_alcohol
            v = r[i]
            if v in (None, ''):
                continue
            if re.search(r'\d', str(v)):
                nr[nix['_is_alcohol_pct']] = nr[nix['_is_alcohol_pct']] or v
            else:
                nr[nix['_is_alcohol']] = nr[nix['_is_alcohol']] or v
        o = None
        if raw_rows:
            o = raw_match(g('Source URL'), {g('Original Title'), g('Product Name')})
            stats['original_fra_raat_ark' if o else 'original_uden_raat_match'] += 1
        for h in ORIG + orig_extra:  # kolonner det raa ark ikke har, tages fra arket selv
            nr[nix[h]] = o[rix[h]] if o is not None and h in rix else g(h)
            if o is not None and h in rix and (nr[nix[h]] or None) != (g(h) or None):
                stats[f'original_genskabt:{h}'] += 1
        if raw_rows:  # markeringer som omlaeg_ark.py skrev ind i Parse Status (PET_FOOD, BRAND_GUESS)
            old = {t.strip() for t in str(o[rix['Parse Status']] or '').split('|') if t.strip()} if o is not None else set()
            new = [t.strip() for t in str(g('Parse Status') or '').split('|') if t.strip() and t.strip() not in old] if o is not None else []
            flags.append(' | '.join(new) or None)
        for h in own_extra:
            nr[nix[h]] = g(h)
        # strip
        for i, v in enumerate(nr):
            if isinstance(v, str) and i not in orig_ix:
                nr[i] = v.strip() or None
        bc = nr[nix['barcode']]  # stregkode altid tekst (ellers viser Excel 5,70604E+12)
        if isinstance(bc, (int, float)) and not isinstance(bc, bool):
            nr[nix['barcode']] = str(int(bc))
            stats['stregkode_til_tekst'] += 1
        # Yes -> beskrivende ord
        for h, i in nix.items():
            if h.startswith('_is_') and isinstance(nr[i], str) and nr[i].strip().lower() in YES:
                nr[i] = WORD[lang].get(h, h[4:].replace('_', ' '))
                stats['yes_til_ord'] += 1
        if isinstance(nr[nix['_is_alcohol']], str) and nr[nix['_is_alcohol']].strip().lower() in ALC_FREE:
            nr[nix['_is_alcohol_free']] = nr[nix['_is_alcohol_free']] or WORD[lang]['_is_alcohol_free']
            nr[nix['_is_alcohol']] = None
            stats['alkoholfri_flyttet'] += 1
        if o is not None and not nr[nix['category']] and 'Kategori' in rix and o[rix['Kategori']]:  # REMA: butikkens kategori
            nr[nix['category']] = str(o[rix['Kategori']]).strip()
            stats['kategori_fra_raat_ark'] += 1
        trust = cfg.get('frozen_packaging', True)
        rules = cfg.get('rules', True)
        if lang == 'da' and rules:
            danish_rules(nr, nix, stats, master, fr, cats, trust)
        frost_cat = str(nr[nix['category']] or '').strip().lower() in FROST_CATS
        if isinstance(nr[nix['packaging']], str) and nr[nix['packaging']].strip().lower() == 'frozen':  # frost kun i _is_frozen
            if trust or frost_cat:
                nr[nix['_is_frozen']] = 'frozen'
                stats['frost_til_is_frozen'] += 1
            else:
                stats['forkert_frozen_fjernet'] += 1
            nr[nix['packaging']] = None
        if frost_cat and not nr[nix['_is_frozen']]:  # hele kategorien Frost er frost (som Bilka)
            nr[nix['_is_frozen']] = 'frozen'
            stats['frost_fra_kategori'] += 1
        if rules:  # REMA har allerede faaet decimalpunktum
            stats['decimal_punktum'] += decimal_points(nr, out_hdr, orig_ix)
        out_rows.append(nr)
    if any(flags) and not cfg.get('identical'):
        out_hdr.append('Flag')
        for nr, f in zip(out_rows, flags):
            nr.append(f)
    # skriv
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(out_hdr)
    fa, fb = (german_formula(a2), german_formula(b2)) if lang == 'de' else (a2, b2)
    for rn, nr in enumerate(out_rows, 2):
        nr[0] = row_formula(fa, rn)
        nr[1] = ArrayFormula(f'B{rn}', row_formula(fb, rn))
        ws.append(nr)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{L(len(out_hdr))}{len(out_rows) + 1}'
    for c, w in (('A', 40), ('B', 40)):
        ws.column_dimensions[c].width = w
    wb.calculation.fullCalcOnLoad = True
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP'))
    os.close(fd)
    wb.save(tmp)
    if dry:
        print(f'{name}: {len(out_rows)} raekker (proeve) -> {tmp}')
    else:
        if os.path.getmtime(src) != mtime or os.path.exists(lock):
            print(f'{name}: SPRUNGET OVER - {os.path.basename(src)} blev aendret/aabnet imens')
            os.remove(tmp)
            return None
        if 'read' not in cfg:  # ved --fra=<backup> findes backuppen allerede
            os.makedirs(BACKUP, exist_ok=True)
            stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
            shutil.copy2(src, os.path.join(BACKUP, f'{os.path.splitext(os.path.basename(src))[0]}_{stamp}_foer_bilka-kolonner.xlsx'))
        shutil.move(tmp, src)
        print(f'{name}: {len(out_rows)} raekker -> {src}')
    print('   ekstra bagerst:', orig_extra, '+', own_extra + (['Flag'] if any(flags) else []))
    print('   ', dict(stats))
    return tmp if dry else src


def decimal_pass(name, path):
    """kun decimalreglen paa et ark der allerede har Bilka-kolonnerne (koer ikke hele omlaegningen to gange)"""
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock):
        print(f'{name}: SPRUNGET OVER - aaben i Excel')
        return
    mtime = os.path.getmtime(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    orig_ix = {i for i, h in enumerate(hdr) if h in ORIG or h in ORIG_EXTRA}
    n = 0
    for row in ws.iter_rows(min_row=2):
        vals = [c.value for c in row]
        before = list(vals)
        n += decimal_points(vals, hdr, orig_ix)
        for c, a, b in zip(row, before, vals):
            if a != b:
                c.value = b
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP'))
    os.close(fd)
    wb.save(tmp)
    if os.path.getmtime(path) != mtime or os.path.exists(lock):
        print(f'{name}: SPRUNGET OVER - aendret/aabnet imens')
        os.remove(tmp)
        return
    shutil.move(tmp, path)
    print(f'{name}: {n} celler med decimalpunktum -> {path}')


def main(args):
    dry = '--dry' in args
    names = [a for a in args if not a.startswith('--')] or list(STORES)
    if '--kun-decimaler' in args:
        for n in names:
            decimal_pass(n, STORES[n]['src'])
        return
    fra = [a.split('=', 1)[1] for a in args if a.startswith('--fra=')]
    bhdr, a2, b2, master = bilka_master()
    fr = T.load_frida_plurals()
    for n in names:
        cfg = dict(STORES[n], read=os.path.abspath(fra[0])) if fra and len(names) == 1 else STORES[n]
        convert(n, cfg, bhdr, a2, b2, master, fr, dry)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1:])
