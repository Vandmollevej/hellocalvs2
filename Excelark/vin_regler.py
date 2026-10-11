# -*- coding: utf-8 -*-
r"""Anvend vin-resultater på arkene.
Brug:  py -3 apply_wine.py <results.json> [--sheets bilka,spar,...] [--dry]
results.json = liste af {id, sheet, row, url, ean, section, productType, misfiled, brand, subbrand, variation, keywords,
               packaging, country, organic, alcoholFree, alcoholPct, ...}
Protokol: læs aktuelt ark -> match på Source URL (+EAN) -> kun egne celler -> kun hvis rækken ikke er ændret siden udtrækket ->
backup i Backup\ -> skriv. Åbne arbejdsbøger (Excel) skrives celle for celle via COM med kontrol af gammel værdi; lukkede ark
via openpyxl (temp + omdøb, kontrol af mtime).
"""
import collections, datetime, io, json, os, re, shutil, sys, tempfile, time

sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vinlib as V

BASE = r"C:\Users\Peter\Desktop\Hello Cal"
BACKUP = os.path.join(BASE, 'Backup')
SHEETS = {
    'bilka': (r"Excelark\bilka.xlsx", 'da'), 'rema': (r"Produkter\rema1000_version 2.xlsx", 'da'),
    'spar': (r"Produkter\SPAR\spar.xlsx", 'da'), 'nemlig': (r"Produkter\Nemlig\nemlig.xlsx", 'da'),
    'wolt': (r"Excelark\wolt.xlsx", 'da'), 'drk': (r"Excelark\drk.xlsx", 'da'), 'aarstiderne': (r"Excelark\aarstiderne.xlsx", 'da'),
    'dm': (r"Excelark\dm.xlsx", 'de'), 'edeka': (r"Excelark\edeka.xlsx", 'de'), 'rewe': (r"Excelark\rewe.xlsx", 'de'),
}
OWNED = ['Product Type', 'Product type plural', 'Brand', 'Subbrand', 'Variation', 'Keyword 1', 'Keyword 2', 'Keyword 3', 'Keyword 4',
         'Keyword 5', 'Packaging', '_is_country_of_origen', '_is_organic', '_is_alcohol_free', 'ALC_PCT']
WORDS = {'da': dict(organic='økologisk', alcfree='alkoholfri'), 'de': dict(organic='bio', alcfree='alkoholfrei')}


def s(v):
    if v is None: return ''
    if isinstance(v, float) and v == int(v): v = int(v)
    return str(v).strip()


def header_index(hdr):
    idx, alc = {}, []
    for i, h in enumerate(hdr):
        if h == '_is_alcohol': alc.append(i)
        if h is not None: idx.setdefault(h, i)
    if len(alc) > 1: idx['ALC_PCT'] = alc[1]
    return idx


def read_current(path):
    """returnér (header, rows) fra en snapshot-kopi af filen (kopi, så Excel-låse ikke generer)"""
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP')); os.close(fd)
    shutil.copy2(path, tmp)
    try:
        wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        rows = [list(r) for r in it]
        wb.close()
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    return hdr, rows


def target_for(rec, out, cur, lang):
    """kortlæg agentens resultat til ønskede celleværdier ('' = ryd). cur = dict kolonne->nuværende værdi (str)."""
    want = {}
    sec = out['section']
    if sec == 'other':
        if out.get('misfiled') and out['productType']:
            want['Product Type'] = out['productType']
            pl = V.plural_for(out['productType'], lang)
            want['Product type plural'] = out['productType'] if pl is None else pl
        return want
    pt = out['productType'] or cur['Product Type']
    want['Product Type'] = pt
    pl = V.plural_for(pt, lang)
    if pl is not None:
        want['Product type plural'] = pl
    want['Brand'] = out['brand'] or ''
    want['Subbrand'] = out['subbrand'] or ''
    want['Variation'] = out['variation'] or ''
    kws = [k for k in (out['keywords'] or []) if k][:5]
    for i in range(5):
        want[f'Keyword {i+1}'] = kws[i] if i < len(kws) else ''
    if out['packaging'] and not cur['Packaging']:
        want['Packaging'] = out['packaging']
    c = out['country'] or ''
    if c:
        if not cur['_is_country_of_origen'] or V.fold(cur['_is_country_of_origen']) == V.fold(c):
            want['_is_country_of_origen'] = c
    elif cur['_is_country_of_origen'] and lang == 'da' and cur['_is_country_of_origen'][:1].islower():
        want['_is_country_of_origen'] = cur['_is_country_of_origen'][:1].upper() + cur['_is_country_of_origen'][1:]
    if out['organic'] and not cur['_is_organic']:
        want['_is_organic'] = WORDS[lang]['organic']
    if out['alcoholFree'] and not cur['_is_alcohol_free']:
        want['_is_alcohol_free'] = WORDS[lang]['alcfree']
    if out['alcoholPct'] and not cur['ALC_PCT']:
        want['ALC_PCT'] = out['alcoholPct']
    return want


def norm_cmp(a, b):
    return s(a) == s(b)


def plan_sheet(name, results, scope):
    path = os.path.join(BASE, SHEETS[name][0])
    lang = SHEETS[name][1]
    hdr, rows = read_current(path)
    idx = header_index(hdr)
    ucol, ecol = idx['Source URL'], idx['EAN']
    by_url = collections.defaultdict(list)
    for i, r in enumerate(rows):
        by_url[s(r[ucol])].append(i)
    changes, conflicts, unmatched, stats = [], [], [], collections.Counter()
    for out in results:
        rec = scope[out['id']]
        cand = by_url.get(rec['url'], [])
        cand = [i for i in cand if s(rows[i][ecol]) == rec['ean']] or cand
        if len(cand) != 1:
            unmatched.append((out['id'], len(cand))); continue
        i = cand[0]
        r = rows[i]
        cur = {c: s(r[idx[c]]) if c in idx else '' for c in OWNED}
        # er rækken ændret af andre siden udtrækket?
        snap = {'Product Type': rec['pt'], 'Brand': rec['brand'], 'Subbrand': rec['sub'], 'Variation': rec['var']}
        diff = [c for c, v in snap.items() if c in idx and s(r[idx[c]]) != v]
        if diff:
            conflicts.append((out['id'], diff)); continue
        want = target_for(rec, out, cur, lang)
        n = 0
        for c, v in want.items():
            if c not in idx: continue
            if cur[c] != v:
                changes.append(dict(id=out['id'], row=i + 2, col=idx[c], name=c, old=cur[c], new=v, url=rec['url']))
                n += 1
        stats['rows_changed' if n else 'rows_same'] += 1
    return dict(path=path, lang=lang, idx=idx, hdr=hdr, ucol=ucol, changes=changes, conflicts=conflicts, unmatched=unmatched, stats=stats)


def text_like(v):
    """strenge Excel ellers ville tolke som tal/dato/formel (2021, 13.5%, 1/2, =x, +x, -x)"""
    return bool(re.fullmatch(r"[\d.,%\s+\-=/:]+", v)) or bool(re.match(r"[=+\-@]", v))


COM_SHEETS = {'bilka', 'rema'}   # brugerens egne arbejdsark: skrives i Excel (bevarer gemte formelværdier); resten som lukket fil


def excel_running():
    try:
        import subprocess
        out = subprocess.run(['tasklist', '/FI', 'IMAGENAME eq EXCEL.EXE'], capture_output=True, text=True).stdout
        return 'EXCEL.EXE' in out
    except Exception:
        return True


def get_workbook(path):
    import win32com.client
    try:
        return win32com.client.GetObject(path), False
    except Exception:
        pass
    try:
        app = win32com.client.GetActiveObject('Excel.Application')
    except Exception:
        app = win32com.client.Dispatch('Excel.Application')
    app.Visible = True
    app.DisplayAlerts = False
    wb = app.Workbooks.Open(path, UpdateLinks=0)
    app.DisplayAlerts = True
    return wb, True


def apply_com(plan, tries=40):
    wb, opened = get_workbook(plan['path'])
    app = wb.Application
    ws = wb.Worksheets(1)
    old_calc = None
    try:
        old_calc = app.Calculation
        app.Calculation = -4135   # manuel beregning mens vi skriver
    except Exception:
        old_calc = None
    done, skipped = 0, []
    for ch in plan['changes']:
        for attempt in range(tries):
            try:
                if s(ws.Cells(ch['row'], plan['ucol'] + 1).Value) != ch['url']:
                    skipped.append((ch, 'url')); break
                cell = ws.Cells(ch['row'], ch['col'] + 1)
                if s(cell.Value) != ch['old']:
                    skipped.append((ch, 'old=' + s(cell.Value))); break
                if ch['new'] == '':
                    cell.ClearContents()
                else:
                    cell.Value = ("'" + ch['new']) if text_like(ch['new']) else ch['new']
                done += 1
                break
            except Exception as e:
                time.sleep(1.5)
        else:
            skipped.append((ch, 'excel optaget'))
    if old_calc is not None:
        for attempt in range(tries):
            try:
                app.Calculation = old_calc; break
            except Exception:
                time.sleep(1.5)
    for attempt in range(tries):
        try:
            wb.Save(); break
        except Exception:
            time.sleep(2)
    return done, skipped


def apply_closed(plan, name):
    path = plan['path']
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock) and excel_running():
        return None, [('lock', 'åben i Excel')]
    mtime = os.path.getmtime(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    done, skipped = 0, []
    for ch in plan['changes']:
        c = ws.cell(row=ch['row'], column=ch['col'] + 1)
        if s(ws.cell(row=ch['row'], column=plan['ucol'] + 1).value) != ch['url']:
            skipped.append((ch, 'url')); continue
        if s(c.value) != ch['old']:
            skipped.append((ch, 'old=' + s(c.value))); continue
        c.value = ch['new'] if ch['new'] != '' else None
        done += 1
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP')); os.close(fd)
    wb.save(tmp)
    wb.close()
    if os.path.getmtime(path) != mtime or (os.path.exists(lock) and excel_running()):
        os.remove(tmp)
        return -1, [('changed', 'filen er ændret imens')]
    stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
    os.makedirs(BACKUP, exist_ok=True)
    shutil.copy2(path, os.path.join(BACKUP, f'{os.path.splitext(os.path.basename(path))[0]}_{stamp}_foer_vin.xlsx'))
    shutil.move(tmp, path)
    return done, skipped


def main(argv):
    res_path = argv[0]
    dry = '--dry' in argv
    only = None
    if '--sheets' in argv:
        only = argv[argv.index('--sheets') + 1].split(',')
    results = json.load(open(res_path, encoding='utf-8'))
    scope = {r['id']: r for r in json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'scope_rows.json'), encoding='utf-8'))}
    by_sheet = collections.defaultdict(list)
    for r in results:
        by_sheet[r['id'].split(':')[0]].append(r)
    report = {}
    for name, res in by_sheet.items():
        if only and name not in only: continue
        plan = plan_sheet(name, res, scope)
        print(f"{name}: {len(res)} resultater | ændringer {len(plan['changes'])} i {plan['stats']['rows_changed']} rækker | uændret {plan['stats']['rows_same']} | "
              f"konflikter {len(plan['conflicts'])} | ikke fundet {len(plan['unmatched'])}")
        with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), f'changes_{name}.json'), 'w', encoding='utf-8') as f:
            json.dump(dict(changes=plan['changes'], conflicts=plan['conflicts'], unmatched=plan['unmatched']), f, ensure_ascii=False)
        if dry: continue
        if not plan['changes']:
            continue
        for attempt in range(6):
            lock = os.path.join(os.path.dirname(plan['path']), '~$' + os.path.basename(plan['path']))
            use_com = name in COM_SHEETS or (os.path.exists(lock) and excel_running())
            if use_com:
                stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
                os.makedirs(BACKUP, exist_ok=True)
                shutil.copy2(plan['path'], os.path.join(BACKUP, f'{os.path.splitext(os.path.basename(plan["path"]))[0]}_{stamp}_foer_vin.xlsx'))
                done, skipped = apply_com(plan)
                break
            done, skipped = apply_closed(plan, name)
            if done is None or done == -1:
                print(f'   forsøg {attempt+1}: {skipped[0][1]} - planlægger igen om lidt')
                time.sleep(20)
                plan = plan_sheet(name, res, scope)
                continue
            break
        print(f"   skrevet {done}, sprunget over {len(skipped)}")
        for sk in skipped[:10]: print('     ', sk)


if __name__ == '__main__':
    main(sys.argv[1:])
