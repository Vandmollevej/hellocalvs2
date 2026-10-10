"""Bilka/REMA: fjern Hello Cal-titel, indfoer ental/flertal-titelformler som i Frida-ark/frida.xlsx
plus brugerens regler 2026-10-09 (alt smaat, sukkerfri i sukkerkolonne, ingen dobbeltord).
Brug: python titel_ental_flertal.py bilka|rema   (originalerne roeres ikke; skriver <navn>_ny.xlsx)
"""
import re, sys, shutil, os, collections
import openpyxl
from openpyxl.utils import get_column_letter as L

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRIDA = os.path.join(ROOT, 'Frida-ark', 'frida.xlsx')

CFG = {
    'bilka': dict(src=os.path.join(ROOT, 'Excelark', 'bilka.xlsx'), out=os.path.join(ROOT, 'Excelark', 'bilka_ny.xlsx'),
                  title='HelloCal_Title', type='Product Type', var=['Variation'], kw=['Keyword 1', 'Keyword 2', 'Keyword 3', 'Keyword 4', 'Keyword 5'],
                  sugar='_is_sugar_free', vegan='_is_vegan', alc='_is_alcohol', gluten='_is_glutenfree', lactose='_is_lactose_free',
                  fat='_is_fat', alcfree='_is_alcohol_free', pack='Packaging'),
    'rema': dict(src=os.path.join(ROOT, 'Produkter', 'rema1000_version 2.xlsx'), out=os.path.join(ROOT, 'Produkter', 'rema1000_version 2_ny.xlsx'),
                 title='Hello Cal product title', type='Product type', var=['Variant', 'taste'], kw=['Keyword 1'],
                 sugar='is_sugar_free', vegan='is_vegan', alc=None, gluten='is_gluten_free', lactose='is_lactose_free',
                 fat='fat', alcfree='is_alcohol_free', pack=None),
}
DBNAMES = {  # arkets kolonne -> navn i serverens database (prisma Product/Barcode)
    'Brand': 'brand', 'Subbrand': 'subbrand', 'Product Type': 'productType', 'Variation': 'variant', 'Quantity': 'packageSizeText',
    'Pack Count': 'packCount', 'Packaging': 'packaging', 'Category': 'category', 'EAN': 'barcode',
    'Keyword 1': 'keyword1', 'Keyword 2': 'keyword2', 'Keyword 3': 'keyword3', 'Keyword 4': 'keyword4', 'Keyword 5': 'keyword5',
    'Variant': 'variant', 'taste': 'flavor', 'Amount': 'packCount', 'Product type': 'productType',
}
TAIL = ['Original Title', 'Product Name', 'Subtitle', 'Source URL', 'Source url', 'Image File', 'Image file', 'Parse Status', 'Parse status']
JA = {'is_vegan': 'vegansk', 'is_lactose_free': 'laktosefri', 'is_gluten_free': 'glutenfri', 'is_sugar_free': 'sukkerfri',
      'is_alcohol_free': 'alkoholfri', 'is_biological': 'økologisk', 'is_whole_grain': 'fuldkorn'}
STOP = {'og', 'med', 'uden', 'af', 'i', 'på', 'fra', 'tilsat', 'eller', 'til', 'en', 'et', 'de', 'den', 'det', 'for', 'er', 'ikke', 'som'}
PROPER = {'beluga', 'hokkaido', 'thüringer', 'schwarzwalder', 'serrano', 'marmite', 'curacao'}
SUGAR_KW = {'tilsat sukker': 'med tilsat sukker', 'indeholder sukker': 'med tilsat sukker', 'indeholder tilsat sukker': 'med tilsat sukker',
            'lavt sukkerindhold': 'med lavt sukkerindhold', 'tilsat sødestof': 'med tilsat kunstig sødestof',
            'tilsat sødestoffer': 'med tilsat kunstig sødestof', 'tilsat kunstigt sødestof': 'med tilsat kunstig sødestof',
            'tilsat kunstig sødestof': 'med tilsat kunstig sødestof', 'indeholder sødestoffer': 'med tilsat kunstig sødestof'}
SUGAR_RE = re.compile(r'\b(sukkerfri[a-zæøå]*|uden tilsat sukker|ikke tilsat sukker)\b', re.I)


def lower_phrase(s):
    out = []
    for w in re.split(r'(\s+)', s):
        if w.strip() and (w.upper() == w and len(re.sub(r'[^A-Za-zÆØÅæøå]', '', w)) >= 2):
            out.append(w)
        elif w.lower().strip(',.') in PROPER:
            out.append(w)
        else:
            out.append(w.lower())
    return ''.join(out)


def tokens(s):
    return re.findall(r"[\wæøåÆØÅéüö]+", s.lower())


def clean_phrase(s):
    s = re.sub(r'\s*,\s*', ', ', s)
    s = re.sub(r'(,\s*){2,}', ', ', s)
    words = s.split()
    changed = True
    while changed and words:
        changed = False
        w0 = words[0].strip(',').lower()
        if w0 in STOP and w0 not in ('uden', 'med', 'i', 'af', 'på', 'fra', 'tilsat', 'ikke') or words[0] in (',', ):
            words.pop(0); changed = True; continue
        wl = words[-1].strip(',').lower()
        if wl in STOP or words[-1] == ',':
            words.pop(); changed = True; continue
        for i in range(len(words) - 1):
            a, b = words[i].strip(',').lower(), words[i + 1].strip(',').lower()
            if a in STOP and b in STOP and a == b:
                words.pop(i + 1); changed = True; break
    s = ' '.join(words).strip(' ,')
    return s


def dedupe(s, seen):
    """fjern ord (ikke funktionsord), der allerede er brugt i et tidligere felt eller tidligere i feltet"""
    if not s:
        return s
    out = []
    removed = False
    for w in s.split(' '):
        t = tokens(w)
        if t and all(x not in STOP and x in seen for x in t) and not re.search(r'\d', w):
            removed = True
            continue
        out.append(w)
        for x in t:
            if x not in STOP:
                seen.add(x)
    r = ' '.join(out)
    return clean_phrase(r) if removed else r


# ---------- flertal ----------
def load_frida_plurals():
    ws = openpyxl.load_workbook(FRIDA, read_only=True).active
    m = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] and r[24] and r[0] != 'A38':
            m[str(r[0]).strip().lower()] = str(r[24]).strip()
    return m


MASS = ('mel', 'sukker', 'olie', 'salt', 'smør', 'mælk', 'saft', 'vand', 'øl', 'brød', 'kød', 'pålæg', 'slik', 'juice', 'kaffe', 'te', 'ris',
        'pasta', 'ketchup', 'sennep', 'yoghurt', 'skyr', 'mayonnaise', 'remoulade', 'gær', 'cider', 'sirup', 'honning', 'ost', 'fisk', 'fløde',
        'is', 'chips', 'müsli', 'mysli', 'granola', 'havregryn', 'tofu', 'tempeh', 'kakao', 'eddike', 'bouillon', 'fond', 'suppe', 'sovs', 'sauce',
        'dressing', 'pesto', 'hummus', 'marmelade', 'syltetøj', 'nutella', 'snack', 'cola', 'sodavand', 'drik', 'drink', 'smoothie', 'kaffe',
        'chokolade', 'lakrids', 'karamel', 'nougat', 'marcipan', 'kiks', 'knækbrød', 'rugbrød', 'toast', 'pizza', 'dej', 'lage', 'pulver')
ADJ = {'rød': 'røde', 'grøn': 'grønne', 'hvid': 'hvide', 'sort': 'sorte', 'gul': 'gule', 'blandet': 'blandede', 'tørret': 'tørrede', 'kogt': 'kogte',
       'hakket': 'hakkede', 'frisk': 'friske', 'stor': 'store', 'lille': 'små', 'rå': 'rå', 'fersk': 'ferske', 'røget': 'røgede', 'syltet': 'syltede',
       'stegt': 'stegte', 'sød': 'søde', 'sur': 'sure', 'blå': 'blå', 'lys': 'lyse', 'mørk': 'mørke', 'blød': 'bløde', 'skimmelmodnet': 'skimmelmodnede', 'modnet': 'modnede', 'fuldkorn': 'fuldkorn'}
PREP = {'med', 'af', 'i', 'til', 'uden', 'på', 'og', 'fra', 'for'}
IRREG = {'mand': 'mænd', 'datter': 'døtre', 'bog': 'bøger', 'ben': 'ben', 'barn': 'børn', 'æg': 'æg', 'ærte': 'ærter', 'bær': 'bær',
         'løg': 'løg', 'knoglemarv': 'knoglemarv', 'sandwich': 'sandwiches', 'burger': 'burgere', 'cookie': 'cookies', 'pommes frites': 'pommes frites',
         'tomat': 'tomater', 'agurk': 'agurker', 'kartoffel': 'kartofler', 'æble': 'æbler', 'banan': 'bananer', 'appelsin': 'appelsiner',
         'citron': 'citroner', 'gulerod': 'gulerødder', 'rod': 'rødder', 'bolle': 'boller', 'hotdog': 'hotdogs', 'wrap': 'wraps', 'nudel': 'nudler',
         'kapsel': 'kapsler', 'pose': 'poser', 'pølse': 'pølser', 'bar': 'barer', 'knækbrød': 'knækbrød', 'rugbrød': 'rugbrød',
         'tyggegummi': 'tyggegummi', 'vin': 'vine', 'gin': 'gin', 'rom': 'rom', 'vodka': 'vodka', 'whisky': 'whisky', 'cognac': 'cognac'}


SUFFIX_PL = {  # kun sikre, taelleligt navneord (ellers: ental = flertal)
    'pølse': 'pølser', 'kage': 'kager', 'bolle': 'boller', 'frikadelle': 'frikadeller', 'filet': 'fileter', 'bryst': 'bryster',
    'stang': 'stænger', 'stykke': 'stykker', 'skive': 'skiver', 'tomat': 'tomater', 'agurk': 'agurker', 'gulerod': 'gulerødder',
    'kartoffel': 'kartofler', 'æble': 'æbler', 'pære': 'pærer', 'banan': 'bananer', 'appelsin': 'appelsiner', 'citron': 'citroner',
    'lime': 'limer', 'ring': 'ringe', 'kapsel': 'kapsler', 'pind': 'pinde', 'kugle': 'kugler', 'tablet': 'tabletter', 'pastil': 'pastiller',
    'bønne': 'bønner', 'nød': 'nødder', 'mandel': 'mandler', 'rosin': 'rosiner', 'dadel': 'dadler', 'figen': 'figner', 'burger': 'burgere',
    'sandwich': 'sandwiches', 'wrap': 'wraps', 'pizza': 'pizzaer', 'flaske': 'flasker', 'dåse': 'dåser', 'rulle': 'ruller', 'snegl': 'snegle',
    'croissant': 'croissanter', 'bagel': 'bagels', 'muffin': 'muffins', 'cookie': 'cookies', 'bøf': 'bøffer', 'steak': 'steaks',
    'kotelet': 'koteletter', 'vinge': 'vinger', 'plade': 'plader', 'tærte': 'tærter', 'kiwi': 'kiwier', 'avocado': 'avocadoer',
    'peberfrugt': 'peberfrugter', 'tortilla': 'tortillaer', 'pita': 'pitaer', 'bar': 'barer', 'nudel': 'nudler', 'ost': 'oste',
    'vin': 'vine', 'chili': 'chilier', 'paprika': 'paprikaer',
}
SHORT_SUF = {'ost', 'vin', 'bar', 'ring', 'stang', 'plade', 'bryst', 'bønne', 'nød', 'lime', 'pind', 'kiwi', 'pita', 'chili'}
ADJ = {'rød': 'røde', 'grøn': 'grønne', 'hvid': 'hvide', 'sort': 'sorte', 'gul': 'gule', 'blandet': 'blandede', 'tørret': 'tørrede', 'kogt': 'kogte',
       'hakket': 'hakkede', 'frisk': 'friske', 'stor': 'store', 'lille': 'små', 'fersk': 'ferske', 'røget': 'røgede', 'syltet': 'syltede',
       'stegt': 'stegte', 'sød': 'søde', 'sur': 'sure', 'lys': 'lyse', 'mørk': 'mørke', 'blød': 'bløde', 'skimmelmodnet': 'skimmelmodnede', 'modnet': 'modnede', 'fyldt': 'fyldte', 'saltet': 'saltede',
       'krydret': 'krydrede', 'marineret': 'marinerede', 'paneret': 'panerede'}
PREP = {'med', 'af', 'i', 'til', 'uden', 'på', 'og', 'fra', 'for', 'eller'}


def plural_word(w):
    lw = w.lower()
    for suf in sorted(SUFFIX_PL, key=len, reverse=True):
        if lw.endswith(suf) and (len(lw) > len(suf) + 2 or lw == suf or suf in SHORT_SUF):
            p = lw[:len(lw) - len(suf)] + SUFFIX_PL[suf]
            return w[0] + p[1:] if w[0].isupper() else p
    return w


def plural_phrase(t, fr):
    key = t.strip().lower()
    if key in fr:
        return fr[key]
    words = t.strip().split(' ')
    idx = len(words)
    for i, w in enumerate(words):
        if i > 0 and (w.lower() in PREP or w.startswith('(') or re.search(r'\d', w) or w.endswith(',')):
            idx = i
            break
    head = idx - 1
    hw = words[head].rstrip(',')
    out = words[:]
    pl = plural_word(hw)
    if pl == hw:
        return t
    out[head] = pl + words[head][len(hw):]
    for i in range(head):
        a = words[i].lower()
        if a in ADJ:
            out[i] = ADJ[a] if words[i][0].islower() else ADJ[a].capitalize()
    return ' '.join(out)


# ---------- formler ----------
def load_templates():
    ws = openpyxl.load_workbook(FRIDA).active
    return ws['B2'].value, ws['C2'].value


def build_formula(tpl, row, m, extra_fl=None, var_expr=None):
    def sub(mo):
        col = mo.group(1)
        v = m.get(col, '""')
        if v.startswith('"'):
            return v
        return f'{v}{row}'
    t = tpl
    if extra_fl:
        assert t.count('LOWER(Z2))') == 1
        t = t.replace('LOWER(Z2))', 'LOWER(Z2),' + ','.join(f'LOWER(@@X{n}@@)' for n in range(len(extra_fl))) + ')')
    t = re.sub(r'(?<![A-Za-z_.])([A-Z]{1,2})2(?![0-9A-Za-z])', sub, t)
    for n, c in enumerate(extra_fl or []):
        t = t.replace(f'@@X{n}@@', f'{c}{row}')
    return t


COOKED = {'stegt', 'friturestegt', 'tørret', 'kogt', 'syltet', 'tørristet', 'ristet', 'dampet', 'grillet', 'bagt'}
RAW = {'rå', 'fersk'}
OWN_KW = ['vild', 'raffinol', 'hydrogeneret', 'grove', 'grov', 'parboiled', 'tør', 'sød']
OWN_PHRASES = [('på dåse/ konserves', r'på dåse(?:\s*/\s*konserves)?'), ('på glas', r'på glas'), ('i saltlage', r'i saltlage'),
               ('uden sten', r'uden sten'), ('hele eller knækkede', r'hele eller knækkede')]


def fmt_list(parts):
    parts = [p.strip() for p in parts if p.strip()]
    if len(parts) <= 1:
        return ''.join(parts)
    return ', '.join(parts[:-1]) + ' og ' + parts[-1]


def put_kw(nr, ki, val):
    if any(isinstance(nr[i], str) and nr[i].lower() == val.lower() for i in ki):
        return
    for i in ki:
        if not nr[i]:
            nr[i] = val
            return


def frida_rules(nr, vi, ki, nix, cfg, stats):
    ti = nix[cfg['type']]
    fi, ai, di = nix[cfg['fat']], nix[cfg['alcfree']], nix['_is_decaf']
    ci, ri, zi = nix['_is_cooked'], nix['_is_raw'], nix['_is_frozen']
    if cfg.get('pack') and isinstance(nr[nix[cfg['pack']]], str) and nr[nix[cfg['pack']]].lower() == 'frozen' and not nr[zi]:
        nr[zi] = 'frozen'
    if cfg['alc'] and isinstance(nr[nix[cfg['alc']]], str) and nr[nix[cfg['alc']]].lower() in ('alkoholfri', 'uden alkohol'):
        nr[ai] = 'alkoholfri'
        nr[nix[cfg['alc']]] = None
        stats['alkoholfri_flyttet'] += 1
    if isinstance(nr[ti], str) and re.search(r'\S\s+blade\b', nr[ti], re.I):
        nr[ti] = re.sub(r'(\S)\s+(blade)\b', r'\1\2', nr[ti], flags=re.I)
        stats['blade_sammen'] += 1
    for i in vi + ki:
        v = nr[i]
        if not isinstance(v, str):
            continue
        v = re.sub(r'\bm\.\s*', 'med ', v)
        v = re.sub(r'\bm/\s*', 'med ', v)
        v = re.sub(r'\bu\.\s*', 'uden ', v)
        v = re.sub(r'\bu/\s*', 'uden ', v)
        v = re.sub(r'\busukrede\b', 'uden sukker', v)
        v = re.sub(r'\bsukrede\b', 'med sukker', v)
        v = re.sub(r'(\S)\s+(blade)\b', r'\1\2', v)
        mo = re.search(r'(\d+)\s*[/,.]\s*(\d+)\s*%\s*(?:fedt)?', v) or re.search(r'(\d+)()\s*%\s*fedt', v)
        if mo and (re.search(r'fedt', v) or mo.group(2)):
            val = mo.group(1) + (',' + mo.group(2) if mo.group(2) else '') + '%'
            if not nr[fi]:
                nr[fi] = val
            v = v[:mo.start()] + v[mo.end():]
            v = re.sub(r'\bfedt\b', '', v)
            stats['fedt_flyttet'] += 1
        v2 = re.sub(r'\b\d+(?:[,.]\d+)?\s*(?:kg|gram|g)\b', '', v)
        if v2 != v:
            stats['gram_fjernet'] += 1
            v = v2
        if re.search(r'\bkoffeinfr[it]*\b', v):
            nr[di] = 'koffeinfri'
            v = re.sub(r'\bkoffeinfr[it]*\b', '', v)
            stats['koffeinfri_flyttet'] += 1
        if re.search(r'\b(alkoholfri|uden alkohol)\b', v):
            nr[ai] = 'alkoholfri'
            v = re.sub(r'\b(alkoholfri|uden alkohol)\b', '', v)
            stats['alkoholfri_flyttet'] += 1
        if re.search(r'\buht\b', v, re.I):
            v = re.sub(r'\buht\b', 'langtidsholdbar', v, flags=re.I)
            stats['uht'] += 1
        if '/' in v and ' men ' not in v and not re.search(r'konserves', v):
            v = fmt_list(re.split(r'\s*/\s*', v))
            stats['liste'] += 1
        elif ',' in v and i in vi:
            v = fmt_list(re.split(r'\s*,\s*', v))
            stats['liste'] += 1
        nr[i] = clean_phrase(v.strip()) or None
    for i in vi:
        v = nr[i]
        if not isinstance(v, str):
            continue
        for kw, pat in OWN_PHRASES:
            if re.search(r'(?<!\w)' + pat + r'(?!\w)', v):
                v = re.sub(r'(?<!\w)' + pat + r'(?!\w)', '', v)
                put_kw(nr, ki, kw)
                stats['keyword_ud_af_variation'] += 1
        for w in OWN_KW:
            if re.search(r'(?<!\w)' + w + r'(?!\w)', v):
                v = re.sub(r'(?<!\w)' + w + r'(?!\w)', '', v)
                put_kw(nr, ki, w)
                stats['keyword_ud_af_variation'] += 1
        nr[i] = clean_phrase(v.strip()) or None
    for i in vi + ki:
        v = nr[i]
        if not isinstance(v, str):
            continue
        t = v.strip().lower()
        if t in COOKED:
            nr[ci] = t if not nr[ci] or nr[ci] == t else nr[ci] + ' og ' + t
            nr[i] = None
            stats['cooked_flyttet'] += 1
        elif t in RAW:
            nr[ri] = nr[ri] or t
            nr[i] = None
            stats['raw_flyttet'] += 1
        elif t in ('frosset', 'frost', 'dybfrost', 'dybfrossen'):
            nr[zi] = nr[zi] or 'frozen'
            nr[i] = None
        elif t.startswith('vitamin'):
            nr[i] = None
            stats['vitamin_keyword_fjernet'] += 1
    kv = [nr[i] for i in ki if nr[i]]
    for j, i in enumerate(ki):
        nr[i] = kv[j] if j < len(kv) else None


def main(which):
    cfg = dict(CFG[which])
    if os.environ.get('SRC_' + which.upper()):
        cfg['src'] = os.environ['SRC_' + which.upper()]
    fr = load_frida_plurals()
    tplB, tplC = load_templates()
    os.makedirs(os.path.join(ROOT, 'Backup'), exist_ok=True)
    if not os.environ.get('SRC_' + which.upper()):
        shutil.copy2(cfg['src'], os.path.join(ROOT, 'Backup', os.path.basename(cfg['src']).replace('.xlsx', '_foer_ental-flertal.xlsx')))
    wb = openpyxl.load_workbook(cfg['src'])
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    ix = {}
    for i, h in enumerate(hdr):
        ix.setdefault(h, i)  # foerste forekomst
    # kolonner der skal bruges
    tcol = ix[cfg['title']]
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    # nye kolonner: fjern titel, tilfoej foran ental+flertal, bagest frozen/raw/cooked/plural
    new_hdr = ['Product title singular', 'Product title plural'] + [h for i, h in enumerate(hdr) if i != tcol]
    for extra in ['_is_frozen', '_is_raw', '_is_cooked', '_is_decaf', cfg['alcfree'], 'Product type plural']:
        if extra not in new_hdr:
            new_hdr.append(extra)
    nix = {}
    for i, h in enumerate(new_hdr):
        nix.setdefault(h, i)
    new_rows = []
    is_cols = [i for i, h in enumerate(new_hdr) if h and (h.startswith('_is_') or h.startswith('is_')) and h not in ('_is_alcohol_pct',)]
    stats = collections.Counter()
    for r in rows:
        nr = [None, None] + [v for i, v in enumerate(r) if i != tcol] + [None] * (len(new_hdr) - 2 - (len(hdr) - 1))
        # strip alle tekstceller
        for i in range(len(nr)):
            if isinstance(nr[i], str):
                nr[i] = nr[i].strip() or None
        # _is_ -> smaat, 'Ja' -> beskrivende ord
        for i in is_cols:
            v = nr[i]
            if isinstance(v, str):
                h = new_hdr[i]
                if v.lower() == 'ja' and h in JA:
                    v = JA[h]
                elif v.lower() == 'ja':
                    v = h.replace('_is_', '').replace('is_', '')
                if v != v.lower():
                    stats['is_smaat'] += 1
                nr[i] = v.lower()
        # variation + keywords
        vi = [nix[h] for h in cfg['var'] if h in nix]
        ki = [nix[h] for h in cfg['kw'] if h in nix]
        # 'Uden' + 'Tilsat sukker' (fejlagtigt delt) -> 'uden tilsat sukker'
        for n in range(len(ki) - 1):
            x, y = nr[ki[n]], nr[ki[n + 1]]
            if isinstance(x, str) and isinstance(y, str) and x.strip(' ,/').lower() == 'uden' and y.strip().lower().startswith('tilsat'):
                nr[ki[n]] = 'uden ' + y.strip().lower()
                nr[ki[n + 1]] = None
                stats['uden_tilsat_samlet'] += 1
        si = nix[cfg['sugar']]
        for i in vi + ki:
            v = nr[i]
            if isinstance(v, str):
                nr[i] = lower_phrase(v)
        frida_rules(nr, vi, ki, nix, cfg, stats)
        # sukker-ord som i Frida: 'tilsat sukker' o.l. er keyword med 'med' foran, sødestof -> _is_sweeteners
        swi = nix.get('_is_sweeteners')
        for i in ki:
            v = nr[i]
            if isinstance(v, str) and v.strip().lower() in SUGAR_KW:
                nr[i] = SUGAR_KW[v.strip().lower()]
                stats['sukker_keyword_normaliseret'] += 1
                if 'sødestof' in nr[i] and swi is not None and not nr[swi]:
                    nr[swi] = 'sødestoffer'
        if isinstance(nr[nix[cfg['sugar']]], str) and nr[nix[cfg['sugar']]].lower().startswith('ikke tilsat'):
            nr[nix[cfg['sugar']]] = 'uden tilsat sukker'
        # sukkerfri -> sukkerkolonne
        for i in vi + ki:
            v = nr[i]
            if isinstance(v, str) and SUGAR_RE.search(v):
                mo = SUGAR_RE.search(v)
                val = 'sukkerfri' if mo.group(1).lower().startswith('sukkerfri') else 'uden tilsat sukker'
                if not nr[si]:
                    nr[si] = val
                nr[i] = clean_phrase(SUGAR_RE.sub('', v)) or None
                stats['sukkerfri_flyttet'] += 1
        # dobbeltord
        seen = set()
        for t in tokens(str(nr[nix[cfg['type']]] or '')):
            seen.add(t)
        for i in is_cols:
            if isinstance(nr[i], str) and new_hdr[i] not in ('_is_fat', 'fat', '%', '_is_alcohol_pct') and not re.search(r'\d', nr[i]):
                for t in tokens(nr[i]):
                    seen.add(t)
        for i in vi + ki:
            v = nr[i]
            if isinstance(v, str):
                nv = dedupe(v, seen) or None
                if nv != v:
                    stats['dobbeltord_fjernet'] += 1
                nr[i] = nv
        # komprimer keywords mod venstre
        kv = [nr[i] for i in ki if nr[i]]
        for j, i in enumerate(ki):
            nr[i] = kv[j] if j < len(kv) else None
        # flertal
        ty = nr[nix[cfg['type']]]
        if ty:
            nr[nix['Product type plural']] = plural_phrase(str(ty), fr)
        new_rows.append(nr)
    # endelig raekkefoelge: de tre tekstkolonner bagerst, databasenavne
    order = [i for i, h in enumerate(new_hdr) if h not in TAIL] + [new_hdr.index(h) for h in TAIL if h in new_hdr]
    final_hdr = [DBNAMES.get(new_hdr[i], new_hdr[i]) for i in order]
    new_rows = [[nr[i] for i in order] for nr in new_rows]
    pos = {new_hdr[i]: k for k, i in enumerate(order)}
    nix = pos
    new_hdr = final_hdr
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.value = None
    ws.delete_rows(2, ws.max_row)
    for j, h in enumerate(new_hdr, 1):
        ws.cell(1, j, h)
    col = lambda name: L(nix[name] + 1)
    m = {'A': col(cfg['type']), 'F': col('_is_frozen'), 'G': col('_is_raw'), 'H': col('_is_cooked'), 'J': '""', 'X': col(cfg['alcfree']), 'Z': col('_is_decaf'),
         'S': col(cfg['vegan']), 'T': col(cfg['sugar']), 'Y': col('Product type plural'),
         'N': col(cfg['kw'][0]), 'O': col(cfg['kw'][1]) if len(cfg['kw']) > 1 else '""', 'P': col(cfg['kw'][2]) if len(cfg['kw']) > 2 else '""',
         'K': col(cfg['alc']) if cfg['alc'] else '""', 'D': col(cfg['var'][0])}
    extra_fl = [col(cfg['gluten']), col(cfg['lactose'])]
    var_expr = None
    if len(cfg['var']) > 1:
        a, b = col(cfg['var'][0]), col(cfg['var'][1])
        m['D'] = '"@@VAR@@"'
        var_expr = lambda row: f'_xlfn.TEXTJOIN(" ",TRUE,{a}{row},{b}{row})'
    for r, nr in enumerate(new_rows, 2):
        for j, v in enumerate(nr, 1):
            if v is not None:
                ws.cell(r, j, v)
        fb = build_formula(tplB, r, m, extra_fl)
        fc = build_formula(tplC, r, m, extra_fl)
        if var_expr:
            fb = fb.replace('"@@VAR@@"', var_expr(r))
            fc = fc.replace('"@@VAR@@"', var_expr(r))
        ws.cell(r, 1, fb)
        ws.cell(r, 2, fc)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{L(len(new_hdr))}{len(new_rows) + 1}'
    for c, w in (('A', 40), ('B', 40)):
        ws.column_dimensions[c].width = w
    wb.save(cfg['out'])
    print(which, len(new_rows), 'raekker ->', cfg['out'])
    print(dict(stats))


if __name__ == '__main__':
    main(sys.argv[1])
