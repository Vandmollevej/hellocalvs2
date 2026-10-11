"""Optimering af dm, EDEKA, REWE og Wolt efter Bilka-moensteret (brugerens opgave 2026-10-10).
Brug: py -3 optimer_ark.py [dm edeka rewe wolt] [--dry]   (uden butik: alle fire; --dry skriver kun til TEMP, rører ikke arket)

Principper (Excelark/NAVNEREGLER.md):
- Originalkolonnerne (Original Title, Product Name, Subtitle, Source URL, Manufacturer, Image File, Parse Status, Servings)
  roeres ALDRIG - kun laeses. Titelformlerne i A/B bevares. Ingen nye arkfiler: arket gemmes oven i sig selv efter en
  backup i Excelark/backup.
- Praecision foerst: kun entydige rettelser. Tomme felter udfyldes; eksisterende vaerdier aendres kun naar reglen
  er en ren normalisering (stavning, store/smaa bogstaver, dobbeltord, brand/maengde/procent i feltet).
- Raekker med alkohol (vin, sekt, spiritus, oel) ejes af vin-sessionen og faar kun hygiejne (mellemrum, EAN som tekst,
  Quantity-stil, smaa bogstaver i _is_-felter).
- Tyske ark: tyske ord (bio, glutenfrei, zuckerfrei ...). Wolt er dansk og faar de danske ord.
- Alt hvad der aendres logges i en CSV i Excelark/backup (butik, raekke, kolonne, foer, efter, regel).
"""
import collections, csv, datetime, os, re, shutil, sys, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import omlaeg_til_bilka_kolonner as O  # noqa: E402  (norm_qty, bilka_master, plural_for)
import tyske_is_og_brand as TB  # noqa: E402  (danish_brands, norm, STOP)
import titel_ental_flertal as T  # noqa: E402

BACKUP = os.path.join(HERE, 'backup')
SHEETS = {
    'dm': (os.path.join(HERE, 'dm.xlsx'), 'de'),
    'edeka': (os.path.join(HERE, 'edeka.xlsx'), 'de'),
    'rewe': (os.path.join(HERE, 'rewe.xlsx'), 'de'),
    'wolt': (os.path.join(HERE, 'wolt.xlsx'), 'da'),
}
ORIG = {'Original Title', 'Product Name', 'Subtitle', 'Source URL', 'Manufacturer', 'Image File', 'Parse Status', 'Servings'}
FORMULA_COLS = {'Product title singular', 'Product title plural'}

WB = r'(?<![\wäöüßÄÖÜæøåÆØÅ])'
WA = r'(?![\wäöüßÄÖÜæøåÆØÅ])'
FUNC = {'de': r'(?:mit|ohne|in|im|aus|von|vom|für|auf|zum|zur|nach|bei|und|oder|als|ab|an|am|zu|über|unter|nur)',
        'da': r'(?:med|uden|i|af|på|fra|til|og|eller|som|for)'}


def ph(phrase):
    return re.compile(WB + re.escape(phrase) + WA, re.I)


def tidy(s):
    """oprydning efter at et ord/udtryk er fjernet"""
    if not isinstance(s, str):
        return s
    s = re.sub(r'\(\s*\)', ' ', s)
    s = re.sub(r'\s+([,;:.])', r'\1', s)
    s = re.sub(r'\s{2,}', ' ', s).strip(' ,;:-–/&+')
    prev = None
    while prev != s:  # haengende funktionsord i enderne
        prev = s
        s = re.sub(r'(?i)^(?:und|oder|mit|ohne|von|vom|aus|in|im|für|auf|og|med|uden|af|i|på|fra|&)\s+', '', s)
        s = re.sub(r'(?i)\s+(?:und|oder|mit|ohne|von|vom|aus|in|im|für|auf|og|med|uden|af|i|på|fra|&|-)$', '', s)
        s = s.strip(' ,;:-–/&+')
    return s


def cap1(s):
    return s[0].upper() + s[1:] if s else s


def lower1(s):
    return s[0].lower() + s[1:] if s else s


# ---------- maengde og pakker ----------
UNIT_RE = r'(kg|g|gr|mg|ml|cl|l|ltr|lt|liter|litre)'
QX = re.compile(r'(?<![\w.,/])(\d{1,3})\s*[xX]\s*(\d+(?:[.,]\d+)?)\s*' + UNIT_RE + r'(?![\wäöüßæøå])', re.I)
Q1 = re.compile(r'(?<![\w.,/-])(\d+(?:[.,]\d+)?)\s*' + UNIT_RE + r'(?![\wäöüßæøå])(?!\s*(?:vol|%))', re.I)
PACK = {
    'de': re.compile(r'(?<![\w.,/])(\d{1,3})\s*(?:er[- ]?(?:Pack|Set|Packung)|Stück|Stk\.?|ST|St\.?|Beutel|Portionen)(?![\wäöüß])'),
    'da': re.compile(r'(?<![\w.,/])(\d{1,3})\s*(?:-?\s*(?:stk|pk|pak|pack|poser?|bdt)\.?)(?![\wæøå])', re.I),
}
UNIT_MAP = {'g': 'g', 'gr': 'g', 'kg': 'kg', 'mg': 'mg', 'ml': 'ml', 'cl': 'cl', 'l': 'l', 'ltr': 'l', 'lt': 'l', 'liter': 'l', 'litre': 'l'}


def find_qty(texts):
    """(quantity, packcount, matched fragments) udledt af titlen - kun hvis entydigt"""
    quants, packs, frags = set(), set(), []
    for t in texts:
        if not t:
            continue
        t = str(t)
        m = QX.search(t)
        if m:
            q = f"{m.group(2).replace(',', '.')} {UNIT_MAP[m.group(3).lower()]}"
            quants.add(q)
            packs.add(int(m.group(1)))
            frags.append(m.group(0))
            continue
        for m in Q1.finditer(t):
            u = UNIT_MAP[m.group(2).lower()]
            if u == 'mg':
                continue
            quants.add(f"{m.group(1).replace(',', '.')} {u}")
            frags.append(m.group(0))
    for lang in PACK:
        for t in texts:
            if not t:
                continue
            for m in PACK[lang].finditer(str(t)):
                packs.add(int(m.group(1)))
                frags.append(m.group(0))
    q = next(iter(quants)) if len(quants) == 1 else None
    p = next(iter(packs)) if len(packs) == 1 else None
    return q, p, frags


FRAG_QTY = {
    'de': re.compile(r'(?i)\s*(?:,\s*)?(?:ergibt\s+)?(?:\d{1,3}\s*[x×]\s*)?\d+(?:[.,]\d+)?(?:\s*-\s*\d+(?:[.,]\d+)?)?\s*'
                     r'(?:kg|g|gr|ml|cl|l|ltr|liter|stück|stk\.?|st\.?|beutel|portionen)(?![\wäöüß])'),
    'da': re.compile(r'(?i)\s*(?:,\s*)?(?:\d{1,3}\s*[x×]\s*)?\d+(?:[.,]\d+)?\s*(?:kg|g|gr|ml|cl|l|ltr|lt|liter|stk\.?|pk|pak|pack|bdt)(?![\wæøå])'),
}

# ---------- _is_-ord ----------
WORD = {
    'de': dict(organic='bio', glutenfree='glutenfrei', lactose='laktosefrei', vegan='vegan', wholegrain='vollkorn',
               sugar='zuckerfrei', sugar_added='ohne zuckerzusatz', alcfree='alkoholfrei', decaf='entkoffeiniert',
               sweet='süßungsmittel', alcohol='enthält alkohol', frozen='frozen', welfare='haltungsform'),
    'da': dict(organic='økologisk', glutenfree='glutenfri', lactose='laktosefri', vegan='vegansk', wholegrain='fuldkorn',
               sugar='sukkerfri', sugar_added='uden tilsat sukker', alcfree='alkoholfri', decaf='koffeinfri',
               sweet='sødestoffer', alcohol='indeholder alkohol', frozen='frozen', welfare=''),
}
COUNTRY_DE = ['Deutschland', 'Italien', 'Spanien', 'Frankreich', 'Griechenland', 'Österreich', 'Niederlande', 'Polen', 'Irland',
              'Norwegen', 'Schottland', 'Neuseeland', 'Argentinien', 'Chile', 'Peru', 'Marokko', 'Ägypten', 'Israel', 'Südafrika',
              'Türkei', 'Portugal', 'Belgien', 'Schweiz', 'Dänemark', 'Thailand', 'Indien', 'Mexiko', 'Japan', 'China', 'Vietnam']
COUNTRY_RE = re.compile(WB + r'aus (?:der |den |dem )?(' + '|'.join(COUNTRY_DE) + r')' + WA, re.I)
IS_FLAG = [
    # (kolonne, vaerdinoegle, moenster paa titel/variation/keywords)
    ('_is_glutenfree', 'glutenfree', re.compile(WB + r'(?:glutenfrei\w*|ohne gluten|glutenfri\w*|uden gluten)' + WA, re.I)),
    ('_is_lactose_free', 'lactose', re.compile(WB + r'(?:la[ck]tosefrei\w*|ohne la[ck]tose|laktosefri\w*|uden laktose)' + WA, re.I)),
    ('_is_vegan', 'vegan', re.compile(WB + r'(?:vegan(?:e|er|es|en|em)?|vegansk(?:e)?)' + WA, re.I)),
    ('_is_whole_grain', 'wholegrain', re.compile(r'vollkorn|fuldkorn', re.I)),
    ('_is_decaf', 'decaf', re.compile(WB + r'(?:entkoffeiniert\w*|koffeinfrei\w*|ohne koffein|decaf|koffeinfri\w*)' + WA, re.I)),
    ('_is_sweeteners', 'sweet', re.compile(WB + r'(?:mit süßungsmittel\w*|mit süßstoff\w*|mit stevia|mit zuckerersatz|med sødestof\w*)' + WA, re.I)),
]
ORGANIC_RE = re.compile(WB + r'(?:bio|öko|ökologisch\w*|biologisch\w*|demeter|bioland|naturland|økologisk\w*|øko)' + WA, re.I)
ORGANIC_BRANDS = {'dmbio', 'alnatura', 'rapunzel', 'lebensbaum', 'davert', 'bio zentrale', 'bio gourmet', 'naturata', 'byodo', 'allos',
                  'bohlsener mühle', 'bauckhof', 'rewe bio', 'edeka bio'}
SUGAR_FREE_RE = re.compile(WB + r'(?:zuckerfrei\w*|sukkerfri\w*|zero sugar|zero zucker)' + WA, re.I)
SUGAR_ADDED_RE = re.compile(WB + r'(?:ohne (?:zugesetzten )?zucker(?:zusatz)?|ohne zusatz von zucker|ungesüßt\w*|uden tilsat sukker|usødet\w*)' + WA, re.I)
ALCFREE_RE = re.compile(WB + r'(?:alkoholfrei\w*|ohne alkohol|alkoholfri\w*|uden alkohol)' + WA + r'|0[.,]0\s*%\s*vol', re.I)
ALCPCT_RE = re.compile(r'(?<![\d.,])(\d{1,2}(?:[.,]\d{1,2})?)\s*%\s*vol', re.I)
FAT_EXPLICIT = re.compile(r'(?<![\d.,])(\d{1,2}(?:[.,]\d{1,2})?)\s*%\s*(?:fett|fedt)(?![a-zäöüßæøå])(?!\s*i\.?\s*tr)', re.I)
FAT_DAIRY_WORD = re.compile(r'milch|joghurt|jogurt|quark|sahne|schmand|kefir|frischkäse|mælk|yoghurt|fløde|skyr', re.I)
FAT_DAIRY_PCT = re.compile(r'(?<![\d.,])(\d(?:[.,]\d{1,2})?|10)\s*%(?!\s*[a-zäöüßæøå])', re.I)
FAT_NOT = re.compile(r'schokolade|kakao|riegel|milchreis|\beis\b|pulver|kaffee|protein|likör', re.I)
WELFARE_RE = re.compile(r'haltungsform\s*(\d)', re.I)
FROZEN_CAT = re.compile(r'tiefk(?:ü|ue)hl|\btk\b|gefroren|^eis im becher$|^eis am stiel$|^eis-großpackungen$|^eis-desserts$|frost', re.I)
FROZEN_TITLE = re.compile(WB + r'(?:tk|tiefgekühlt\w*|tiefgefroren\w*|gefroren\w*|frosne|frossen\w*|dybfrost\w*)' + WA, re.I)
ALC_CAT_RE = re.compile(WB + r'(?:wein\w*|rotwein|weißwein|weisswein|rosé\w*|sekt|prosecco|champagner|spirituos\w*|likör\w*|liköre|bitter|whisk(?:e)?y|'
                        r'gin|rum|vodka|wodka|brandy|weinbrand|cognac|grappa|tequila|wermut|portwein|sherry|schnaps|brände|met|aperitif|'
                        r'glühwein|punsch|cider|bier\w*|pils|lager|vin|spiritus|øl|rødvin|hvidvin|rosévin|mousserende)' + WA, re.I)

# keywords der egentlig er en anden kolonne
KW_FLAG = {  # noegle (smaat) -> (kolonne, vaerdinoegle)
    'ohne zuckerzusatz': ('_is_sugar_free', 'sugar_added'), 'ohne zucker': ('_is_sugar_free', 'sugar_added'),
    'ungesüßt': ('_is_sugar_free', 'sugar_added'), 'zuckerfrei': ('_is_sugar_free', 'sugar'),
    'ohne alkohol': ('_is_alcohol_free', 'alcfree'), 'alkoholfrei': ('_is_alcohol_free', 'alcfree'),
    'entkoffeiniert': ('_is_decaf', 'decaf'), 'koffeinfrei': ('_is_decaf', 'decaf'),
    'bio': ('_is_organic', 'organic'), 'vegan': ('_is_vegan', 'vegan'), 'glutenfrei': ('_is_glutenfree', 'glutenfree'),
    'laktosefrei': ('_is_lactose_free', 'lactose'), 'vollkorn': ('_is_whole_grain', 'wholegrain'),
    'mit süßungsmittel': ('_is_sweeteners', 'sweet'), 'mit süßungsmitteln': ('_is_sweeteners', 'sweet'),
    'uden tilsat sukker': ('_is_sugar_free', 'sugar_added'), 'sukkerfri': ('_is_sugar_free', 'sugar'),
    'økologisk': ('_is_organic', 'organic'), 'vegansk': ('_is_vegan', 'vegan'), 'glutenfri': ('_is_glutenfree', 'glutenfree'),
    'laktosefri': ('_is_lactose_free', 'lactose'), 'alkoholfri': ('_is_alcohol_free', 'alcfree'), 'fuldkorn': ('_is_whole_grain', 'wholegrain'),
}
PACK_KW = {'dose': 'Dose', 'dosen': 'Dose', 'glas': 'Glas', 'pet': 'PET', 'flasche': 'Flasche', 'tube': 'Tube', 'beutel': 'Beutel',
           'karton': 'Karton', 'tetra pak': 'Tetra Pak', 'tetrapak': 'Tetra Pak', 'schale': 'Schale', 'becher': 'Becher',
           'bakke': 'Bakke', 'dåse': 'Dåse', 'pose': 'Pose', 'flaske': 'Flaske', 'glas ': 'Glas'}
THEME = {  # butikkens temakategorier - ikke en produkttype
    'abverkauf', 'edeka-welt', 'elkos', 'edeka herzstücke', 'edeka bio', 'bio-produkte', 'sportnahrung', 'nahrungsergänzung', 'mehr',
    'süß & salzig', 'weihnachtszeit', 'alkoholfrei', 'schottland', 'deutschland', 'frankreich', 'italien', 'irland', 'usa', 'spanien',
    'international', 'filter', 'fürs büro', 'exklusiv: coffeeb', 'gepa', 'hipp bio', 'bio gourmet', 'rettergut', 'little lunch', 'bauckhof',
    'besh bio-konserven', 'bio marken', 'edeka genussmomente', 'vegane ernährung', 'laktosefreie ernährung', 'glutenfreie ernährung',
    'internationale küche', 'mediterrane küche', 'mexikanische küche', 'asiatische küche', 'getränke', 'frühstück & snacks',
}
MEAT_WORD = {'de': {'unbekannt'}, 'da': {'ukendt'}}


class Sheet:
    def __init__(self, name, path, lang):
        self.name, self.path, self.lang = name, path, lang


def lower_flag(v, col):
    """_is_-vaerdier med smaat; lande med stort begyndelsesbogstav; forkortelser (USA, MSC) beholdes"""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if col == '_is_country_of_origen':
        if s.islower() or s.isupper():
            if s.isupper() and len(s) <= 4:
                return s
            return ' '.join(w.capitalize() if w not in ('/', 'und', 'and', 'og') else w for w in s.split(' '))
        return s
    out = []
    for w in s.split(' '):
        out.append(w if (w.isupper() and 2 <= len(w) <= 5 and w.isalpha()) else w.lower())
    return ' '.join(out)


def run_sheet(sh, brands, by_len, plural_ctx, dry, audit, report):
    path, lang = sh.path, sh.lang
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock) and not dry:
        print(f'{sh.name}: SPRUNGET OVER - {os.path.basename(path)} er aaben i Excel')
        return None
    mtime = os.path.getmtime(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[1]]
    ix = {}
    for i, h in enumerate(hdr):
        ix.setdefault(h, i)
    alc_cols = [i for i, h in enumerate(hdr) if h == '_is_alcohol']
    I = lambda h: ix[h]  # noqa: E731
    cpct = alc_cols[1]
    own_text = [i for i, h in enumerate(hdr) if h not in ORIG and h not in FORMULA_COLS and h]
    W = WORD[lang]
    stats = collections.Counter()
    examples = collections.defaultdict(list)
    pc_is_int = None

    def log(rn, col, old, new, rule):
        stats[rule] += 1
        audit.append([sh.name, rn, col, old, new, rule])
        if len(examples[rule]) < 6:
            examples[rule].append((rn, col, old, new))

    rows = list(ws.iter_rows(min_row=2))
    # pack count: heltal eller tekst?
    ints = sum(isinstance(r[I('Pack Count')].value, int) for r in rows)
    strs = sum(isinstance(r[I('Pack Count')].value, str) for r in rows)
    pc_is_int = ints >= strs

    for row in rows:
        rn = row[0].row
        vals = [c.value for c in row]
        if not any(v not in (None, '') for v in vals):
            continue
        new = list(vals)

        def get(h, k=None):
            v = new[ix[h] if k is None else k]
            return v.strip() if isinstance(v, str) else v

        def put(h, v, rule, k=None):
            i = ix[h] if k is None else k
            old = new[i]
            if (old or None) == (v or None):
                return
            new[i] = v
            log(rn, hdr[i] if k is None else f'{hdr[i]}#{alc_cols.index(k) + 1 if k in alc_cols else k}', old, v, rule)

        # --- hygiejne (alle raekker) ---
        for i in own_text:
            v = new[i]
            if isinstance(v, str):
                nv = re.sub(r'\s{2,}', ' ', v.replace(' ', ' ')).strip() or None
                if nv != v:
                    new[i] = nv
                    log(rn, hdr[i], v, nv, 'hygiejne_mellemrum')
        ean = new[I('EAN')]
        if isinstance(ean, (int, float)) and not isinstance(ean, bool):
            put('EAN', str(int(ean)), 'ean_som_tekst')
        q = new[I('Quantity')]
        if isinstance(q, str) and O.norm_qty(q) != q:
            put('Quantity', O.norm_qty(q), 'quantity_bilka_stil')
        for h, i in ix.items():
            if h.startswith('_is_') and i != cpct and isinstance(new[i], str):
                nv = lower_flag(new[i], h)
                if nv != new[i]:
                    put(h, nv, 'is_smaat_bogstaver')
        v = new[cpct]
        if isinstance(v, str):
            nv = re.sub(r'(?<=\d),(?=\d)', '.', v).replace(' %', '%')
            if nv != v:
                put('_is_alcohol', nv, 'decimalpunktum', k=cpct)
        for h in ('_is_fat',):
            v = new[ix[h]]
            if isinstance(v, str) and re.search(r'\d,\d', v):
                put(h, re.sub(r'(?<=\d),(?=\d)', '.', v), 'decimalpunktum')
        # --- alkohol-raekker: kun hygiejne ---
        cat = str(new[I('Category')] or '')
        ptx = str(new[I('Product Type')] or '')
        alc_row = bool(new[alc_cols[0]] or new[cpct]) or bool(ALC_CAT_RE.search(cat) or ALC_CAT_RE.search(ptx))
        if alc_row and not new[I('_is_alcohol_free')] and not ALCFREE_RE.search(str(new[I('Original Title')] or '') + ' ' + str(new[I('Product Name')] or '')):
            for i in range(len(new)):
                if new[i] != vals[i]:
                    row[i].value = new[i]
            continue
        # --- tekstgrundlag ---
        ot, pn = str(vals[I('Original Title')] or ''), str(vals[I('Product Name')] or '')
        title_txt = ot + ' | ' + pn
        brand0 = get('Brand') or ''
        # --- brand ---
        if isinstance(brand0, str) and brand0 and (len(brand0) > 70 or len(brand0.split()) > 9 or '*' in brand0 or 'Nielsen' in brand0):
            put('Brand', None, 'brand_skrald_fjernet')
        if not get('Brand'):
            tn = [TB.norm(ot), TB.norm(pn), TB.norm(re.sub(r'(?i)^(?:bio|neu)\s+', '', ot))]
            hit = None
            for k in by_len:
                if any(t == k or t.startswith(k + ' ') for t in tn if t):
                    hit = k
                    break
            if hit:
                b, subs = brands[hit]
                put('Brand', b, 'brand_udfyldt_fra_titel')
                if not get('Subbrand'):
                    rest = [t[len(hit):] if t.startswith(hit) else t for t in tn if t]
                    for sk in sorted(subs, key=len, reverse=True):
                        if any(re.search(r'(?<![a-z0-9])' + re.escape(sk) + r'(?![a-z0-9])', t) for t in rest):
                            put('Subbrand', subs[sk], 'subbrand_udfyldt_fra_titel')
                            break
        else:  # samme maerke, anden stavemaade (kun store/smaa bogstaver og tegnsaetning) -> flertalsstavemaaden
            k = TB.norm(get('Brand'))
            if k in brands and brands[k][0] != get('Brand'):
                a_, b_ = get('Brand'), brands[k][0]
                if re.sub(r'[^\wäöüßæøåé]', '', a_.lower()) == re.sub(r'[^\wäöüßæøåé]', '', b_.lower()):
                    put('Brand', b_, 'brand_stavemaade')
        brand, sub = get('Brand') or '', get('Subbrand') or ''
        if brand and brand.lower() in ORGANIC_BRANDS and not get('_is_organic') and lang == 'de':
            put('_is_organic', W['organic'], 'is_bio_fra_oekomaerke')
        # --- maengde og pakke ---
        if not get('Quantity') or not get('Pack Count'):
            qv, pv, frags = find_qty([ot, pn])
            if qv and not get('Quantity'):
                put('Quantity', O.norm_qty(qv), 'quantity_fra_titel')
            if pv and not get('Pack Count') and (pv > 1 or lang == 'de'):
                put('Pack Count', pv if pc_is_int else str(pv), 'packcount_fra_titel')
        # --- produkttype ---
        pt = get('Product Type') or ''
        var = get('Variation') or ''
        kws = [get(f'Keyword {n}') or '' for n in range(1, 6)]
        pt0, var0, kws0 = pt, var, list(kws)
        # procent/alkohol/fedt ud af tekstfelterne
        def pct_pass(text):
            nonlocal_changes = {}
            t = re.sub(r'(?<=\d)\s*/\s*(?=\d+\s*%)', '.', text)  # "2/ 3%" er 2,3% (komma blev tidligere til skraastreg)
            m = re.search(r'(?i)alkoholgehalt:?\s*(\d+(?:[.,]\d+)?)\s*%\s*vol\.?', t)
            if m:
                nonlocal_changes['alc'] = m.group(1).replace(',', '.') + '%'
                t = t.replace(m.group(0), ' ')
            m = re.search(r'(?i)(\d{2})\s*%\s*fett\s*i\.?\s*tr\.?', t)
            if m:
                nonlocal_changes['cheese'] = m.group(1) + '+'
                t = t.replace(m.group(0), ' ')
            m = re.search(r'(?i)(\d{1,2}(?:[.,]\d{1,2})?)\s*%\s*fett(?![a-zäöüß])(?!\s*i\.?\s*tr)', t)
            if m:
                nonlocal_changes['fat'] = m.group(1).replace(',', '.') + '%'
                t = t.replace(m.group(0), ' ')
            t2 = re.sub(r'\(?\s*\d+(?:[.,]\d+)?\s*(?:%|prozent)\s*\)?', ' ', t, flags=re.I)
            if t2 != t:
                nonlocal_changes['ingr'] = True
                t = t2
            return tidy(t), nonlocal_changes

        for fld in ('pt', 'var'):
            text = pt if fld == 'pt' else var
            if text and re.search(r'%|prozent|alkoholgehalt', text, re.I):
                nt, ch = pct_pass(text)
                if 'alc' in ch and not get('_is_alcohol', cpct):
                    put('_is_alcohol', ch['alc'], 'alkoholprocent_fra_tekst', k=cpct)
                    if not get('_is_alcohol', alc_cols[0]):
                        put('_is_alcohol', W['alcohol'], 'alkoholflag_fra_procent', k=alc_cols[0])
                if 'fat' in ch and not get('_is_fat'):
                    put('_is_fat', ch['fat'], 'fedtprocent_fra_tekst')
                if 'cheese' in ch and fld == 'var':
                    nt = (ch['cheese'] + (', ' + nt if nt else ''))
                elif 'cheese' in ch and fld == 'pt':
                    var = ch['cheese'] + (', ' + var if var else '')
                if fld == 'pt':
                    pt = nt
                else:
                    var = nt
        # maengde-brudstykker ud af typen/varianten (flyttes til Quantity/Pack Count hvis tomme)
        for fld in ('pt', 'var'):
            text = pt if fld == 'pt' else var
            if not text:
                continue
            m = FRAG_QTY[lang].search(text)
            if m and (len(text) - len(m.group(0))) >= 3:
                q2, p2, _ = find_qty([m.group(0)])
                if q2 and not get('Quantity'):
                    put('Quantity', O.norm_qty(q2), 'quantity_fra_produkttype')
                if p2 and p2 > 1 and not get('Pack Count'):
                    put('Pack Count', p2 if pc_is_int else str(p2), 'packcount_fra_produkttype')
                nt = tidy(FRAG_QTY[lang].sub(' ', text))
                if fld == 'pt':
                    pt = nt
                else:
                    var = nt
        # brand/subbrand ud af produkttypen
        for b in (brand, sub):
            if pt and b and len(b) >= 3 and pt.strip().lower() != b.strip().lower() and b.strip().lower() not in {'dinkel', 'wiener', 'highland', 'bio', 'frisk'}:
                if ph(b).search(pt):
                    nt = tidy(ph(b).sub(' ', pt))
                    # resten skal stadig indeholde et rigtigt ord (ikke kun "6-pak (daase)")
                    words = [w for w in re.findall(r'[A-Za-zäöüßæøåÄÖÜÆØÅ]{3,}', nt) if w.lower() not in {'pak', 'dåse', 'dose', 'pack', 'stk', 'flaske'}]
                    if nt and words:
                        pt = cap1(nt)
        # produkttype = bare brandet, eller en temakategori -> titel uden brand/maengde
        bad_pt = (not pt) or bool(brand and pt.lower() == brand.lower()) or pt.lower() in THEME
        pt_rule = 'produkttype_oprydning'
        if bad_pt and lang == 'de':
            nm = re.sub(r'\s*\([^()]*\)\s*$', '', pn)
            nm = re.sub(r',\s*(?:\d+\s*[x×]\s*)?\d+(?:[.,]\d+)?\s*(?:kg|g|ml|cl|l|st|stück)\s*$', '', nm, flags=re.I)
            for b in (brand, sub):
                if b and len(b) >= 2 and re.match(re.escape(b) + r'\b', nm, re.I):
                    nm = nm[len(b):]
            nm = tidy(FRAG_QTY['de'].sub(' ', nm))
            nm = tidy(re.sub(WB + r'bio' + WA, ' ', nm, flags=re.I))  # bio -> _is_organic (ordet slettes fra egne kolonner)
            cand = nm
            if not cand and cat and cat.lower() not in THEME and cat.lower() != brand.lower():
                cand = cat
            if cand and cand != pt:
                pt = cap1(cand)
                pt_rule = 'produkttype_fra_navn_eller_kategori'
        # dobbeltord i typen og variationen (samme token to gange): det andet ord fjernes sammen med et smaat
        # ord lige foran det (adjektiv/funktionsord), saa der ikke staar "würzige" eller "in" alene tilbage
        def dedupe_phrase(text):
            toks = text.split(' ')
            seen, out = set(), []
            for w in toks:
                k = w.lower().strip(',.')
                if len(k) >= 3 and k in seen and k not in {'und', 'mit', 'ohne', 'for', 'med'}:
                    if out and out[-1][:1].islower() and len(out) > 1:
                        out.pop()
                    elif out and re.fullmatch(r'(?i)' + FUNC[lang], out[-1].strip(',.')):
                        out.pop()
                    continue
                seen.add(k)
                out.append(w)
            return ' '.join(out) if out != toks else text
        if pt:
            pt = tidy(dedupe_phrase(pt)) or pt
        # variation: skraastreger -> opremsning, smaa funktionsord
        def slash_list(v):
            if '/' not in v or re.search(r'\d\s*/\s*\d', v):
                return v
            parts = [p.strip() for p in re.split(r'\s*/\s*', v) if p.strip()]
            if len(parts) < 2:
                return tidy(v.replace('/', ' '))
            items = []
            for p in parts:
                if items and re.match(r'(?i)' + FUNC[lang] + r'\b', p):
                    items[-1] += ' ' + p
                else:
                    items.append(p)
            if len(items) == 1:
                return items[0]
            und = 'und' if lang == 'de' else 'og'
            return (', '.join(items[:-1]) + f' {und} ' + items[-1]) if len(items) > 2 else f'{items[0]} {und} {items[1]}'
        if var:
            var = slash_list(var)
            m = re.match(r'(?i)^(' + FUNC[lang] + r')\b', var)
            if m and var[0].isupper():
                var = lower1(var)
        # variation: kun hvis ALLE betydningsbaerende ord allerede staar i typen, er variationen overfloedig
        if var:
            var = tidy(dedupe_phrase(var)) or var
        if var and pt:
            ptoks = {w.lower().strip(',.') for w in re.split(r'[\s-]+', pt) if len(w) >= 4}
            vtoks = [w.lower().strip(',.') for w in re.split(r'[\s-]+', var) if len(w) >= 4]
            if vtoks and all(t in ptoks for t in vtoks):
                var = ''
        if var and var.lower() == pt.lower():
            var = ''
        # keywords: flag-ord -> egen kolonne, pakker -> Packaging, procent/kategori vk, dobbeltord, luk huller
        kept = []
        own_words = {w.lower().strip(',.') for w in re.split(r'[\s-]+', ' '.join([pt, var, brand, sub])) if len(w) >= 3}
        for kwv in kws:
            if not kwv:
                continue
            kl = kwv.strip().lower()
            if kl in KW_FLAG:
                col, wk = KW_FLAG[kl]
                if col in ix and not get(col):
                    put(col, W[wk], 'keyword_til_egen_kolonne')
                else:
                    stats['keyword_flag_dublet_fjernet'] += 1
                continue
            mw = WELFARE_RE.fullmatch(kl)
            if mw and '_is_animal_welfare' in ix:
                if not get('_is_animal_welfare'):
                    put('_is_animal_welfare', f'haltungsform {mw.group(1)}', 'keyword_til_egen_kolonne')
                continue
            if kl in PACK_KW and 'Packaging' in ix:
                if not get('Packaging'):
                    put('Packaging', PACK_KW[kl], 'emballage_keyword_til_packaging')
                    continue
                if get('Packaging').lower() == PACK_KW[kl].lower():
                    continue
            if re.search(r'%|prozent', kl) and not re.search(r'fett|alkohol|vol', kl):
                stats['keyword_ingrediensprocent_fjernet'] += 1
                continue
            if kl in THEME or kl == (get('Category') or '').lower():
                stats['keyword_er_kategori_fjernet'] += 1
                continue
            toks = [w.lower().strip(',.') for w in re.split(r'[\s-]+', kwv) if len(w) >= 3]
            if toks and all(t in own_words for t in toks):
                stats['keyword_dobbeltord_fjernet'] += 1
                continue
            if re.match(r'(?i)^(' + FUNC[lang] + r')\b', kwv) and kwv[0].isupper():
                kwv = lower1(kwv)
            if kwv.lower() in {k.lower() for k in kept}:
                continue
            kept.append(kwv)
        kws_new = (kept + [''] * 5)[:5]
        # --- skriv felter ---
        if (pt or None) != (pt0 or None):
            put('Product Type', pt or None, pt_rule)
        if (var or None) != (var0 or None):
            put('Variation', var or None, 'variation_oprydning')
        for n in range(1, 6):
            if (kws_new[n - 1] or None) != (kws0[n - 1] or None):
                put(f'Keyword {n}', kws_new[n - 1] or None, 'keywords_oprydning')
        # --- flag fra tekst (kun tomme felter, kun titel/variation/keywords) ---
        txt = ' | '.join([ot, pn, str(new[I('Variation')] or '')] + [str(new[I(f'Keyword {n}')] or '') for n in range(1, 6)])
        if not get('_is_organic') and ORGANIC_RE.search(txt) and lang == 'de':
            put('_is_organic', W['organic'], 'is_bio_fra_tekst')
        elif not get('_is_organic') and ORGANIC_RE.search(txt) and lang == 'da':
            put('_is_organic', W['organic'], 'is_oeko_fra_tekst')
        for col, wk, rx in IS_FLAG:
            if col in ix and not get(col) and rx.search(txt):
                put(col, W[wk], 'is_fra_tekst:' + col)
        if not get('_is_sugar_free'):
            if SUGAR_FREE_RE.search(txt):
                put('_is_sugar_free', W['sugar'], 'is_fra_tekst:_is_sugar_free')
            elif SUGAR_ADDED_RE.search(txt):
                put('_is_sugar_free', W['sugar_added'], 'is_fra_tekst:_is_sugar_free')
        if not get('_is_alcohol_free') and ALCFREE_RE.search(txt):
            put('_is_alcohol_free', W['alcfree'], 'is_fra_tekst:_is_alcohol_free')
        if not get('_is_fat'):
            m = FAT_EXPLICIT.search(ot)
            if m:
                put('_is_fat', m.group(1).replace(',', '.') + '%', 'fedtprocent_fra_titel')
            elif FAT_DAIRY_WORD.search(ot) and not FAT_NOT.search(ot):
                m = FAT_DAIRY_PCT.search(ot)
                if m:
                    put('_is_fat', m.group(1).replace(',', '.') + '%', 'fedtprocent_mejeri_fra_titel')
        if lang == 'de' and '_is_animal_welfare' in ix and not get('_is_animal_welfare'):
            m = WELFARE_RE.search(ot + ' ' + ' '.join(kws0))
            if m:
                put('_is_animal_welfare', f'haltungsform {m.group(1)}', 'is_fra_tekst:_is_animal_welfare')
        if lang == 'de' and not get('_is_country_of_origen'):
            m = COUNTRY_RE.search(ot)
            if m:
                put('_is_country_of_origen', next(c for c in COUNTRY_DE if c.lower() == m.group(1).lower()), 'is_fra_tekst:_is_country_of_origen')
        if not get('_is_frozen') and (FROZEN_CAT.search(cat) or FROZEN_TITLE.search(ot + ' ' + pn)):
            put('_is_frozen', W['frozen'], 'frost_fra_kategori_eller_titel')
        if get('_is_alcohol', cpct) and not get('_is_alcohol', alc_cols[0]) and re.search(r'[1-9]', str(get('_is_alcohol', cpct))):
            put('_is_alcohol', W['alcohol'], 'alkoholflag_fra_procent', k=alc_cols[0])
        # Clean Subtitle = brand | quantity (Bilkas moenster) - kun hvor den fulgte moensteret
        if 'Clean Subtitle' in ix:
            b0, q0 = vals[I('Brand')], vals[I('Quantity')]
            exp0 = ' | '.join(str(x) for x in (b0, q0) if x)
            exp1 = ' | '.join(str(x) for x in (get('Brand'), get('Quantity')) if x)
            cs = get('Clean Subtitle')
            if exp1 != exp0 and (cs in (None, '') or cs == exp0 or cs == (str(b0) if b0 else None)):
                put('Clean Subtitle', exp1 or None, 'clean_subtitle_opdateret')
        # Wolt: flertal af produkttype genberegnes naar typen er aendret
        if lang == 'da' and (pt or None) != (pt0 or None) and pt:
            master, fr = plural_ctx
            pl = O.plural_for(pt, master, fr)
            if pl:
                pl = pt if pl.lower() == pt.lower() else (pl[0].upper() if pt[0].isupper() else pl[0].lower()) + pl[1:]
            put('Product type plural', pl, 'flertal_genberegnet')
        for i in range(len(new)):
            if new[i] != vals[i]:
                row[i].value = new[i]
    wb.calculation.fullCalcOnLoad = True
    fd, tmp = tempfile.mkstemp(suffix='.xlsx', dir=os.environ.get('TEMP'))
    os.close(fd)
    wb.save(tmp)
    wb.close()
    if dry:
        print(f'{sh.name}: proeve -> {tmp}')
    else:
        if os.path.getmtime(path) != mtime or os.path.exists(lock):
            print(f'{sh.name}: SPRUNGET OVER - {os.path.basename(path)} blev aendret/aabnet imens - koer igen')
            os.remove(tmp)
            return None
        os.makedirs(BACKUP, exist_ok=True)
        stamp = datetime.datetime.now().strftime('%Y-%m-%d_%H%M')
        shutil.copy2(path, os.path.join(BACKUP, f'{os.path.splitext(os.path.basename(path))[0]}_{stamp}_foer_optimering.xlsx'))
        shutil.move(tmp, path)
        print(f'{sh.name}: gemt -> {path}')
    for k, v in sorted(stats.items()):
        print(f'   {k:<48}{v:>6}')
    report[sh.name] = (dict(stats), dict(examples), len(rows))
    return tmp if dry else path


def build_brands():
    dk = TB.danish_brands()  # norm -> (stavemaade, {sub})
    spell = collections.defaultdict(collections.Counter)
    for n, (p, lang) in SHEETS.items():
        if lang != 'de' or not os.path.exists(p):
            continue
        ws = openpyxl.load_workbook(p, read_only=True).active
        it = ws.iter_rows(values_only=True)
        h = [str(x).strip() if x else x for x in next(it)]
        bi = h.index('Brand')
        for r in it:
            b = r[bi] if bi < len(r) else None
            if isinstance(b, str) and b.strip() and len(b.strip()) <= 45:
                k = TB.norm(b)
                if len(k) >= 3 and k not in TB.STOP and re.search(r'[a-z]', k):
                    spell[k][b.strip()] += 1
    out = dict(dk)
    for k, c in spell.items():
        tot = sum(c.values())
        if k not in out and tot >= 10 and k not in {'demeter', 'bioland', 'naturland', 'hokkaido', 'dryk', 'drik'}:
            out[k] = (c.most_common(1)[0][0], {})
    # almindelige ord ("Frisk", "Original", "Gode") er ikke maerker: enkeltords-noegler der staar oftere i
    # produkttype/variation end som brand i arkene frasorteres
    pv, bc = collections.Counter(), collections.Counter()
    paths = [p for p in TB.DANISH if os.path.exists(p)] + [p for p, lang in SHEETS.values() if lang == 'de' and os.path.exists(p)]
    for path in paths:
        try:
            ws = openpyxl.load_workbook(path, read_only=True).active
            it = ws.iter_rows(values_only=True)
            h = [str(x).strip() if x else x for x in next(it)]
            ci = {n: h.index(n) for n in ('Brand', 'Product Type', 'Variation') if n in h}
            if 'Brand' not in ci:
                continue
            for r in it:
                b = r[ci['Brand']] if ci['Brand'] < len(r) else None
                if isinstance(b, str) and b.strip():
                    bc[TB.norm(b)] += 1
                txt = ' '.join(str(r[ci[n]]) for n in ('Product Type', 'Variation') if n in ci and ci[n] < len(r) and r[ci[n]])
                for t in set(TB.norm(txt).split()):
                    pv[t] += 1
        except Exception as e:  # et ark der er laast/aaben maa ikke vaelte kørslen
            print('   (kunne ikke laese', path, e, ')')
    dropped = []
    for k in list(out):
        if (' ' not in k and pv[k] >= 5 and pv[k] > 2 * bc[k]) or k in {'demeter', 'bioland', 'naturland', 'hokkaido', 'dryk', 'drik'}:
            dropped.append(k)
            del out[k]
    print(f'   {len(dropped)} almindelige ord frasorteret som maerker, fx {dropped[:12]}')
    return out


def main(args):
    dry = '--dry' in args
    names = [a for a in args if not a.startswith('--')] or list(SHEETS)
    print('Maerker ...')
    brands = build_brands()
    by_len = sorted(brands, key=len, reverse=True)
    print(f'   {len(brands)} maerker')
    master, fr = O.bilka_master()[3], T.load_frida_plurals()
    audit, report = [], {}
    for n in names:
        p, lang = SHEETS[n]
        run_sheet(Sheet(n, p, lang), brands, by_len, (master, fr), dry, audit, report)
    if audit:
        out = os.path.join(BACKUP if not dry else (os.environ.get('TEMP') or tempfile.gettempdir()),
                           f'optimering_{datetime.datetime.now():%Y-%m-%d_%H%M}.csv')
        os.makedirs(os.path.dirname(out), exist_ok=True)
        with open(out, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(['butik', 'raekke', 'kolonne', 'foer', 'efter', 'regel'])
            w.writerows(audit)
        print(f'Alle aendringer ({len(audit)}) -> {out}')
    import json
    rp = os.path.join(os.environ.get('TEMP') or tempfile.gettempdir(), 'optimering_rapport.json')
    with open(rp, 'w', encoding='utf-8') as f:
        json.dump(report, f, ensure_ascii=False, indent=1, default=str)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main(sys.argv[1:])
